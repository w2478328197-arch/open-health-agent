from __future__ import annotations

import json
import os
import sqlite3
import subprocess
import sys
import threading
from argparse import Namespace
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

import health_agent as health_cli
import oha.profile as profile_module
from oha.database import HealthDatabase
from oha.config import atomic_write_json, load_config


ROOT = Path(__file__).resolve().parents[1]
SKILL_ROOT = ROOT / "skills" / "open-health-agent"
CLI = SKILL_ROOT / "scripts" / "health_agent.py"
TEMPLATE = SKILL_ROOT / "assets" / "health-ledger.xlsx"
FIXTURES = Path(__file__).parent / "fixtures" / "ghealth"


def assert_private_mode(path: Path) -> None:
    assert path.exists()
    if os.name != "nt":
        assert path.stat().st_mode & 0o777 == 0o600


def run_cli(*arguments: str, expected_code: int = 0) -> tuple[subprocess.CompletedProcess[str], dict]:
    environment = os.environ.copy()
    if "--fixture-dir" in arguments:
        environment["OPEN_HEALTH_AGENT_TEST_MODE"] = "1"
    completed = subprocess.run(
        [sys.executable, str(CLI), *arguments],
        cwd=ROOT,
        check=False,
        capture_output=True,
        text=True,
        env=environment,
    )
    assert completed.returncode == expected_code, completed.stderr or completed.stdout
    stream = completed.stdout if expected_code == 0 else completed.stderr
    return completed, json.loads(stream)


def test_cli_help_lists_core_workflows() -> None:
    completed = subprocess.run(
        [sys.executable, str(CLI), "--help"],
        cwd=ROOT,
        check=False,
        capture_output=True,
        text=True,
    )
    assert completed.returncode == 0
    for command in ("init", "sync", "record", "goal", "context", "scheduler"):
        assert command in completed.stdout

    record_help = subprocess.run(
        [sys.executable, str(CLI), "record", "--help"],
        cwd=ROOT,
        check=False,
        capture_output=True,
        text=True,
    )
    assert record_help.returncode == 0
    for field in ("second_value", "consumed=true", "summary_nutrients", "original_text"):
        assert field in record_help.stdout


def test_scheduler_install_uses_fresh_one_shot_consent_and_redacts_paths(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        health_cli,
        "require_scheduler_install_convergence",
        lambda *_args, **_kwargs: None,
    )
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    fingerprint = "a" * 64
    static_fingerprint = "c" * 64
    batch_id = "sync_ghealth_manual_20260101T000000Z_synthetic"
    with HealthDatabase(home / "health.sqlite3", create=False) as database:
        database.begin_sync(batch_id, "2026-01-01", "2026-01-02")
        database.finish_sync(
            batch_id,
            "success",
            1,
            0,
            0,
            None,
            [],
        )
        persisted = database.get_sync_run(batch_id)
    assert persisted and persisted["finished_at"]
    atomic_write_json(
        home / "state.json",
        {
            "onboarding_explained_at": now,
            "onboarding_explanation_delivery_confirmed_at": now,
            "onboarding_consents": {
                "google-health": {"granted_at": now, "policy_version": 1},
                "scheduler": {
                    "granted_at": persisted["finished_at"],
                    "policy_version": 1,
                },
            },
            "last_manual_ghealth_sync": {
                "version": 2,
                "batch_id": batch_id,
                "status": "success",
                "finished_at": persisted["finished_at"],
                "timezone": "UTC",
                "static_runtime_fingerprint": static_fingerprint,
                "runtime_fingerprint": fingerprint,
            },
        },
    )
    executable = tmp_path / "ghealth"
    executable.write_text("synthetic", encoding="utf-8")
    executable.chmod(0o700)
    monkeypatch.setattr(
        health_cli, "resolve_ghealth_executable", lambda _command: executable
    )
    monkeypatch.setattr(
        health_cli, "ghealth_runtime_fingerprint", lambda _config: fingerprint
    )
    monkeypatch.setattr(
        health_cli, "static_runtime_fingerprint", lambda _config: static_fingerprint
    )

    class MatchingAdapter:
        def require_matching_timezone(self) -> None:
            return None

    monkeypatch.setattr(
        health_cli,
        "GHealthAdapter",
        lambda *_args, **_kwargs: MatchingAdapter(),
    )
    calls: list[dict] = []

    def fake_install(_config, interval, **kwargs):
        calls.append({"interval": interval, **kwargs})
        return {"backend": "launchd", "path": "/synthetic/private/job.plist"}

    monkeypatch.setattr(health_cli, "install_scheduler", fake_install)
    arguments = Namespace(
        home=str(home),
        action="install",
        interval_seconds=3600,
        inherit_proxy_env=False,
        proxy_env_file=None,
        verbose_paths=False,
    )

    result = health_cli.command_scheduler(arguments)

    assert result == {"backend": "launchd"}
    assert (
        calls[0]["expected_static_runtime_fingerprint"]
        == static_fingerprint
    )
    assert calls[0]["expected_runtime_fingerprint"] == fingerprint
    state = json.loads((home / "state.json").read_text(encoding="utf-8"))
    assert state["onboarding_consents"]["scheduler"]["consumed_at"]
    authorization = state["onboarding_consents"]["scheduler"][
        "installed_authorization"
    ]
    assert authorization["static_runtime_fingerprint"] == static_fingerprint
    assert authorization["runtime_fingerprint"] == fingerprint
    with pytest.raises(ValueError, match="fresh scheduler consent"):
        health_cli.command_scheduler(arguments)


def test_scheduler_consent_must_follow_the_qualifying_manual_sync(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        health_cli,
        "require_scheduler_install_convergence",
        lambda *_args, **_kwargs: None,
    )
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    old_grant = (datetime.now(timezone.utc) - timedelta(days=1)).isoformat(
        timespec="seconds"
    )
    completed_at = datetime.now(timezone.utc).isoformat(timespec="seconds")
    recent_pre_sync_grant = (
        datetime.now(timezone.utc) - timedelta(seconds=5)
    ).isoformat(timespec="seconds")
    atomic_write_json(
        home / "state.json",
        {
            "onboarding_explanation_delivery_confirmed_at": old_grant,
            "onboarding_consents": {
                "google-health": {"granted_at": old_grant, "policy_version": 1},
                "scheduler": {
                    "granted_at": recent_pre_sync_grant,
                    "policy_version": 1,
                },
            },
        },
    )
    monkeypatch.setattr(
        health_cli,
        "_recent_manual_ghealth_sync_fingerprint",
        lambda _config: {
            "static_runtime_fingerprint": "c" * 64,
            "runtime_fingerprint": "a" * 64,
            "finished_at": completed_at,
        },
    )
    installs: list[str] = []
    monkeypatch.setattr(
        health_cli,
        "install_scheduler",
        lambda *_args, **_kwargs: installs.append("installed") or {},
    )

    with pytest.raises(ValueError, match="after the qualifying manual sync"):
        health_cli.command_scheduler(
            Namespace(
                home=str(home),
                action="install",
                interval_seconds=3600,
                inherit_proxy_env=False,
                proxy_env_file=None,
                verbose_paths=False,
            )
        )

    assert installs == []
    state = json.loads((home / "state.json").read_text(encoding="utf-8"))
    assert "consumed_at" not in state["onboarding_consents"]["scheduler"]


def test_legacy_explanation_timestamp_cannot_authorize_new_scopes(tmp_path: Path) -> None:
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    atomic_write_json(
        home / "state.json",
        {"onboarding_explained_at": "2026-01-01T00:00:00+00:00"},
    )

    _, failure = run_cli(
        "--home",
        str(home),
        "onboarding",
        "grant-consent",
        "--scope",
        "scheduler",
        expected_code=1,
    )

    assert failure["code"] == "invalid_input"
    assert not health_cli._has_scoped_consent(load_config(home), "scheduler")


def test_scheduler_rejects_expired_manual_sync_evidence(tmp_path: Path) -> None:
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    expired = (datetime.now(timezone.utc) - timedelta(hours=25)).isoformat(
        timespec="seconds"
    )
    atomic_write_json(
        home / "state.json",
        {
            "last_manual_ghealth_sync": {
                "version": 2,
                "batch_id": "sync_ghealth_manual_20260101T000000Z_expired",
                "status": "success",
                "finished_at": expired,
                "timezone": "UTC",
                "static_runtime_fingerprint": "c" * 64,
                "runtime_fingerprint": "a" * 64,
            },
        },
    )

    assert health_cli._recent_manual_ghealth_sync_fingerprint(load_config(home)) is None


def test_scheduler_gate_reports_a_stable_actionable_error() -> None:
    for error in (
        ValueError(
            "a successful real manual ghealth sync with the current profile and timezone is required within 30 minutes before scheduler install"
        ),
        RuntimeError(
            "the current ghealth executable, profile, or timezone differs from the last successful manual sync"
        ),
    ):
        payload = health_cli._public_exception_payload(error)
        assert payload["code"] == "recent_manual_ghealth_sync_required"
        assert "30 minutes" in payload["error"]
        assert "Fixture and scheduled runs do not qualify" in payload["error"]


def test_recent_manual_sync_evidence_requires_the_current_runtime(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    batch_id = "sync_ghealth_manual_20260101T000000Z_synthetic"
    with HealthDatabase(home / "health.sqlite3", create=False) as database:
        database.begin_sync(batch_id, "2026-01-01", "2026-01-02")
        database.finish_sync(batch_id, "success", 1, 0, 0, None, [])
        persisted = database.get_sync_run(batch_id)
    assert persisted and persisted["finished_at"]
    atomic_write_json(
        home / "state.json",
        {
            "last_manual_ghealth_sync": {
                "version": 2,
                "batch_id": batch_id,
                "status": "success",
                "finished_at": persisted["finished_at"],
                "timezone": "UTC",
                "static_runtime_fingerprint": "c" * 64,
                "runtime_fingerprint": "a" * 64,
            }
        },
    )
    monkeypatch.setattr(
        health_cli, "ghealth_runtime_fingerprint", lambda _config: "b" * 64
    )
    monkeypatch.setattr(
        health_cli, "static_runtime_fingerprint", lambda _config: "c" * 64
    )

    assert health_cli._recent_manual_ghealth_sync_fingerprint(load_config(home)) is None


def test_revoking_scheduler_consent_stops_the_managed_job_first(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    atomic_write_json(
        home / "state.json",
        {
            "onboarding_explanation_delivery_confirmed_at": now,
            "onboarding_consents": {
                "scheduler": {"granted_at": now, "policy_version": 1}
            },
        },
    )
    calls: list[str] = []
    monkeypatch.setattr(
        health_cli, "scheduler_status", lambda _config: {"backend": "launchd"}
    )
    def successful_removal(_config):
        state_during_removal = json.loads(
            (home / "state.json").read_text(encoding="utf-8")
        )
        assert state_during_removal["scheduler_job_removal_pending"] is True
        calls.append("uninstalled")
        return {"removed": True}

    monkeypatch.setattr(health_cli, "uninstall_scheduler", successful_removal)

    result = health_cli.command_onboarding(
        Namespace(
            home=str(home),
            action="revoke-consent",
            scope="scheduler",
            delivery_confirmed=False,
        )
    )

    assert calls == ["uninstalled"]
    assert result == {
        "status": "consent_revoked",
        "scope": "scheduler",
        "job_removal": "completed",
    }
    state = json.loads((home / "state.json").read_text(encoding="utf-8"))
    assert state["onboarding_consents"]["scheduler"]["revoked_at"]
    assert state["scheduler_job_removal_pending"] is False
    assert state["scheduler_job_removal_error_code"] is None
    assert not health_cli._has_scoped_consent(load_config(home), "scheduler")


def test_new_scheduler_grant_cannot_reauthorize_an_orphaned_revoked_job(
    tmp_path: Path,
) -> None:
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    old_static = "c" * 64
    old_runtime = "a" * 64
    atomic_write_json(
        home / "state.json",
        {
            "onboarding_explanation_delivery_confirmed_at": now,
            "onboarding_consents": {
                "google-health": {"granted_at": now, "policy_version": 1},
                "scheduler": {
                    "granted_at": now,
                    "consumed_at": now,
                    "revoked_at": now,
                    "policy_version": 1,
                    "installed_authorization": {
                        "version": 1,
                        "authorized_at": now,
                        "static_runtime_fingerprint": old_static,
                        "runtime_fingerprint": old_runtime,
                    },
                },
            },
        },
    )

    granted = health_cli.command_onboarding(
        Namespace(
            home=str(home),
            action="grant-consent",
            scope="scheduler",
            delivery_confirmed=False,
        )
    )

    assert granted["consents"]["scheduler"]["active"] is True
    state = json.loads((home / "state.json").read_text(encoding="utf-8"))
    assert "installed_authorization" not in state["onboarding_consents"]["scheduler"]
    assert not health_cli._has_scheduled_runtime_authorization(
        load_config(home), old_static, old_runtime
    )
    with pytest.raises(ValueError, match="successful installed runtime"):
        health_cli.command_sync(
            Namespace(
                home=str(home),
                from_date="2026-01-01",
                to_date="2026-01-02",
                lookback_days=None,
                fixture_dir=None,
                scheduled=True,
                static_runtime_fingerprint=old_static,
                runtime_fingerprint=old_runtime,
                quiet=True,
                verbose_path=False,
            )
        )


def test_scheduler_revocation_persists_when_job_removal_needs_retry(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    atomic_write_json(
        home / "state.json",
        {
            "onboarding_consents": {
                "scheduler": {"granted_at": now, "policy_version": 1}
            }
        },
    )
    monkeypatch.setattr(
        health_cli, "scheduler_status", lambda _config: {"backend": "launchd"}
    )

    def removal_failure(_config):
        raise RuntimeError("launchctl synthetic failure with /private/path")

    monkeypatch.setattr(health_cli, "uninstall_scheduler", removal_failure)

    result = health_cli.command_onboarding(
        Namespace(
            home=str(home),
            action="revoke-consent",
            scope="scheduler",
            delivery_confirmed=False,
        )
    )

    assert result["status"] == "consent_revoked_job_removal_pending"
    assert result["job_removal"] == "pending"
    assert "/private/path" not in json.dumps(result)
    assert not health_cli._has_scoped_consent(load_config(home), "scheduler")
    state = json.loads((home / "state.json").read_text(encoding="utf-8"))
    assert state["scheduler_job_removal_pending"] is True
    assert state["scheduler_job_removal_error_code"] == "scheduler_operation_failed"
    _, doctor = run_cli("--home", str(home), "doctor", "--offline")
    removal_check = next(
        item
        for item in doctor["checks"]
        if item["check"] == "scheduler job removal complete"
    )
    assert removal_check["ok"] is False
    assert removal_check["detail"]["error_code"] == "scheduler_operation_failed"

    monkeypatch.setattr(
        health_cli,
        "uninstall_scheduler",
        lambda _config: {"backend": "launchd", "removed": True},
    )
    retried = health_cli.command_scheduler(
        Namespace(
            home=str(home),
            action="uninstall",
            verbose_paths=False,
            proxy_env_file=None,
            inherit_proxy_env=False,
        )
    )
    assert retried["consent_revoked"] is True
    state = json.loads((home / "state.json").read_text(encoding="utf-8"))
    assert state["scheduler_job_removal_pending"] is False
    assert state["scheduler_job_removal_error_code"] is None


def test_scheduler_install_never_starts_if_one_shot_consumption_fails(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(
        health_cli,
        "require_scheduler_install_convergence",
        lambda *_args, **_kwargs: None,
    )
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    atomic_write_json(
        home / "state.json",
        {
            "onboarding_explanation_delivery_confirmed_at": now,
            "onboarding_consents": {
                "google-health": {"granted_at": now, "policy_version": 1},
                "scheduler": {"granted_at": now, "policy_version": 1},
            },
        },
    )
    executable = tmp_path / "ghealth"
    executable.write_text("synthetic", encoding="utf-8")
    executable.chmod(0o700)
    monkeypatch.setattr(
        health_cli,
        "_recent_manual_ghealth_sync_fingerprint",
        lambda _config: {
            "static_runtime_fingerprint": "c" * 64,
            "runtime_fingerprint": "a" * 64,
            "finished_at": now,
        },
    )
    monkeypatch.setattr(
        health_cli, "resolve_ghealth_executable", lambda _command: executable
    )

    class MatchingAdapter:
        def require_matching_timezone(self) -> None:
            return None

    monkeypatch.setattr(
        health_cli, "GHealthAdapter", lambda *_args, **_kwargs: MatchingAdapter()
    )
    installs: list[str] = []
    monkeypatch.setattr(
        health_cli,
        "install_scheduler",
        lambda *_args, **_kwargs: installs.append("installed") or {},
    )

    def failed_consumption(_config, _scope):
        raise OSError("synthetic state write failure")

    monkeypatch.setattr(health_cli, "_consume_scoped_consent", failed_consumption)

    with pytest.raises(OSError, match="state write failure"):
        health_cli.command_scheduler(
            Namespace(
                home=str(home),
                action="install",
                interval_seconds=3600,
                inherit_proxy_env=False,
                proxy_env_file=None,
                verbose_paths=False,
            )
        )

    assert installs == []


def test_emit_round_trips_unicode_through_an_ascii_only_console(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    class AsciiOnlyStream:
        def __init__(self) -> None:
            self.value = ""

        def write(self, value: str) -> int:
            value.encode("ascii")
            self.value += value
            return len(value)

        def flush(self) -> None:
            return None

    stream = AsciiOnlyStream()
    monkeypatch.setattr(sys, "stdout", stream)

    health_cli.emit({"工作簿": "健康档案.xlsx"})

    assert json.loads(stream.value) == {"工作簿": "健康档案.xlsx"}


@pytest.mark.skipif(not TEMPLATE.exists(), reason="workbook asset is still being generated")
def test_cli_private_end_to_end_with_synthetic_ghealth(tmp_path: Path) -> None:
    home = tmp_path / "private-health-home"
    workbook_path = tmp_path / "健康档案.xlsx"
    base = ("--home", str(home))

    _, initialized = run_cli(
        *base,
        "init",
        "--workbook",
        str(workbook_path),
        "--timezone",
        "UTC",
        "--ghealth-command",
        "missing-ghealth-for-synthetic-test",
        "--verbose-paths",
    )
    assert initialized["status"] == "ready"
    assert Path(initialized["database"]).exists()
    assert workbook_path.exists()
    assert_private_mode(home / "AGENTS.md")
    assert_private_mode(home / "profile.json")

    _, doctor = run_cli(*base, "doctor", "--offline")
    assert doctor["status"] == "ok"
    ghealth_check = next(
        item for item in doctor["checks"] if item["check"] == "ghealth command"
    )
    assert ghealth_check["ok"] is False
    assert ghealth_check["required"] is False
    assert next(
        item for item in doctor["checks"] if item["check"] == "workbook managed schema"
    )["ok"] is True
    assert next(
        item for item in doctor["checks"] if item["check"] == "goal projections"
    )["ok"] is True

    private_before = {
        path.name: path.read_bytes()
        for path in (
            home / "config.json",
            home / "AGENTS.md",
            home / "profile.json",
            home / "state.json",
            home / "health.sqlite3",
            workbook_path,
        )
    }
    _, repeated_init = run_cli(*base, "init")
    assert repeated_init["existing_configuration_reused"] is True
    assert repeated_init["workbook_exported"] is False
    assert "database" not in repeated_init
    assert str(tmp_path) not in json.dumps(repeated_init)
    private_after = {
        path.name: path.read_bytes()
        for path in (
            home / "config.json",
            home / "AGENTS.md",
            home / "profile.json",
            home / "state.json",
            home / "health.sqlite3",
            workbook_path,
        )
    }
    assert private_after == private_before

    _, consent_error = run_cli(*base, "onboarding", "grant-consent", expected_code=1)
    assert consent_error["code"] == "invalid_input"
    assert "Skill explanation" in consent_error["error"]
    _, unconfirmed = run_cli(
        *base, "onboarding", "mark-explained", expected_code=1
    )
    assert unconfirmed["code"] == "explanation_delivery_unconfirmed"
    run_cli(*base, "onboarding", "mark-explained", "--delivery-confirmed")
    _, missing_scope = run_cli(
        *base, "onboarding", "grant-consent", expected_code=1
    )
    assert missing_scope["code"] == "consent_scope_required"
    _, consent = run_cli(
        *base,
        "onboarding",
        "grant-consent",
        "--scope",
        "google-health",
    )
    assert consent["explanation_delivered"] is True
    assert consent["consents"]["google-health"]["active"] is True
    assert consent["consents"]["google-health"]["policy_version"] == 1

    for _ in range(2):
        _, synced = run_cli(
            *base,
            "sync",
            "--from-date",
            "2026-01-01",
            "--to-date",
            "2026-01-03",
            "--fixture-dir",
            str(FIXTURES),
        )
        assert synced["status"] == "success"
        assert synced["counts"] == {"daily": 1, "measurements": 1, "workouts": 1}
    fixture_state = json.loads((home / "state.json").read_text(encoding="utf-8"))
    assert "last_manual_ghealth_sync" not in fixture_state

    food_payload = {
        "consumed": True,
        "date": "2026-01-02",
        "time": "12:30",
        "meal": "午餐",
        "food_name": "合成测试餐",
        "grams_low": 250,
        "grams_high": 350,
        "summary_nutrients": {
            "energy_kcal": 600,
            "protein_g": 35,
            "fat_g": 20,
            "carbohydrate_g": 70,
        },
        "source": "synthetic test estimate",
        "original_text": "我吃了这份合成测试餐",
    }
    _, recorded = run_cli(
        *base,
        "record",
        "--kind",
        "food",
        "--json",
        json.dumps(food_payload, ensure_ascii=False),
    )
    assert recorded["record"] == {"record_id": recorded["record"]["record_id"]}

    exact_goal = "合成测试目标：每周规律完成训练"
    _, goal = run_cli(
        *base,
        "goal",
        "set",
        "--text",
        exact_goal,
        "--priority",
        "1",
        "--effective-date",
        "2026-01-02",
    )
    assert goal["goal"]["status"] == "active"
    assert exact_goal in (home / "AGENTS.md").read_text(encoding="utf-8")

    run_cli(*base, "profile", "set", "--key", "lean_mass_kg", "--value", "60")
    _, context = run_cli(*base, "context", "--date", "2026-01-02")
    assert context["profile"]["active_goals"][0]["original_text"] == exact_goal
    assert context["today"]["health"]["steps"] == 4321
    assert context["today"]["food_items"][0]["food_name"] == "合成测试餐"
    assert context["energy"]["estimated_ree_kcal"] == 1666

    connection = sqlite3.connect(home / "health.sqlite3")
    try:
        counts = dict(
            connection.execute(
                "SELECT kind, COUNT(*) FROM records GROUP BY kind"
            ).fetchall()
        )
    finally:
        connection.close()
    assert counts["daily"] == 1
    assert counts["measurement"] == 1
    assert counts["workout"] == 1
    assert counts["food"] == 1
    assert counts["goal"] == 1

    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        assert workbook["健康日报"].max_row == 2
        assert workbook["健康测量"].max_row == 2
        assert workbook["训练记录"].max_row == 2
        assert workbook["饮食记录"].max_row == 2
        assert workbook["目标历史"].max_row == 2
    finally:
        workbook.close()


@pytest.mark.parametrize(
    ("key", "value"),
    [
        ("timezone", '"Not/A_Real_Timezone"'),
        ("activity_energy_semantics", '"ambiguous"'),
    ],
)
def test_invalid_operational_profile_value_leaves_no_false_pending_marker(
    tmp_path: Path, key: str, value: str
) -> None:
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    config_before = (home / "config.json").read_bytes()
    profile_before = (home / "profile.json").read_bytes()
    state_before = (home / "state.json").read_bytes()

    _, error = run_cli(
        "--home",
        str(home),
        "profile",
        "set",
        "--key",
        key,
        "--value",
        value,
        expected_code=1,
    )

    assert error["code"] == "invalid_input"
    assert (home / "config.json").read_bytes() == config_before
    assert (home / "profile.json").read_bytes() == profile_before
    assert (home / "state.json").read_bytes() == state_before


def test_cli_profile_partial_success_is_explicit_and_journaled(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    original_profile = (home / "profile.json").read_bytes()
    monkeypatch.setattr(
        profile_module,
        "save_profile",
        lambda _config, _profile: (_ for _ in ()).throw(
            OSError("synthetic projection failure")
        ),
    )

    with pytest.raises(health_cli.ProfileProjectionError) as failure:
        health_cli.command_profile_set(
            Namespace(
                home=str(home),
                key="timezone",
                value='"Asia/Shanghai"',
                file=None,
                source="user-confirmed",
            )
        )

    public = health_cli._public_exception_payload(failure.value)
    assert public["code"] == "profile_projection_pending"
    assert public["operational_config"] == "succeeded"
    assert load_config(home).timezone == "Asia/Shanghai"
    assert (home / "profile.json").read_bytes() == original_profile
    state = json.loads((home / "state.json").read_text(encoding="utf-8"))
    assert state["profile_projection_pending"] is True
    assert state["profile_projection_pending_field"] == "timezone"


def test_cli_goal_partial_success_keeps_sqlite_and_outbox(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    monkeypatch.setattr(
        profile_module,
        "update_agents_goals",
        lambda _config, _goals: (_ for _ in ()).throw(
            OSError("synthetic agents projection failure")
        ),
    )

    with pytest.raises(health_cli.GoalProjectionError) as failure:
        health_cli.command_goal_set(
            Namespace(
                home=str(home),
                text="合成目标投影失败测试",
                file=None,
                effective_date="2026-01-02",
                priority=1,
                safety_constraint="",
            )
        )

    public = health_cli._public_exception_payload(failure.value)
    assert public["code"] == "goal_projection_pending"
    assert public["database_write"] == "succeeded"
    with HealthDatabase(home / "health.sqlite3", create=False) as database:
        goals = database.list_records("goal")
    assert len(goals) == 1
    state = json.loads((home / "state.json").read_text(encoding="utf-8"))
    assert state["workbook_export_pending"] is True
    assert str(state["workbook_export_pending_operation"]).startswith("goal:set:")


@pytest.mark.skipif(not TEMPLATE.exists(), reason="workbook asset is still being generated")
def test_force_init_restores_existing_config_when_export_fails(tmp_path: Path) -> None:
    home = tmp_path / "private-health-home"
    original_workbook = tmp_path / "健康档案.xlsx"
    base = ("--home", str(home))

    _, initialized = run_cli(
        *base,
        "init",
        "--workbook",
        str(original_workbook),
        "--timezone",
        "UTC",
    )
    original_config = json.loads((home / "config.json").read_text(encoding="utf-8"))

    incompatible_workbook = tmp_path / "不兼容健康档案.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "健康日报"
    sheet.append(["错误表头"])
    workbook.save(incompatible_workbook)
    workbook.close()

    _, failed = run_cli(
        *base,
        "init",
        "--force",
        "--workbook",
        str(incompatible_workbook),
        expected_code=1,
    )
    assert failed["code"] == "workbook_incompatible"
    assert failed["type"] == "RuntimeError"
    assert "健康日报" not in json.dumps(failed, ensure_ascii=False)

    restored_config = json.loads((home / "config.json").read_text(encoding="utf-8"))
    assert restored_config == original_config
    assert restored_config["workbook_path"] == str(original_workbook)


@pytest.mark.skipif(not TEMPLATE.exists(), reason="workbook asset is still being generated")
def test_operational_commands_never_recreate_a_missing_canonical_database(
    tmp_path: Path,
) -> None:
    home = tmp_path / "private-health-home"
    workbook_path = tmp_path / "健康档案.xlsx"
    base = ("--home", str(home))
    _, initialized = run_cli(
        *base,
        "init",
        "--workbook",
        str(workbook_path),
        "--timezone",
        "UTC",
    )
    workbook_before = workbook_path.read_bytes()
    (home / "health.sqlite3").unlink()

    for arguments in (
        ("context", "--date", "2026-01-02"),
        (
            "record",
            "--kind",
            "measurement",
            "--json",
            json.dumps(
                {
                    "date": "2026-01-02",
                    "metric": "体重",
                    "value": 70,
                    "unit": "kg",
                }
            ),
        ),
        ("export",),
        ("backup",),
    ):
        _, failure = run_cli(*base, *arguments, expected_code=1)
        assert failure["code"] == "canonical_database_missing"
    assert not (home / "health.sqlite3").exists()
    assert workbook_path.read_bytes() == workbook_before


def test_existing_init_rechecks_database_after_local_file_initialization(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    home = tmp_path / "private-health-home"
    workbook = tmp_path / "health.xlsx"
    run_cli(
        "--home",
        str(home),
        "init",
        "--workbook",
        str(workbook),
        "--timezone",
        "UTC",
    )
    config_before = (home / "config.json").read_bytes()
    workbook_before = workbook.read_bytes()
    original_initialize = health_cli.initialize_local_files

    def remove_database_after_first_check(*args, **kwargs):
        result = original_initialize(*args, **kwargs)
        (home / "health.sqlite3").unlink()
        return result

    monkeypatch.setattr(
        health_cli, "initialize_local_files", remove_database_after_first_check
    )

    with pytest.raises(health_cli.CanonicalDatabaseMissingError):
        health_cli.command_init(
            Namespace(
                home=str(home),
                workbook=None,
                timezone="Asia/Shanghai",
                ghealth_command=None,
                force=True,
                verbose_paths=False,
            )
        )

    assert not (home / "health.sqlite3").exists()
    assert (home / "config.json").read_bytes() == config_before
    assert workbook.read_bytes() == workbook_before


@pytest.mark.skipif(not TEMPLATE.exists(), reason="workbook asset is still being generated")
def test_sqlite_backups_are_consistent_private_and_bounded(tmp_path: Path) -> None:
    home = tmp_path / "private-health-home"
    base = ("--home", str(home))
    run_cli(*base, "init", "--timezone", "UTC")
    config_path = home / "config.json"
    config = json.loads(config_path.read_text(encoding="utf-8"))
    config["backup_retention"] = 2
    config_path.write_text(json.dumps(config), encoding="utf-8")
    if os.name != "nt":
        config_path.chmod(0o600)

    _, manual = run_cli(*base, "backup")
    assert manual == {
        "backup_created": True,
        "retention": 2,
        "status": "backup_created",
    }

    for index in range(3):
        payload = {
            "date": "2026-01-02",
            "time": f"08:0{index}",
            "metric": "synthetic_weight",
            "value": 70 + index,
            "unit": "kg",
            "source_event_id": f"synthetic-message-{index}",
        }
        run_cli(
            *base,
            "record",
            "--kind",
            "measurement",
            "--json",
            json.dumps(payload),
        )

    backups = sorted((home / "backups").glob("health-ledger-db-*.sqlite3"))
    assert len(backups) == 2
    for path in backups:
        assert_private_mode(path)
        with sqlite3.connect(path) as connection:
            assert connection.execute("PRAGMA quick_check").fetchone()[0] == "ok"


@pytest.mark.skipif(not TEMPLATE.exists(), reason="workbook asset is still being generated")
def test_record_correction_reuses_stable_id_and_missing_id_fails(
    tmp_path: Path,
) -> None:
    home = tmp_path / "private-health-home"
    base = ("--home", str(home))
    run_cli(*base, "init", "--timezone", "UTC")
    initial_payload = {
        "date": "2026-01-02",
        "time": "08:00",
        "metric": "体重",
        "value": 70,
        "unit": "kg",
        "original_text": "今早体重七十公斤",
        "source": "weixin",
        "method": "wechat-text",
        "source_event_id": "synthetic-message-1",
        "source_event_item_id": "item-1",
    }
    _, first = run_cli(
        *base,
        "record",
        "--kind",
        "measurement",
        "--json",
        json.dumps(initial_payload, ensure_ascii=False),
    )
    record_id = first["record"]["record_id"]
    _, corrected = run_cli(
        *base,
        "record",
        "--kind",
        "measurement",
        "--record-id",
        record_id,
        "--json",
        json.dumps(initial_payload | {"value": 69.8}, ensure_ascii=False),
    )
    assert corrected["record"]["record_id"] == record_id
    with HealthDatabase(home / "health.sqlite3") as database:
        rows = database.list_records("measurement")
    assert len(rows) == 1
    assert rows[0]["value"] == 69.8

    for provenance_field in ("source_event_id", "source_event_item_id"):
        _, rejected = run_cli(
            *base,
            "record",
            "--kind",
            "measurement",
            "--record-id",
            record_id,
            "--json",
            json.dumps({"value": 69.7, provenance_field: ""}),
            expected_code=1,
        )
        assert rejected["code"] in {
            "conflicting_source_event_id",
            "invalid_source_event_item_id",
        }

    _, missing = run_cli(
        *base,
        "record",
        "--kind",
        "measurement",
        "--record-id",
        "missing-record",
        "--json",
        json.dumps(initial_payload),
        expected_code=1,
    )
    assert missing["code"] == "record_not_found"


@pytest.mark.skipif(not TEMPLATE.exists(), reason="workbook asset is still being generated")
def test_durable_mutations_report_export_pending_instead_of_false_failure(
    tmp_path: Path,
) -> None:
    home = tmp_path / "private-health-home"
    workbook_path = tmp_path / "健康档案.xlsx"
    base = ("--home", str(home))
    run_cli(
        *base,
        "init",
        "--workbook",
        str(workbook_path),
        "--timezone",
        "UTC",
    )

    incompatible = Workbook()
    sheet = incompatible.active
    sheet.title = "健康日报"
    sheet.append(["错误表头"])
    incompatible.save(workbook_path)
    incompatible.close()

    payload = {
        "date": "2026-01-02",
        "metric": "体重",
        "value": 70,
        "unit": "kg",
        "original_text": "合成测试记录",
    }
    _, recorded = run_cli(
        *base,
        "record",
        "--kind",
        "measurement",
        "--json",
        json.dumps(payload, ensure_ascii=False),
    )
    assert recorded["status"] == "recorded_export_pending"
    assert recorded["database_write"] == "succeeded"
    assert recorded["workbook_export"] == "failed"
    record_id = recorded["record"]["record_id"]
    with HealthDatabase(home / "health.sqlite3") as database:
        assert database.get("measurement", record_id) is not None

    _, goal = run_cli(
        *base,
        "goal",
        "set",
        "--text",
        "合成测试目标",
        "--effective-date",
        "2026-01-02",
    )
    assert goal["status"] == "goal_saved_export_pending"
    assert goal["database_write"] == "succeeded"

    _, deleted = run_cli(
        *base,
        "delete",
        "--kind",
        "measurement",
        "--record-id",
        record_id,
        "--reason",
        "synthetic correction",
    )
    assert deleted["status"] == "deleted_export_pending"
    assert deleted["database_write"] == "succeeded"

    _, doctor = run_cli(*base, "doctor", "--offline")
    pending = next(
        item for item in doctor["checks"] if item["check"] == "workbook export current"
    )
    assert doctor["status"] == "needs_attention"
    assert pending["ok"] is False
    assert str(tmp_path) not in json.dumps(doctor)

    workbook_path.unlink()
    _, exported = run_cli(*base, "export")
    assert exported["status"] == "exported"
    _, repaired = run_cli(*base, "doctor", "--offline")
    repaired_pending = next(
        item for item in repaired["checks"] if item["check"] == "workbook export current"
    )
    assert repaired_pending["ok"] is True


@pytest.mark.skipif(not TEMPLATE.exists(), reason="workbook asset is still being generated")
def test_cloud_workbook_exports_require_active_cloud_consent(
    tmp_path: Path,
) -> None:
    home = tmp_path / "private-health-home"
    cloud_dir = (
        tmp_path
        / "Library"
        / "Mobile Documents"
        / "com~apple~CloudDocs"
        / "Open Health Agent"
    )
    cloud_dir.mkdir(parents=True)
    workbook_path = cloud_dir / "health.xlsx"
    base = ("--home", str(home))
    _, initialized = run_cli(
        *base,
        "init",
        "--workbook",
        str(workbook_path),
        "--timezone",
        "UTC",
    )
    assert initialized["workbook_export"] == "blocked_by_consent"
    assert initialized["required_consent_scope"] == "cloud-workbook"
    assert not workbook_path.exists()
    payload = {
        "date": "2026-01-02",
        "metric": "synthetic",
        "value": 1,
        "unit": "test",
    }

    _, recorded = run_cli(
        *base,
        "record",
        "--kind",
        "measurement",
        "--json",
        json.dumps(payload),
    )

    assert recorded["status"] == "recorded_export_pending"
    assert recorded["workbook_export"] == "blocked_by_consent"
    assert recorded["required_consent_scope"] == "cloud-workbook"
    assert not workbook_path.exists()
    state = json.loads((home / "state.json").read_text(encoding="utf-8"))
    assert state["workbook_export_pending"] is True
    assert state["workbook_export_blocked_by_consent"] is True
    with HealthDatabase(home / "health.sqlite3", create=False) as database:
        assert len(database.list_records("measurement")) == 1

    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    state["onboarding_explanation_delivery_confirmed_at"] = now
    atomic_write_json(home / "state.json", state)
    health_cli.command_onboarding(
        Namespace(
            home=str(home),
            action="grant-consent",
            scope="cloud-workbook",
            delivery_confirmed=False,
        )
    )
    _, exported = run_cli(*base, "export")

    assert exported["status"] == "exported"
    assert workbook_path.is_file()
    state = json.loads((home / "state.json").read_text(encoding="utf-8"))
    assert state["workbook_export_pending"] is False
    assert state["workbook_export_blocked_by_consent"] is False
    assert state["workbook_export_pending_operation"] is None


@pytest.mark.skipif(not TEMPLATE.exists(), reason="workbook asset is still being generated")
def test_cloud_workbook_consent_is_bound_to_the_exact_target(
    tmp_path: Path,
) -> None:
    home = tmp_path / "private-health-home"
    cloud_root = (
        tmp_path
        / "Library"
        / "Mobile Documents"
        / "com~apple~CloudDocs"
        / "Open Health Agent"
    )
    first_workbook = cloud_root / "first.xlsx"
    second_workbook = cloud_root / "second.xlsx"
    base = ("--home", str(home))
    run_cli(*base, "init", "--workbook", str(first_workbook), "--timezone", "UTC")
    run_cli(*base, "onboarding", "mark-explained", "--delivery-confirmed")
    _, granted = run_cli(
        *base, "onboarding", "grant-consent", "--scope", "cloud-workbook"
    )

    assert granted["consents"]["cloud-workbook"]["active"] is True
    assert health_cli._has_scoped_consent(load_config(home), "cloud-workbook")

    _, reconfigured = run_cli(
        *base,
        "init",
        "--force",
        "--workbook",
        str(second_workbook),
    )

    assert reconfigured["workbook_export"] == "blocked_by_consent"
    assert reconfigured["required_consent_scope"] == "cloud-workbook"
    assert not health_cli._has_scoped_consent(load_config(home), "cloud-workbook")


def test_cloud_private_home_consent_is_bound_to_the_exact_provider(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    state = json.loads((home / "state.json").read_text(encoding="utf-8"))
    state["onboarding_explanation_delivery_confirmed_at"] = now
    atomic_write_json(home / "state.json", state)

    def provider_a(path: Path) -> str | None:
        resolved = Path(path).resolve()
        return (
            "iCloud Drive"
            if resolved == home.resolve() or home.resolve() in resolved.parents
            else None
        )

    monkeypatch.setattr(health_cli, "_sync_storage_provider", provider_a)
    health_cli.command_onboarding(
        Namespace(
            home=str(home),
            action="grant-consent",
            scope="cloud-private-home",
            delivery_confirmed=False,
        )
    )
    assert health_cli._has_scoped_consent(load_config(home), "cloud-private-home")

    def provider_b(path: Path) -> str | None:
        resolved = Path(path).resolve()
        return (
            "Dropbox"
            if resolved == home.resolve() or home.resolve() in resolved.parents
            else None
        )

    monkeypatch.setattr(health_cli, "_sync_storage_provider", provider_b)
    config = load_config(home)
    assert not health_cli._has_scoped_consent(config, "cloud-private-home")
    with pytest.raises(ValueError, match="cloud-private-home"):
        health_cli._require_private_home_storage_consent(config)


@pytest.mark.skipif(not TEMPLATE.exists(), reason="workbook asset is still being generated")
def test_force_init_cannot_project_existing_health_rows_to_cloud_without_consent(
    tmp_path: Path,
) -> None:
    home = tmp_path / "private-health-home"
    local_workbook = tmp_path / "local.xlsx"
    base = ("--home", str(home))
    run_cli(
        *base,
        "init",
        "--workbook",
        str(local_workbook),
        "--timezone",
        "UTC",
    )
    run_cli(
        *base,
        "record",
        "--kind",
        "measurement",
        "--json",
        json.dumps(
            {
                "date": "2026-01-02",
                "metric": "synthetic",
                "value": 1,
                "unit": "test",
            }
        ),
    )
    cloud_workbook = (
        tmp_path
        / "Library"
        / "CloudStorage"
        / "Box-Box"
        / "private-health.xlsx"
    )
    cloud_workbook.parent.mkdir(parents=True)

    _, initialized = run_cli(
        *base,
        "init",
        "--force",
        "--workbook",
        str(cloud_workbook),
    )

    assert initialized["workbook_export"] == "blocked_by_consent"
    assert not cloud_workbook.exists()
    with HealthDatabase(home / "health.sqlite3", create=False) as database:
        assert len(database.list_records("measurement")) == 1
    state = json.loads((home / "state.json").read_text(encoding="utf-8"))
    assert state["workbook_export_pending"] is True
    assert state["workbook_export_pending_operation"] == "init:workbook"


@pytest.mark.parametrize(
    ("relative_path", "provider"),
    [
        ("Library/CloudStorage/Box-Box/folder/health.xlsx", "Box"),
        (
            "Library/CloudStorage/GoogleDrive-user@example.test/Shared drives/team/health.xlsx",
            "Google Drive",
        ),
    ],
)
def test_cloud_provider_recognizes_common_macos_cloudstorage_paths(
    tmp_path: Path, relative_path: str, provider: str
) -> None:
    assert health_cli._sync_storage_provider(tmp_path / relative_path) == provider


@pytest.mark.skipif(not TEMPLATE.exists(), reason="workbook asset is still being generated")
def test_sync_keeps_local_truth_but_not_cloud_view_without_cloud_consent(
    tmp_path: Path,
) -> None:
    home = tmp_path / "private-health-home"
    cloud_dir = (
        tmp_path
        / "Library"
        / "Mobile Documents"
        / "com~apple~CloudDocs"
    )
    cloud_dir.mkdir(parents=True)
    workbook_path = cloud_dir / "health.xlsx"
    base = ("--home", str(home))
    _, initialized = run_cli(
        *base,
        "init",
        "--workbook",
        str(workbook_path),
        "--timezone",
        "UTC",
    )
    assert initialized["workbook_export"] == "blocked_by_consent"
    assert not workbook_path.exists()

    _, synced = run_cli(
        *base,
        "sync",
        "--from-date",
        "2026-01-01",
        "--to-date",
        "2026-01-03",
        "--fixture-dir",
        str(FIXTURES),
    )

    assert synced["status"] == "success"
    assert synced["workbook_export"] == "blocked_by_consent"
    assert not workbook_path.exists()
    with HealthDatabase(home / "health.sqlite3", create=False) as database:
        assert database.list_records("daily")
        run = database.list_sync_runs(limit=1)[0]
    assert run["status"] == "success"
    state = json.loads((home / "state.json").read_text(encoding="utf-8"))
    assert state["last_sync_status"] == "success_export_blocked_by_consent"
    assert state["workbook_export_pending"] is True


@pytest.mark.skipif(not TEMPLATE.exists(), reason="workbook asset is still being generated")
def test_cli_doctor_detects_and_repairs_goal_projection_drift(tmp_path: Path) -> None:
    home = tmp_path / "private-health-home"
    base = ("--home", str(home))
    run_cli(
        *base,
        "init",
        "--workbook",
        str(tmp_path / "健康档案.xlsx"),
        "--timezone",
        "UTC",
        "--ghealth-command",
        "missing-optional-ghealth",
    )
    exact_goal = "用于目标投影修复的合成目标"
    run_cli(
        *base,
        "goal",
        "set",
        "--text",
        exact_goal,
        "--effective-date",
        "2026-01-02",
    )

    profile_path = home / "profile.json"
    profile = json.loads(profile_path.read_text(encoding="utf-8"))
    profile["goals"] = []
    profile_path.write_text(json.dumps(profile, ensure_ascii=False), encoding="utf-8")
    agents_path = home / "AGENTS.md"
    agents_path.write_text(
        agents_path.read_text(encoding="utf-8").replace(exact_goal, "错误投影"),
        encoding="utf-8",
    )

    _, before = run_cli(*base, "doctor", "--offline")
    assert before["status"] == "needs_attention"
    projection = next(
        item for item in before["checks"] if item["check"] == "goal projections"
    )
    assert projection["ok"] is False

    _, repaired = run_cli(*base, "goal", "repair")
    assert repaired["status"] == "goal_views_rebuilt"
    assert repaired["projection"]["ok"] is True
    assert exact_goal in agents_path.read_text(encoding="utf-8")

    _, after = run_cli(*base, "doctor", "--offline")
    assert after["status"] == "ok"


@pytest.mark.skipif(not TEMPLATE.exists(), reason="workbook asset is still being generated")
def test_partial_sync_marks_existing_daily_only_row_stale(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("OPEN_HEALTH_AGENT_TEST_MODE", "1")
    home = tmp_path / "private-health-home"
    workbook_path = tmp_path / "健康档案.xlsx"
    run_cli(
        "--home",
        str(home),
        "init",
        "--workbook",
        str(workbook_path),
        "--timezone",
        "UTC",
    )
    with HealthDatabase(home / "health.sqlite3") as database:
        database.upsert(
            "daily",
            {
                "record_id": "2026-01-02",
                "date": "2026-01-02",
                "steps": 1000,
                "sources": "prior.steps.source",
                "data_until": "2026-01-02T20:00:00+00:00",
                "quality": "automatic ghealth import; blanks mean unavailable",
                "batch_id": "prior",
                "imported_at": "2026-01-02T20:05:00+00:00",
            },
        )

    class FailedStepsAdapter:
        def fetch(self, _start: str, _end: str, _batch_id: str) -> dict:
            return {
                "daily": [],
                "measurements": [],
                "workouts": [],
                "data_until": None,
                "errors": ["steps: synthetic failure"],
                "failed_query_keys": ["steps"],
            }

    monkeypatch.setattr(
        health_cli,
        "GHealthAdapter",
        lambda *_args, **_kwargs: FailedStepsAdapter(),
    )
    result = health_cli.command_sync(
        Namespace(
            home=str(home),
            from_date="2026-01-02",
            to_date="2026-01-02",
            fixture_dir=str(tmp_path),
            quiet=False,
        )
    )

    assert result["status"] == "failed"
    with HealthDatabase(home / "health.sqlite3") as database:
        daily = database.get("daily", "2026-01-02")
    assert daily is not None
    assert daily["steps"] == 1000
    assert daily["sources"] == "prior.steps.source"
    assert "partial; failed daily queries: steps" in daily["quality"]
    assert "stale fields retained from prior sync: steps" in daily["quality"]


@pytest.mark.skipif(not TEMPLATE.exists(), reason="workbook asset is still being generated")
def test_sync_summary_records_workbook_export_failure(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("OPEN_HEALTH_AGENT_TEST_MODE", "1")
    home = tmp_path / "private-health-home"
    workbook_path = tmp_path / "健康档案.xlsx"
    run_cli(
        "--home",
        str(home),
        "init",
        "--workbook",
        str(workbook_path),
        "--timezone",
        "UTC",
    )
    private_marker = "private-health-detail-must-not-leak"

    def fail_export(*_args, **_kwargs):
        raise RuntimeError(private_marker)

    monkeypatch.setattr(health_cli, "export_workbook", fail_export)
    with pytest.raises(RuntimeError, match=private_marker):
        health_cli.command_sync(
            Namespace(
                home=str(home),
                from_date="2026-01-01",
                to_date="2026-01-03",
                fixture_dir=str(FIXTURES),
                quiet=False,
            )
        )

    summaries = list((home / "raw").glob("*.summary.json"))
    assert len(summaries) == 1
    summary = json.loads(summaries[0].read_text(encoding="utf-8"))
    assert summary["status"] == "failed_export"
    assert private_marker not in json.dumps(summary, ensure_ascii=False)
    with HealthDatabase(home / "health.sqlite3") as database:
        sync_run = database.list_sync_runs(1)[0]
    assert sync_run["status"] == "failed_export"
    assert json.loads((home / "state.json").read_text(encoding="utf-8"))[
        "last_sync_status"
    ] == "failed_export"


def test_cli_runtime_error_redacts_health_text_and_private_paths(
    monkeypatch: pytest.MonkeyPatch,
    capsys: pytest.CaptureFixture[str],
) -> None:
    private_health_text = "我的血压是非常私密的 188/120"
    private_database = "/" + "Users/" + "private-person/.open-health-agent/health.sqlite3"

    def fail_with_private_exception(_arguments) -> dict:
        raise RuntimeError(f"record={private_health_text}; database={private_database}")

    monkeypatch.setattr(health_cli, "command_context", fail_with_private_exception)
    monkeypatch.setattr(sys, "argv", ["health-agent", "context"])

    assert health_cli.main() == 1
    payload = json.loads(capsys.readouterr().err)
    serialized = json.dumps(payload, ensure_ascii=False)
    assert payload == {
        "status": "error",
        "code": "operation_failed",
        "type": "RuntimeError",
        "error": "The operation could not complete; run doctor and retry.",
    }
    assert private_health_text not in serialized
    assert private_database not in serialized


def test_cli_parser_error_does_not_echo_unrecognized_private_arguments() -> None:
    private_health_text = "我刚刚吃了私密测试餐"
    private_database = "/" + "Users/" + "private-person/.open-health-agent/health.sqlite3"
    completed = subprocess.run(
        [
            sys.executable,
            str(CLI),
            "context",
            "--unexpected-private-value",
            private_health_text,
            private_database,
        ],
        cwd=ROOT,
        check=False,
        capture_output=True,
        text=True,
    )

    assert completed.returncode == 1
    payload = json.loads(completed.stderr)
    assert payload["code"] == "invalid_command"
    assert payload["type"] == "CLIUsageError"
    assert private_health_text not in completed.stderr
    assert private_database not in completed.stderr


def test_fixture_sync_is_unavailable_without_the_repository_test_guard(
    tmp_path: Path,
) -> None:
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    environment = os.environ.copy()
    environment.pop("OPEN_HEALTH_AGENT_TEST_MODE", None)

    completed = subprocess.run(
        [
            sys.executable,
            str(CLI),
            "--home",
            str(home),
            "sync",
            "--from-date",
            "2026-01-01",
            "--to-date",
            "2026-01-03",
            "--fixture-dir",
            str(FIXTURES),
        ],
        cwd=ROOT,
        check=False,
        capture_output=True,
        text=True,
        env=environment,
    )

    assert completed.returncode == 1
    assert json.loads(completed.stderr)["code"] == "fixture_sync_disabled"
    with HealthDatabase(home / "health.sqlite3", create=False) as database:
        assert database.list_sync_runs() == []


def test_onboarding_status_projects_only_consent_metadata(tmp_path: Path) -> None:
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    private_marker = "synthetic-private-operational-marker"
    atomic_write_json(
        home / "state.json",
        {
            "onboarding_explanation_delivery_confirmed_at": now,
            "onboarding_consents": {
                "google-health": {"granted_at": now, "policy_version": 1}
            },
            "last_sync_batch_id": private_marker,
            "last_manual_ghealth_sync": {
                "version": 2,
                "batch_id": private_marker,
                "status": "success",
                "finished_at": now,
                "timezone": "UTC",
                "static_runtime_fingerprint": "b" * 64,
                "runtime_fingerprint": "a" * 64,
            },
            "workbook_export_pending_since": now,
            "workbook_export_pending_operation": private_marker,
        },
    )

    result = health_cli.command_onboarding(
        Namespace(home=str(home), action="status", scope=None, delivery_confirmed=False)
    )
    serialized = json.dumps(result)

    assert result["explanation_delivered"] is True
    assert result["consents"]["google-health"]["active"] is True
    assert private_marker not in serialized
    for operational_key in (
        "last_sync_batch_id",
        "last_manual_ghealth_sync",
        "runtime_fingerprint",
        "workbook_export_pending_since",
    ):
        assert operational_key not in serialized


def test_invalid_state_shape_fails_closed_and_doctor_reports_it(
    tmp_path: Path,
) -> None:
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    (home / "state.json").write_text("[]\n", encoding="utf-8")

    _, status_error = run_cli(
        "--home", str(home), "onboarding", "status", expected_code=1
    )
    _, doctor = run_cli("--home", str(home), "doctor", "--offline")

    assert status_error["code"] == "state_invalid"
    state_check = next(
        item for item in doctor["checks"] if item["check"] == "state JSON schema"
    )
    assert doctor["status"] == "needs_attention"
    assert state_check["ok"] is False
    assert not health_cli._has_scoped_consent(load_config(home), "google-health")


@pytest.mark.parametrize(
    "invalid_state",
    [
        {"workbook_export_pending": "yes"},
        {"profile_projection_pending": 1},
    ],
)
def test_invalid_known_state_field_types_make_context_freshness_unknown(
    tmp_path: Path, invalid_state: dict[str, object]
) -> None:
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    atomic_write_json(home / "state.json", invalid_state)

    _, context = run_cli("--home", str(home), "context")
    _, doctor = run_cli("--home", str(home), "doctor", "--offline")

    assert context["freshness"]["workbook_projection_status"] == "unknown"
    state_check = next(
        item for item in doctor["checks"] if item["check"] == "state JSON schema"
    )
    assert state_check["ok"] is False


def test_missing_state_file_fails_closed_and_doctor_reports_it(tmp_path: Path) -> None:
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    (home / "state.json").unlink()

    _, status_error = run_cli(
        "--home", str(home), "onboarding", "status", expected_code=1
    )
    _, doctor = run_cli("--home", str(home), "doctor", "--offline")

    assert status_error["code"] == "state_invalid"
    state_check = next(
        item for item in doctor["checks"] if item["check"] == "state JSON schema"
    )
    assert state_check["ok"] is False


@pytest.mark.parametrize(
    "timestamp",
    [
        "2026-01-01T00:00:00",
        (datetime.now(timezone.utc) + timedelta(days=1)).isoformat(
            timespec="seconds"
        ),
    ],
)
def test_consent_public_view_and_enforcement_reject_invalid_timezones_or_future_grants(
    tmp_path: Path, timestamp: str
) -> None:
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    atomic_write_json(
        home / "state.json",
        {
            "onboarding_explanation_delivery_confirmed_at": timestamp,
            "onboarding_consents": {
                "google-health": {"granted_at": timestamp, "policy_version": 1}
            },
        },
    )

    with pytest.raises(health_cli.StateInvalidError):
        health_cli.command_onboarding(
            Namespace(
                home=str(home),
                action="status",
                scope=None,
                delivery_confirmed=False,
            )
        )
    assert not health_cli._has_scoped_consent(load_config(home), "google-health")


def test_keep_awake_rejects_an_unused_but_stale_one_shot_grant(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    stale = (datetime.now(timezone.utc) - timedelta(hours=1)).isoformat(
        timespec="seconds"
    )
    atomic_write_json(
        home / "state.json",
        {
            "onboarding_consents": {
                "keep-awake": {"granted_at": stale, "policy_version": 1}
            }
        },
    )
    installs: list[str] = []
    monkeypatch.setattr(
        health_cli,
        "install_macos_ac_keepawake",
        lambda _config: installs.append("installed") or {},
    )

    with pytest.raises(ValueError, match="fresh keep-awake consent"):
        health_cli.command_power(
            Namespace(
                home=str(home),
                action="install",
                verbose_paths=False,
            )
        )

    assert installs == []


def test_keep_awake_removal_failure_is_durable_and_visible_in_status(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    atomic_write_json(
        home / "state.json",
        {
            "onboarding_consents": {
                "keep-awake": {"granted_at": now, "policy_version": 1}
            }
        },
    )
    monkeypatch.setattr(
        health_cli,
        "uninstall_macos_ac_keepawake",
        lambda: (_ for _ in ()).throw(RuntimeError("keep-awake synthetic failure")),
    )

    result = health_cli.command_power(
        Namespace(home=str(home), action="uninstall", verbose_paths=False)
    )

    assert result["status"] == "consent_revoked_job_removal_pending"
    state = json.loads((home / "state.json").read_text(encoding="utf-8"))
    assert state["keep_awake_job_removal_pending"] is True
    assert state["keep_awake_job_removal_error_code"] == "scheduler_operation_failed"
    monkeypatch.setattr(
        health_cli,
        "keepawake_status",
        lambda: {"backend": "launchd", "installed": True},
    )
    status = health_cli.command_power(
        Namespace(home=str(home), action="status", verbose_paths=False)
    )
    assert status["keep_awake_consent_active"] is False
    assert status["job_removal_pending"] is True
    assert status["job_removal_error_code"] == "scheduler_operation_failed"


def test_scheduler_status_does_not_hold_writer_lock_during_external_probe(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    status_started = threading.Event()
    release_status = threading.Event()

    def slow_status(_config):
        status_started.set()
        if not release_status.wait(10):
            raise AssertionError("test did not release scheduler status")
        return {"backend": "synthetic", "installed": False}

    monkeypatch.setattr(health_cli, "scheduler_status", slow_status)
    status_errors: list[BaseException] = []
    record_errors: list[BaseException] = []
    record_finished = threading.Event()

    def run_status() -> None:
        try:
            health_cli.command_scheduler(
                Namespace(
                    home=str(home),
                    action="status",
                    interval_seconds=3600,
                    verbose_paths=False,
                    proxy_env_file=None,
                    inherit_proxy_env=False,
                )
            )
        except BaseException as exc:  # pragma: no cover - asserted below
            status_errors.append(exc)

    def run_record() -> None:
        try:
            health_cli.command_record(
                Namespace(
                    home=str(home),
                    kind="measurement",
                    record_id=None,
                    json_payload=json.dumps(
                        {
                            "date": "2026-01-02",
                            "metric": "synthetic scheduler-status metric",
                            "value": 3,
                            "unit": "test",
                        }
                    ),
                    file=None,
                )
            )
        except BaseException as exc:  # pragma: no cover - asserted below
            record_errors.append(exc)
        finally:
            record_finished.set()

    status_thread = threading.Thread(target=run_status)
    status_thread.start()
    assert status_started.wait(5)
    record_thread = threading.Thread(target=run_record)
    record_thread.start()
    record_completed_while_status_waited = record_finished.wait(5)
    release_status.set()
    record_thread.join(10)
    status_thread.join(10)

    assert record_completed_while_status_waited is True
    assert not status_thread.is_alive()
    assert not record_thread.is_alive()
    assert status_errors == []
    assert record_errors == []


def test_keep_awake_status_does_not_hold_writer_lock_during_external_probe(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    status_started = threading.Event()
    release_status = threading.Event()

    def slow_status():
        status_started.set()
        if not release_status.wait(10):
            raise AssertionError("test did not release keep-awake status")
        return {"backend": "synthetic", "installed": False}

    monkeypatch.setattr(health_cli, "keepawake_status", slow_status)
    status_errors: list[BaseException] = []
    record_errors: list[BaseException] = []
    record_finished = threading.Event()

    def run_status() -> None:
        try:
            health_cli.command_power(
                Namespace(home=str(home), action="status", verbose_paths=False)
            )
        except BaseException as exc:  # pragma: no cover - asserted below
            status_errors.append(exc)

    def run_record() -> None:
        try:
            health_cli.command_record(
                Namespace(
                    home=str(home),
                    kind="measurement",
                    record_id=None,
                    json_payload=json.dumps(
                        {
                            "date": "2026-01-02",
                            "metric": "synthetic keep-awake-status metric",
                            "value": 4,
                            "unit": "test",
                        }
                    ),
                    file=None,
                )
            )
        except BaseException as exc:  # pragma: no cover - asserted below
            record_errors.append(exc)
        finally:
            record_finished.set()

    status_thread = threading.Thread(target=run_status)
    status_thread.start()
    assert status_started.wait(5)
    record_thread = threading.Thread(target=run_record)
    record_thread.start()
    record_completed_while_status_waited = record_finished.wait(5)
    release_status.set()
    record_thread.join(10)
    status_thread.join(10)

    assert record_completed_while_status_waited is True
    assert not status_thread.is_alive()
    assert not record_thread.is_alive()
    assert status_errors == []
    assert record_errors == []


@pytest.mark.parametrize("install_command", ["scheduler", "keep-awake"])
def test_cloud_private_home_install_gate_runs_before_writer_lock(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, install_command: str
) -> None:
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")

    def fake_sync_provider(path: Path) -> str | None:
        resolved = Path(path).resolve()
        return "iCloud Drive" if resolved == home.resolve() else None

    def forbidden_lock(_config):
        raise AssertionError("writer lock must not be created before cloud-home consent")

    monkeypatch.setattr(health_cli, "_sync_storage_provider", fake_sync_provider)
    monkeypatch.setattr(health_cli, "lock_for", forbidden_lock)

    if install_command == "scheduler":
        with pytest.raises(ValueError, match="cloud-private-home"):
            health_cli.command_scheduler(
                Namespace(
                    home=str(home),
                    action="install",
                    interval_seconds=3600,
                    verbose_paths=False,
                    proxy_env_file=None,
                    inherit_proxy_env=False,
                )
            )
    else:
        with pytest.raises(ValueError, match="cloud-private-home"):
            health_cli.command_power(
                Namespace(home=str(home), action="install", verbose_paths=False)
            )


def test_external_scope_revocation_warns_that_host_action_is_still_required(
    tmp_path: Path,
) -> None:
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    atomic_write_json(
        home / "state.json",
        {
            "onboarding_consents": {
                "vision": {"granted_at": now, "policy_version": 1}
            }
        },
    )

    result = health_cli.command_onboarding(
        Namespace(
            home=str(home),
            action="revoke-consent",
            scope="vision",
            delivery_confirmed=False,
        )
    )

    assert result["status"] == "local_consent_revoked_external_disable_required"
    assert result["external_action_required"] is True
    assert "Hermes" in result["action"]
    assert not health_cli._has_scoped_consent(load_config(home), "vision")


def test_unwritable_raw_summary_does_not_undo_a_durable_sync(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("OPEN_HEALTH_AGENT_TEST_MODE", "1")
    home = tmp_path / "private-health-home"
    workbook_path = tmp_path / "synthetic-ledger.xlsx"
    run_cli(
        "--home",
        str(home),
        "init",
        "--workbook",
        str(workbook_path),
        "--timezone",
        "UTC",
    )
    original_atomic_write = health_cli.atomic_write_json

    def fail_only_raw_summary(path, value, *args, **kwargs):
        if Path(path).parent == home / "raw":
            raise OSError("synthetic raw summary failure")
        return original_atomic_write(path, value, *args, **kwargs)

    monkeypatch.setattr(health_cli, "atomic_write_json", fail_only_raw_summary)

    result = health_cli.command_sync(
        Namespace(
            home=str(home),
            from_date="2026-01-01",
            to_date="2026-01-03",
            lookback_days=None,
            fixture_dir=str(FIXTURES),
            scheduled=False,
            runtime_fingerprint=None,
            quiet=False,
            verbose_path=False,
        )
    )

    assert result["status"] == "success"
    assert workbook_path.is_file()
    assert list((home / "raw").glob("*.summary.json")) == []
    with HealthDatabase(home / "health.sqlite3", create=False) as database:
        assert database.list_sync_runs(1)[0]["status"] == "success"
    state = json.loads((home / "state.json").read_text(encoding="utf-8"))
    assert state["last_sync_status"] == "success"
    assert state["workbook_export_pending"] is False


def test_failed_sync_finishes_the_exact_same_second_batch(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("OPEN_HEALTH_AGENT_TEST_MODE", "1")
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    current_batches: list[str] = []

    def tied_begin_sync(
        self,
        batch_id: str,
        from_date: str,
        to_date: str,
        **_kwargs,
    ) -> None:
        current_batches.append(batch_id)
        timestamp = "2026-01-01T00:00:00+00:00"
        self.connection.execute(
            "INSERT INTO sync_runs(batch_id, started_at, from_date, to_date, status) VALUES (?, ?, ?, ?, 'running')",
            ("sync_fixture_manual_prior", timestamp, from_date, to_date),
        )
        self.connection.execute(
            "INSERT INTO sync_runs(batch_id, started_at, from_date, to_date, status) VALUES (?, ?, ?, ?, 'running')",
            (batch_id, timestamp, from_date, to_date),
        )
        self.connection.commit()

    class FailingAdapter:
        def fetch(self, _start: str, _end: str, _batch_id: str) -> dict:
            raise RuntimeError("synthetic fetch failure")

    monkeypatch.setattr(HealthDatabase, "begin_sync", tied_begin_sync)
    monkeypatch.setattr(
        health_cli, "GHealthAdapter", lambda *_args, **_kwargs: FailingAdapter()
    )

    with pytest.raises(RuntimeError, match="synthetic fetch failure"):
        health_cli.command_sync(
            Namespace(
                home=str(home),
                from_date="2026-01-01",
                to_date="2026-01-02",
                lookback_days=None,
                fixture_dir=str(FIXTURES),
                scheduled=False,
                runtime_fingerprint=None,
                quiet=False,
                verbose_path=False,
            )
        )

    assert len(current_batches) == 1
    with HealthDatabase(home / "health.sqlite3", create=False) as database:
        assert database.get_sync_run(current_batches[0])["status"] == "failed"
        assert database.get_sync_run("sync_fixture_manual_prior")["status"] == "running"


def test_scheduled_sync_rejects_runtime_drift_before_fetch(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("GHEALTH_PROFILE", "default")
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    atomic_write_json(
        home / "state.json",
        {
            "onboarding_consents": {
                "google-health": {"granted_at": now, "policy_version": 1},
                "scheduler": {
                    "granted_at": now,
                    "consumed_at": now,
                    "policy_version": 1,
                },
            }
        },
    )
    executable = tmp_path / "ghealth"
    executable.write_text("synthetic", encoding="utf-8")
    executable.chmod(0o700)
    fetched: list[bool] = []

    class MatchingAdapter:
        def require_matching_timezone(self) -> None:
            return None

        def fetch(self, *_args) -> dict:
            fetched.append(True)
            raise AssertionError("fetch must not run after runtime drift")

    monkeypatch.setattr(
        health_cli, "resolve_ghealth_executable", lambda _command: executable
    )
    monkeypatch.setattr(
        health_cli, "GHealthAdapter", lambda *_args, **_kwargs: MatchingAdapter()
    )
    runtime_identity_checks: list[bool] = []

    def account_bound_runtime(_config, _profile=None):
        runtime_identity_checks.append(True)
        return "a" * 64

    monkeypatch.setattr(
        health_cli,
        "ghealth_runtime_fingerprint",
        account_bound_runtime,
    )
    monkeypatch.setattr(
        health_cli, "static_runtime_fingerprint", lambda _config: "c" * 64
    )
    monkeypatch.setattr(
        health_cli,
        "_has_scheduled_runtime_authorization",
        lambda *_args, **_kwargs: True,
    )

    arguments = Namespace(
        home=str(home),
        from_date="2026-01-01",
        to_date="2026-01-02",
        lookback_days=None,
        fixture_dir=None,
        scheduled=True,
        static_runtime_fingerprint="d" * 64,
        runtime_fingerprint="b" * 64,
        quiet=True,
        verbose_path=False,
    )
    with pytest.raises(ValueError, match="static runtime fingerprint"):
        health_cli.command_sync(arguments)
    assert runtime_identity_checks == []

    arguments.static_runtime_fingerprint = "c" * 64
    with pytest.raises(ValueError, match="account-bound runtime fingerprint"):
        health_cli.command_sync(
            arguments
        )

    assert runtime_identity_checks == [True]
    assert fetched == []
    with HealthDatabase(home / "health.sqlite3", create=False) as database:
        assert database.list_sync_runs() == []


def test_manual_sync_pins_profile_for_all_queries_and_rechecks_identity(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    atomic_write_json(
        home / "state.json",
        {
            "onboarding_consents": {
                "google-health": {"granted_at": now, "policy_version": 1}
            }
        },
    )
    executable = tmp_path / "ghealth"
    executable.write_text("synthetic", encoding="utf-8")
    executable.chmod(0o700)
    observed_profiles: list[str | None] = []

    class PinnedRunner:
        def __init__(self, _command: str, *, profile: str | None = None):
            observed_profiles.append(profile)

    class SuccessfulAdapter:
        def require_matching_timezone(self) -> None:
            return None

        def fetch(self, _start: str, _end: str, batch_id: str) -> dict:
            return {
                "daily": [
                    {
                        "record_id": "2026-01-02",
                        "date": "2026-01-02",
                        "steps": 1234,
                        "batch_id": batch_id,
                        "data_until": "2026-01-02T09:00:00+00:00",
                    }
                ],
                "measurements": [],
                "workouts": [],
                "errors": [],
                "failed_query_keys": [],
                "data_until": "2026-01-02T09:00:00+00:00",
            }

    monkeypatch.setattr(
        health_cli, "resolve_ghealth_executable", lambda _command: executable
    )
    monkeypatch.setattr(
        health_cli, "resolve_active_ghealth_profile", lambda _config: "profile-a"
    )
    monkeypatch.setattr(health_cli, "CommandRunner", PinnedRunner)
    monkeypatch.setattr(
        health_cli,
        "GHealthAdapter",
        lambda *_args, **_kwargs: SuccessfulAdapter(),
    )
    monkeypatch.setattr(
        health_cli, "static_runtime_fingerprint", lambda _config: "c" * 64
    )
    identity_profiles: list[str | None] = []

    def stable_identity(_config, profile=None):
        identity_profiles.append(profile)
        return "a" * 64

    monkeypatch.setattr(health_cli, "ghealth_runtime_fingerprint", stable_identity)

    result = health_cli.command_sync(
        Namespace(
            home=str(home),
            from_date="2026-01-02",
            to_date="2026-01-02",
            lookback_days=None,
            fixture_dir=None,
            scheduled=False,
            quiet=False,
            static_runtime_fingerprint=None,
            runtime_fingerprint=None,
            verbose_path=False,
        )
    )

    assert result["status"] == "success"
    assert observed_profiles == ["profile-a"]
    assert identity_profiles == ["profile-a", "profile-a"]


def test_manual_sync_discards_fetched_rows_if_profile_identity_drifts(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    atomic_write_json(
        home / "state.json",
        {
            "onboarding_consents": {
                "google-health": {"granted_at": now, "policy_version": 1}
            }
        },
    )
    executable = tmp_path / "ghealth"
    executable.write_text("synthetic", encoding="utf-8")
    executable.chmod(0o700)

    class DriftAdapter:
        def require_matching_timezone(self) -> None:
            return None

        def fetch(self, _start: str, _end: str, batch_id: str) -> dict:
            return {
                "daily": [
                    {
                        "record_id": "2026-01-02",
                        "date": "2026-01-02",
                        "steps": 1234,
                        "batch_id": batch_id,
                    }
                ],
                "measurements": [],
                "workouts": [],
                "errors": [],
                "failed_query_keys": [],
                "data_until": None,
            }

    monkeypatch.setattr(
        health_cli, "resolve_ghealth_executable", lambda _command: executable
    )
    monkeypatch.setattr(
        health_cli, "resolve_active_ghealth_profile", lambda _config: "profile-a"
    )
    monkeypatch.setattr(health_cli, "CommandRunner", lambda *_args, **_kwargs: object())
    monkeypatch.setattr(
        health_cli, "GHealthAdapter", lambda *_args, **_kwargs: DriftAdapter()
    )
    monkeypatch.setattr(
        health_cli, "static_runtime_fingerprint", lambda _config: "c" * 64
    )
    identities = iter(("a" * 64, "b" * 64))
    monkeypatch.setattr(
        health_cli,
        "ghealth_runtime_fingerprint",
        lambda _config, _profile=None: next(identities),
    )

    with pytest.raises(RuntimeError, match="profile identity changed"):
        health_cli.command_sync(
            Namespace(
                home=str(home),
                from_date="2026-01-02",
                to_date="2026-01-02",
                lookback_days=None,
                fixture_dir=None,
                scheduled=False,
                quiet=False,
                static_runtime_fingerprint=None,
                runtime_fingerprint=None,
                verbose_path=False,
            )
        )

    with HealthDatabase(home / "health.sqlite3", create=False) as database:
        assert database.list_records("daily") == []
        assert database.list_sync_runs(limit=1) == []
    state = json.loads((home / "state.json").read_text(encoding="utf-8"))
    assert state["last_sync_status"] == "failed_postfetch_validation"
    assert state["last_manual_ghealth_sync"]["status"] == (
        "failed_postfetch_validation"
    )


def test_real_sync_fetch_does_not_hold_the_ledger_writer_lock(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    atomic_write_json(
        home / "state.json",
        {
            "onboarding_consents": {
                "google-health": {"granted_at": now, "policy_version": 1}
            }
        },
    )
    executable = tmp_path / "ghealth"
    executable.write_text("synthetic", encoding="utf-8")
    executable.chmod(0o700)
    fetch_started = threading.Event()
    release_fetch = threading.Event()

    class SlowAdapter:
        def require_matching_timezone(self) -> None:
            return None

        def fetch(self, _start: str, _end: str, batch_id: str) -> dict:
            fetch_started.set()
            if not release_fetch.wait(10):
                raise AssertionError("test did not release the synthetic fetch")
            return {
                "daily": [
                    {
                        "record_id": "2026-01-02",
                        "date": "2026-01-02",
                        "steps": 1234,
                        "batch_id": batch_id,
                    }
                ],
                "measurements": [],
                "workouts": [],
                "errors": [],
                "failed_query_keys": [],
                "data_until": None,
            }

    monkeypatch.setattr(
        health_cli, "resolve_ghealth_executable", lambda _command: executable
    )
    monkeypatch.setattr(
        health_cli, "resolve_active_ghealth_profile", lambda _config: "profile-a"
    )
    monkeypatch.setattr(health_cli, "CommandRunner", lambda *_args, **_kwargs: object())
    monkeypatch.setattr(
        health_cli, "GHealthAdapter", lambda *_args, **_kwargs: SlowAdapter()
    )
    monkeypatch.setattr(
        health_cli, "static_runtime_fingerprint", lambda _config: "c" * 64
    )
    monkeypatch.setattr(
        health_cli,
        "ghealth_runtime_fingerprint",
        lambda _config, _profile=None: "a" * 64,
    )
    sync_errors: list[BaseException] = []
    record_errors: list[BaseException] = []
    record_finished = threading.Event()

    def run_sync() -> None:
        try:
            health_cli.command_sync(
                Namespace(
                    home=str(home),
                    from_date="2026-01-02",
                    to_date="2026-01-02",
                    lookback_days=None,
                    fixture_dir=None,
                    scheduled=False,
                    quiet=False,
                    static_runtime_fingerprint=None,
                    runtime_fingerprint=None,
                    verbose_path=False,
                )
            )
        except BaseException as exc:  # pragma: no cover - asserted below
            sync_errors.append(exc)

    def run_record() -> None:
        try:
            health_cli.command_record(
                Namespace(
                    home=str(home),
                    kind="measurement",
                    record_id=None,
                    json_payload=json.dumps(
                        {
                            "date": "2026-01-02",
                            "metric": "synthetic concurrent metric",
                            "value": 1,
                            "unit": "test",
                        }
                    ),
                    file=None,
                )
            )
        except BaseException as exc:  # pragma: no cover - asserted below
            record_errors.append(exc)
        finally:
            record_finished.set()

    sync_thread = threading.Thread(target=run_sync)
    sync_thread.start()
    assert fetch_started.wait(5)
    record_thread = threading.Thread(target=run_record)
    record_thread.start()
    record_completed_while_fetching = record_finished.wait(5)
    release_fetch.set()
    record_thread.join(10)
    sync_thread.join(10)

    assert record_completed_while_fetching is True
    assert not sync_thread.is_alive()
    assert not record_thread.is_alive()
    assert sync_errors == []
    assert record_errors == []
    with HealthDatabase(home / "health.sqlite3", create=False) as database:
        assert len(database.list_records("measurement")) == 1
        assert len(database.list_records("daily")) == 1


def test_real_sync_postfetch_ghealth_validation_does_not_hold_writer_lock(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    atomic_write_json(
        home / "state.json",
        {
            "onboarding_consents": {
                "google-health": {"granted_at": now, "policy_version": 1}
            }
        },
    )
    executable = tmp_path / "ghealth"
    executable.write_text("synthetic", encoding="utf-8")
    executable.chmod(0o700)
    postfetch_validation_started = threading.Event()
    release_postfetch_validation = threading.Event()

    class FastFetchAdapter:
        def require_matching_timezone(self) -> None:
            return None

        def fetch(self, _start: str, _end: str, batch_id: str) -> dict:
            return {
                "daily": [
                    {
                        "record_id": "2026-01-02",
                        "date": "2026-01-02",
                        "steps": 1234,
                        "batch_id": batch_id,
                    }
                ],
                "measurements": [],
                "workouts": [],
                "errors": [],
                "failed_query_keys": [],
                "data_until": None,
            }

    monkeypatch.setattr(
        health_cli, "resolve_ghealth_executable", lambda _command: executable
    )
    monkeypatch.setattr(
        health_cli, "resolve_active_ghealth_profile", lambda _config: "profile-a"
    )
    monkeypatch.setattr(health_cli, "CommandRunner", lambda *_args, **_kwargs: object())
    monkeypatch.setattr(
        health_cli, "GHealthAdapter", lambda *_args, **_kwargs: FastFetchAdapter()
    )
    monkeypatch.setattr(
        health_cli, "static_runtime_fingerprint", lambda _config: "c" * 64
    )
    runtime_checks = 0

    def slow_second_runtime_check(_config, _profile=None):
        nonlocal runtime_checks
        runtime_checks += 1
        if runtime_checks == 2:
            postfetch_validation_started.set()
            if not release_postfetch_validation.wait(10):
                raise AssertionError("test did not release post-fetch validation")
        return "a" * 64

    monkeypatch.setattr(
        health_cli, "ghealth_runtime_fingerprint", slow_second_runtime_check
    )
    sync_errors: list[BaseException] = []
    record_errors: list[BaseException] = []
    record_finished = threading.Event()

    def run_sync() -> None:
        try:
            health_cli.command_sync(
                Namespace(
                    home=str(home),
                    from_date="2026-01-02",
                    to_date="2026-01-02",
                    lookback_days=None,
                    fixture_dir=None,
                    scheduled=False,
                    quiet=False,
                    static_runtime_fingerprint=None,
                    runtime_fingerprint=None,
                    verbose_path=False,
                )
            )
        except BaseException as exc:  # pragma: no cover - asserted below
            sync_errors.append(exc)

    def run_record() -> None:
        try:
            health_cli.command_record(
                Namespace(
                    home=str(home),
                    kind="measurement",
                    record_id=None,
                    json_payload=json.dumps(
                        {
                            "date": "2026-01-02",
                            "metric": "synthetic postfetch metric",
                            "value": 2,
                            "unit": "test",
                        }
                    ),
                    file=None,
                )
            )
        except BaseException as exc:  # pragma: no cover - asserted below
            record_errors.append(exc)
        finally:
            record_finished.set()

    sync_thread = threading.Thread(target=run_sync)
    sync_thread.start()
    assert postfetch_validation_started.wait(5)
    record_thread = threading.Thread(target=run_record)
    record_thread.start()
    record_completed_while_validating = record_finished.wait(5)
    release_postfetch_validation.set()
    record_thread.join(10)
    sync_thread.join(10)

    assert record_completed_while_validating is True
    assert not sync_thread.is_alive()
    assert not record_thread.is_alive()
    assert sync_errors == []
    assert record_errors == []
    with HealthDatabase(home / "health.sqlite3", create=False) as database:
        assert len(database.list_records("measurement")) == 1
        assert len(database.list_records("daily")) == 1


@pytest.mark.parametrize("drift_kind", ["consent", "config"])
def test_real_sync_discards_fetched_rows_after_locked_authority_or_config_drift(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, drift_kind: str
) -> None:
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    atomic_write_json(
        home / "state.json",
        {
            "onboarding_consents": {
                "google-health": {"granted_at": now, "policy_version": 1}
            }
        },
    )
    executable = tmp_path / "ghealth"
    executable.write_text("synthetic", encoding="utf-8")
    executable.chmod(0o700)
    fetch_started = threading.Event()
    release_fetch = threading.Event()

    class SlowAdapter:
        def require_matching_timezone(self) -> None:
            return None

        def fetch(self, _start: str, _end: str, batch_id: str) -> dict:
            fetch_started.set()
            if not release_fetch.wait(10):
                raise AssertionError("test did not release the synthetic fetch")
            return {
                "daily": [
                    {
                        "record_id": "2026-01-02",
                        "date": "2026-01-02",
                        "steps": 1234,
                        "batch_id": batch_id,
                    }
                ],
                "measurements": [],
                "workouts": [],
                "errors": [],
                "failed_query_keys": [],
                "data_until": None,
            }

    monkeypatch.setattr(
        health_cli, "resolve_ghealth_executable", lambda _command: executable
    )
    monkeypatch.setattr(
        health_cli, "resolve_active_ghealth_profile", lambda _config: "profile-a"
    )
    monkeypatch.setattr(health_cli, "CommandRunner", lambda *_args, **_kwargs: object())
    monkeypatch.setattr(
        health_cli, "GHealthAdapter", lambda *_args, **_kwargs: SlowAdapter()
    )
    monkeypatch.setattr(
        health_cli, "static_runtime_fingerprint", lambda _config: "c" * 64
    )
    monkeypatch.setattr(
        health_cli,
        "ghealth_runtime_fingerprint",
        lambda _config, _profile=None: "a" * 64,
    )
    sync_errors: list[BaseException] = []

    def run_sync() -> None:
        try:
            health_cli.command_sync(
                Namespace(
                    home=str(home),
                    from_date="2026-01-02",
                    to_date="2026-01-02",
                    lookback_days=None,
                    fixture_dir=None,
                    scheduled=False,
                    quiet=False,
                    static_runtime_fingerprint=None,
                    runtime_fingerprint=None,
                    verbose_path=False,
                )
            )
        except BaseException as exc:  # pragma: no cover - asserted below
            sync_errors.append(exc)

    sync_thread = threading.Thread(target=run_sync)
    sync_thread.start()
    assert fetch_started.wait(5)
    if drift_kind == "consent":
        health_cli.command_onboarding(
            Namespace(
                home=str(home),
                action="revoke-consent",
                scope="google-health",
                delivery_confirmed=False,
            )
        )
    else:
        config = load_config(home)
        with health_cli.lock_for(config):
            config = load_config(home)
            config.lookback_days += 1
            health_cli.save_config(config)
    release_fetch.set()
    sync_thread.join(10)

    assert not sync_thread.is_alive()
    assert len(sync_errors) == 1
    if drift_kind == "consent":
        assert "consent was withdrawn" in str(sync_errors[0])
    else:
        assert "configuration or date range changed" in str(sync_errors[0])
    with HealthDatabase(home / "health.sqlite3", create=False) as database:
        assert database.list_records("daily") == []
        assert database.list_sync_runs() == []
    state = json.loads((home / "state.json").read_text(encoding="utf-8"))
    assert state["last_sync_status"] == "failed_postfetch_validation"
    assert state["last_manual_ghealth_sync"]["status"] == (
        "failed_postfetch_validation"
    )


def test_real_sync_checks_consent_only_after_entering_the_shared_lock(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    inside_lock = False

    class ObservedLock:
        def __enter__(self):
            nonlocal inside_lock
            inside_lock = True
            return self

        def __exit__(self, *_args):
            nonlocal inside_lock
            inside_lock = False

    monkeypatch.setattr(health_cli, "lock_for", lambda _config: ObservedLock())

    def deny_consent(_config, _scope, **_kwargs):
        assert inside_lock is True
        return False

    monkeypatch.setattr(health_cli, "_has_scoped_consent", deny_consent)

    with pytest.raises(ValueError, match="google-health consent"):
        health_cli.command_sync(
            Namespace(
                home=str(home),
                from_date="2026-01-01",
                to_date="2026-01-02",
                lookback_days=None,
                fixture_dir=None,
                scheduled=False,
                runtime_fingerprint=None,
                quiet=False,
                verbose_path=False,
            )
        )


def test_legacy_quiet_scheduler_invocation_fails_closed_and_is_observable(
    tmp_path: Path,
) -> None:
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")

    _, rejected = run_cli(
        "--home",
        str(home),
        "sync",
        "--quiet",
        expected_code=1,
    )

    assert rejected["code"] == "legacy_scheduler_reinstall_required"
    state = json.loads((home / "state.json").read_text(encoding="utf-8"))
    assert state["last_scheduled_sync_status"] == "legacy_definition_rejected"
    assert (
        state["last_scheduled_sync_error_code"]
        == "legacy_scheduler_reinstall_required"
    )
    assert "last_manual_ghealth_sync" not in state


def test_scheduled_preflight_failure_updates_scheduled_only_status(
    tmp_path: Path,
) -> None:
    home = tmp_path / "private-health-home"
    run_cli("--home", str(home), "init", "--timezone", "UTC")
    now = datetime.now(timezone.utc).isoformat(timespec="seconds")
    atomic_write_json(
        home / "state.json",
        {
            "onboarding_consents": {
                "google-health": {"granted_at": now, "policy_version": 1}
            }
        },
    )

    with pytest.raises(ValueError, match="scheduler consent"):
        health_cli.command_sync(
            Namespace(
                home=str(home),
                from_date="2026-01-01",
                to_date="2026-01-02",
                lookback_days=None,
                fixture_dir=None,
                scheduled=True,
                static_runtime_fingerprint="c" * 64,
                runtime_fingerprint="a" * 64,
                quiet=True,
                verbose_path=False,
            )
        )

    state = json.loads((home / "state.json").read_text(encoding="utf-8"))
    assert state["last_scheduled_sync_status"] == "failed_preflight"
    assert state["last_scheduled_sync_error_code"] == "scheduler_consent_required"
    assert "last_successful_scheduled_sync_at" not in state
