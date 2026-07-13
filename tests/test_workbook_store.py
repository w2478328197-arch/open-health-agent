from __future__ import annotations

import os
from pathlib import Path

from openpyxl import Workbook, load_workbook

from oha.config import create_config
from oha.constants import SHEET_HEADERS
from oha.database import HealthDatabase
from oha.workbook_store import export_workbook, inspect_workbook_schema, prune_backups


def assert_private_mode(path: Path) -> None:
    assert path.exists()
    if os.name != "nt":
        assert path.stat().st_mode & 0o777 == 0o600


def test_export_preserves_non_health_sheet_and_builds_health_views(tmp_path: Path) -> None:
    home = tmp_path / "private-home"
    destination = tmp_path / "健康档案.xlsx"
    workbook = Workbook()
    custom = workbook.active
    custom.title = "我的自定义指标"
    custom.append(["日期", "自定义指标"])
    custom.append(["2026-01-01", 7])
    custom["D4"] = "必须保留"
    workbook.save(destination)
    workbook.close()

    config = create_config(home, workbook=destination, timezone="UTC")
    with HealthDatabase(config.database) as database:
        database.upsert(
            "daily",
            {
                "record_id": "2026-01-02",
                "date": "2026-01-02",
                "steps": 4321,
                "active_energy_kcal": 456,
                "batch_id": "synthetic",
            },
        )
        database.upsert(
            "food",
            {
                "record_id": "food-1",
                "date": "2026-01-02",
                "food_name": "合成测试餐",
                "source": "synthetic",
                "summary_nutrients": {
                    "energy_kcal": 600,
                    "protein_g": 40,
                    "fat_g": 20,
                    "carbohydrate_g": 70,
                },
            },
        )
        result = export_workbook(config, database, tmp_path / "missing-template.xlsx")

    assert result == destination
    assert_private_mode(destination)
    check = load_workbook(destination, data_only=False)
    try:
        assert check["我的自定义指标"]["D4"].value == "必须保留"
        assert set(SHEET_HEADERS).issubset(check.sheetnames)
        daily = check["健康日报"]
        assert daily["A2"].value.strftime("%Y-%m-%d") == "2026-01-02"
        assert daily["B2"].value == 4321
        summary = check["每日营养汇总"]
        assert summary["B2"].value == 1
        assert summary["C2"].value == 600
        assert summary["N2"].value < summary["O2"].value < summary["P2"].value
    finally:
        check.close()


def test_prune_backups_keeps_newest_files(tmp_path: Path) -> None:
    backup_dir = tmp_path / "backups"
    backup_dir.mkdir()
    files = []
    for index in range(5):
        path = backup_dir / f"managed-{index}.xlsx"
        path.write_bytes(b"synthetic")
        os.utime(path, (1000 + index, 1000 + index))
        files.append(path)

    unrelated = backup_dir / "personal.xlsx"
    unrelated.write_bytes(b"must remain")
    prune_backups(backup_dir, "managed", keep=2)

    assert sorted(path.name for path in backup_dir.glob("*.xlsx")) == [
        "managed-3.xlsx",
        "managed-4.xlsx",
        "personal.xlsx",
    ]


def test_read_only_workbook_schema_inspection_reports_managed_header_damage(
    tmp_path: Path,
) -> None:
    path = tmp_path / "health.xlsx"
    workbook = Workbook()
    workbook.active.title = "使用说明"
    for name, headers in SHEET_HEADERS.items():
        sheet = workbook.create_sheet(name)
        sheet.append(headers)
    workbook.save(path)
    workbook.close()
    assert inspect_workbook_schema(path)["ok"] is True

    workbook = load_workbook(path)
    workbook["健康日报"]["A1"] = "损坏的表头"
    workbook.save(path)
    workbook.close()
    result = inspect_workbook_schema(path)
    assert result["ok"] is False
    assert result["reason"] == "managed_schema_mismatch"
    assert result["incompatible_managed_sheets"] == ["健康日报"]


def test_export_escapes_formula_shaped_external_text(tmp_path: Path) -> None:
    config = create_config(tmp_path / "private", workbook=tmp_path / "health.xlsx", timezone="UTC")
    with HealthDatabase(config.database) as database:
        database.upsert(
            "measurement",
            {
                "record_id": "formula-injection-test",
                "date": "2026-01-02",
                "metric": "合成测试",
                "value": 1,
                "unit": "test",
                "source": "@UNTRUSTED",
                "original_text": "  =WEBSERVICE(\"https://invalid.example\")",
            },
        )
        export_workbook(config, database, tmp_path / "missing.xlsx")

    workbook = load_workbook(config.workbook, data_only=False)
    try:
        sheet = workbook["健康测量"]
        assert sheet["H2"].data_type == "s"
        assert sheet["H2"].value == "'@UNTRUSTED"
        assert sheet["K2"].data_type == "s"
        assert sheet["K2"].value.startswith("  '=WEBSERVICE")
        assert not any(
            cell.data_type == "f"
            for row in sheet.iter_rows(min_row=2)
            for cell in row
        )
    finally:
        workbook.close()


def test_export_replaces_xml_control_characters_but_keeps_canonical_record_exact(
    tmp_path: Path,
) -> None:
    config = create_config(tmp_path / "private", workbook=tmp_path / "health.xlsx", timezone="UTC")
    original_text = "早餐\x00包含\x0b设备控制字符\x1f，原始记录必须保留"
    with HealthDatabase(config.database) as database:
        database.upsert(
            "measurement",
            {
                "record_id": "xml-control-character-test",
                "date": "2026-01-02",
                "metric": "合成测试",
                "value": 1,
                "unit": "test",
                "source": "synthetic",
                "original_text": original_text,
            },
        )

        # This used to raise openpyxl IllegalCharacterError after the database
        # transaction had already committed.
        export_workbook(config, database, tmp_path / "missing.xlsx")
        assert database.get("measurement", "xml-control-character-test")["original_text"] == original_text

    workbook = load_workbook(config.workbook, data_only=False)
    try:
        exported = workbook["健康测量"]["K2"]
        assert exported.data_type == "s"
        assert exported.value == "早餐�包含�设备控制字符�，原始记录必须保留"
        assert not any(ord(character) < 0x20 and character not in "\t\r\n" for character in exported.value)
    finally:
        workbook.close()
