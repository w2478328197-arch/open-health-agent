from __future__ import annotations

import hashlib
import json
from argparse import Namespace
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import health_agent as health_cli
from openpyxl import Workbook, load_workbook
import pytest
from oha.config import create_config, initialize_local_files, save_config
from oha.database import HealthDatabase, canonical_json, utc_now
from oha.workbook_store import export_workbook
from oha.runtime_convergence import (
    legacy_writer_status,
    require_scheduler_install_convergence,
)


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "skills" / "open-health-agent" / "assets"
TEMPLATE = ASSETS / "health-ledger.xlsx"


def _synthetic_health_fingerprint(path: Path) -> dict[str, object]:
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        digest = hashlib.sha256()
        counts: dict[str, int] = {}
        for name in ("健康日报", "健康测量", "训练记录", "饮食记录", "饮食营养明细", "每日营养汇总", "健康导入日志"):
            if name not in workbook.sheetnames:
                continue
            populated = [
                list(row)
                for row in workbook[name].iter_rows(values_only=True)
                if any(value not in (None, "") for value in row)
            ]
            counts[name] = max(len(populated) - 1, 0)
            digest.update(canonical_json([name, populated]).encode("utf-8"))
        return {"sha256": digest.hexdigest(), "sheet_rows": counts}
    finally:
        workbook.close()


def _insert_completed_migration_audit(
    database: HealthDatabase,
    config,
    legacy: Path,
    *,
    migration_id: str = "synthetic-migration",
) -> None:
    stat = legacy.stat()
    database.connection.execute(
        """
        INSERT INTO audit_events(occurred_at, action, kind, record_id, payload_json)
        VALUES (?, 'legacy_workbook_migration', NULL, ?, ?)
        """,
        (
            utc_now(),
            migration_id,
            canonical_json(
                {
                    "migration_id": migration_id,
                    "completion": "complete",
                    "target_database": str(config.database),
                    "target_workbook": str(config.workbook),
                    "source": {
                        "path": str(legacy.resolve()),
                        "size": stat.st_size,
                        "mtime_ns": stat.st_mtime_ns,
                        "sha256": hashlib.sha256(legacy.read_bytes()).hexdigest(),
                        "health": _synthetic_health_fingerprint(legacy),
                    },
                }
            ),
        ),
    )
    database.connection.commit()


def _initialized_runtime(home: Path, workbook: Path):
    config = create_config(
        home,
        workbook=workbook,
        timezone="UTC",
        ghealth_command="missing-ghealth-for-synthetic-test",
    )
    save_config(config)
    initialize_local_files(
        config,
        ASSETS / "AGENTS.md.template",
        ASSETS / "profile.example.json",
    )
    with HealthDatabase(config.database) as database:
        export_workbook(config, database, TEMPLATE)
    return config


def test_doctor_warns_for_cloud_home_without_marking_installation_failed(
    tmp_path: Path, monkeypatch
) -> None:
    home = tmp_path / "Library" / "Mobile Documents" / "synthetic-private-home"
    config = _initialized_runtime(home, home / "health.xlsx")
    monkeypatch.setattr(
        health_cli,
        "scheduler_status",
        lambda _config: {"installed": False, "platform": "synthetic"},
    )

    result = health_cli.command_doctor(Namespace(home=str(config.home_path), offline=True))

    cloud = next(item for item in result["checks"] if item["check"] == "storage topology")
    assert result["status"] == "ok"
    assert cloud["required"] is False
    assert cloud["detail"]["private_home_sync_provider"] == "icloud"
    assert cloud["detail"]["canonical_database_sync_provider"] == "icloud"


def test_doctor_reports_confirmed_profile_gaps_without_private_values(
    tmp_path: Path, monkeypatch
) -> None:
    config = _initialized_runtime(tmp_path / "private", tmp_path / "health.xlsx")
    monkeypatch.setattr(
        health_cli,
        "scheduler_status",
        lambda _config: {"installed": False, "platform": "synthetic"},
    )

    result = health_cli.command_doctor(Namespace(home=str(config.home_path), offline=True))

    profile_gap = next(
        item for item in result["checks"] if item["check"] == "confirmed calculation inputs"
    )
    assert result["status"] == "ok"
    assert profile_gap["ok"] is False
    assert profile_gap["required"] is False
    assert profile_gap["detail"] == {
        "missing_fields": ["active_goal", "lean_mass_kg"],
        "effects": [
            "goal-specific planning remains unavailable",
            "FFM-based REE remains unavailable",
        ],
    }
    serialized = str(profile_gap)
    assert "health_constraints" not in serialized
    assert "original_text" not in serialized


def test_doctor_reports_specified_legacy_health_rows_when_database_has_no_sync(
    tmp_path: Path, monkeypatch
) -> None:
    config = _initialized_runtime(tmp_path / "private", tmp_path / "target.xlsx")
    legacy = tmp_path / "legacy.xlsx"
    workbook = Workbook()
    measurement = workbook.active
    measurement.title = "健康测量"
    measurement.append(["记录ID", "日期", "指标", "数值", "单位", "原话"])
    measurement.append(
        [
            "synthetic-id",
            "2026-01-02",
            "synthetic metric",
            42,
            "unit",
            "SYNTHETIC-HEALTH-ROW-MUST-NOT-LEAK",
        ]
    )
    finance = workbook.create_sheet("财务记录")
    finance.append(["SYNTHETIC-FINANCE-CELL-MUST-NOT-LEAK"])
    workbook.save(legacy)
    workbook.close()
    monkeypatch.setattr(
        health_cli,
        "scheduler_status",
        lambda _config: {"installed": False, "platform": "synthetic"},
    )

    result = health_cli.command_doctor(
        Namespace(
            home=str(config.home_path),
            offline=True,
            legacy_workbook=str(legacy),
        )
    )

    coverage = next(
        item for item in result["checks"] if item["check"] == "canonical health sync coverage"
    )
    assert result["status"] == "needs_attention"
    assert coverage["ok"] is False
    assert coverage["required"] is True
    assert coverage["detail"]["health_sheet_rows"] == {"健康测量": 1}
    assert coverage["detail"]["database_sync_run_count"] == 0
    serialized = str(coverage)
    assert "SYNTHETIC-HEALTH-ROW-MUST-NOT-LEAK" not in serialized
    assert "SYNTHETIC-FINANCE-CELL-MUST-NOT-LEAK" not in serialized


def test_doctor_does_not_treat_failed_or_zero_record_sync_as_health_coverage(
    tmp_path: Path, monkeypatch
) -> None:
    config = _initialized_runtime(tmp_path / "private", tmp_path / "target.xlsx")
    legacy = tmp_path / "legacy.xlsx"
    workbook = Workbook()
    workbook.active.title = "健康测量"
    workbook.active.append(["记录ID", "日期", "指标", "数值", "单位"])
    workbook.active.append(
        ["synthetic-id", "2026-01-02", "synthetic metric", 42, "unit"]
    )
    workbook.save(legacy)
    workbook.close()
    with HealthDatabase(config.database) as database:
        database.begin_sync("synthetic-failed", "2026-01-02", "2026-01-02")
        database.finish_sync(
            "synthetic-failed", "failed", 1, 0, 0, None, ["synthetic failure"]
        )
        database.begin_sync("synthetic-empty", "2026-01-02", "2026-01-02")
        database.finish_sync(
            "synthetic-empty", "success", 0, 0, 0, "2026-01-02T00:00:00Z", []
        )
    monkeypatch.setattr(
        health_cli,
        "scheduler_status",
        lambda _config: {"installed": False, "platform": "synthetic"},
    )

    result = health_cli.command_doctor(
        Namespace(
            home=str(config.home_path),
            offline=True,
            legacy_workbook=str(legacy),
        )
    )

    coverage = next(
        item for item in result["checks"] if item["check"] == "canonical health sync coverage"
    )
    assert result["status"] == "needs_attention"
    assert coverage["ok"] is False
    assert coverage["detail"]["database_sync_run_count"] == 2
    assert coverage["detail"]["successful_health_sync_run_count"] == 0


def test_doctor_reports_configured_workbook_mismatch_from_migration_audit(
    tmp_path: Path, monkeypatch
) -> None:
    configured = tmp_path / "configured.xlsx"
    expected = tmp_path / "confirmed-canonical.xlsx"
    legacy = tmp_path / "legacy.xlsx"
    config = _initialized_runtime(tmp_path / "private", configured)
    with HealthDatabase(config.database) as database:
        database.connection.execute(
            """
            INSERT INTO audit_events(occurred_at, action, kind, record_id, payload_json)
            VALUES (?, 'legacy_workbook_migration', NULL, 'synthetic-migration', ?)
            """,
            (
                utc_now(),
                    canonical_json(
                        {
                            "migration_id": "synthetic-migration",
                            "completion": "complete",
                            "target_database": str(config.database),
                            "target_workbook": str(expected.resolve()),
                            "source": {
                                "path": str(legacy.resolve()),
                                "size": 123,
                                "mtime_ns": 456,
                                "sha256": "0" * 64,
                                "health": {"sha256": "1" * 64, "sheet_rows": {}},
                            },
                        }
                ),
            ),
        )
        database.connection.commit()
    monkeypatch.setattr(
        health_cli,
        "scheduler_status",
        lambda _config: {"installed": False, "platform": "synthetic"},
    )

    result = health_cli.command_doctor(Namespace(home=str(config.home_path), offline=True))

    convergence = next(
        item for item in result["checks"] if item["check"] == "workbook convergence"
    )
    assert result["status"] == "needs_attention"
    assert convergence == {
        "check": "workbook convergence",
        "ok": False,
        "required": True,
        "detail": {
            "configured_workbook": str(configured.resolve()),
            "migration_target_workbook": str(expected.resolve()),
            "configured_database": str(config.database),
            "migration_target_database": str(config.database),
            "legacy_source": str(legacy.resolve()),
            "matches_migration_target": False,
            "matches_migration_database": True,
        },
    }


def test_doctor_fails_closed_for_malformed_completed_migration_audit(
    tmp_path: Path, monkeypatch
) -> None:
    config = _initialized_runtime(tmp_path / "private", tmp_path / "target.xlsx")
    with HealthDatabase(config.database) as database:
        database.connection.execute(
            """
            INSERT INTO audit_events(occurred_at, action, kind, record_id, payload_json)
            VALUES (?, 'legacy_workbook_migration', NULL, 'synthetic-corrupt', ?)
            """,
            (utc_now(), "{SYNTHETIC-PRIVATE-CORRUPT-PAYLOAD"),
        )
        database.connection.commit()
    monkeypatch.setattr(
        health_cli,
        "scheduler_status",
        lambda _config: {"installed": False, "platform": "synthetic"},
    )

    result = health_cli.command_doctor(
        Namespace(home=str(config.home_path), offline=True)
    )

    audit = next(
        item for item in result["checks"] if item["check"] == "migration audit integrity"
    )
    assert result["status"] == "needs_attention"
    assert audit == {
        "check": "migration audit integrity",
        "ok": False,
        "required": True,
        "detail": {"present": True, "readable": False, "valid": False},
    }
    assert "SYNTHETIC-PRIVATE-CORRUPT-PAYLOAD" not in str(audit)


def test_doctor_fails_closed_for_structurally_incomplete_migration_audit(
    tmp_path: Path, monkeypatch
) -> None:
    legacy = tmp_path / "legacy.xlsx"
    workbook = Workbook()
    workbook.active.title = "健康测量"
    workbook.active.append(["记录ID", "日期", "指标", "数值", "单位"])
    workbook.save(legacy)
    workbook.close()
    config = _initialized_runtime(tmp_path / "private", tmp_path / "target.xlsx")
    stat = legacy.stat()
    with HealthDatabase(config.database) as database:
        database.connection.execute(
            """
            INSERT INTO audit_events(occurred_at, action, kind, record_id, payload_json)
            VALUES (?, 'legacy_workbook_migration', NULL, 'synthetic-migration', ?)
            """,
            (
                utc_now(),
                canonical_json(
                    {
                        "completion": "complete",
                        "target_database": str(config.database),
                        "target_workbook": str(config.workbook),
                        "source": {
                            "path": str(legacy.resolve()),
                            "size": stat.st_size,
                            "mtime_ns": stat.st_mtime_ns,
                            "sha256": hashlib.sha256(legacy.read_bytes()).hexdigest(),
                            "health": _synthetic_health_fingerprint(legacy),
                        },
                    }
                ),
            ),
        )
        database.connection.commit()
    monkeypatch.setattr(
        health_cli,
        "scheduler_status",
        lambda _config: {"installed": False, "platform": "synthetic"},
    )

    result = health_cli.command_doctor(
        Namespace(home=str(config.home_path), offline=True)
    )

    audit = next(
        item for item in result["checks"] if item["check"] == "migration audit integrity"
    )
    assert result["status"] == "needs_attention"
    assert audit["ok"] is False
    assert audit["required"] is True
    assert audit["detail"] == {"present": True, "readable": True, "valid": False}
    assert not any(
        item["check"] in {"workbook convergence", "legacy workbook freshness"}
        for item in result["checks"]
    )


def test_doctor_fails_closed_while_legacy_migration_postconditions_are_pending(
    tmp_path: Path, monkeypatch
) -> None:
    config = _initialized_runtime(tmp_path / "private", tmp_path / "target.xlsx")
    with HealthDatabase(config.database) as database:
        database.connection.execute(
            """
            INSERT INTO audit_events(occurred_at, action, kind, record_id, payload_json)
            VALUES (?, 'legacy_workbook_migration_pending', NULL, 'synthetic-pending', ?)
            """,
            (
                utc_now(),
                canonical_json(
                    {"private_detail": "SYNTHETIC-PENDING-HEALTH-MUST-NOT-LEAK"}
                ),
            ),
        )
        database.connection.commit()
    monkeypatch.setattr(
        health_cli,
        "scheduler_status",
        lambda _config: {"installed": False, "platform": "synthetic"},
    )

    result = health_cli.command_doctor(
        Namespace(home=str(config.home_path), offline=True)
    )

    pending = next(
        item
        for item in result["checks"]
        if item["check"] == "migration pending postconditions"
    )
    assert result["status"] == "needs_attention"
    assert pending == {
        "check": "migration pending postconditions",
        "ok": False,
        "required": True,
        "detail": {"present": True, "count": 1},
    }
    assert "SYNTHETIC-PENDING-HEALTH-MUST-NOT-LEAK" not in str(pending)


def test_doctor_reports_legacy_freshness_without_leaking_rows(
    tmp_path: Path, monkeypatch
) -> None:
    legacy = tmp_path / "legacy.xlsx"
    workbook = Workbook()
    measurement = workbook.active
    measurement.title = "健康测量"
    measurement.append(["记录ID", "日期", "指标", "数值", "单位", "原话"])
    measurement.append(
        [
            "synthetic-id",
            "2026-01-02",
            "synthetic metric",
            42,
            "unit",
            "SYNTHETIC-ROW-MUST-NOT-LEAK",
        ]
    )
    custom = workbook.create_sheet("财务记录")
    custom.append(["SYNTHETIC-FINANCE-CELL-MUST-NOT-LEAK"])
    workbook.save(legacy)
    workbook.close()

    config = _initialized_runtime(tmp_path / "private", tmp_path / "target.xlsx")
    stat = legacy.stat()
    digest = hashlib.sha256(legacy.read_bytes()).hexdigest()
    with HealthDatabase(config.database) as database:
        database.connection.execute(
            """
            INSERT INTO audit_events(occurred_at, action, kind, record_id, payload_json)
            VALUES (?, 'legacy_workbook_migration', NULL, 'synthetic-migration', ?)
            """,
            (
                utc_now(),
                    canonical_json(
                        {
                            "migration_id": "synthetic-migration",
                            "completion": "complete",
                            "target_database": str(config.database),
                        "target_workbook": str(config.workbook),
                        "source": {
                            "path": str(legacy.resolve()),
                            "size": stat.st_size,
                            "mtime_ns": stat.st_mtime_ns,
                            "sha256": digest,
                            "health": _synthetic_health_fingerprint(legacy),
                        },
                    }
                ),
            ),
        )
        database.connection.commit()
    monkeypatch.setattr(
        health_cli,
        "scheduler_status",
        lambda _config: {"installed": False, "platform": "synthetic"},
    )

    first = health_cli.command_doctor(Namespace(home=str(config.home_path), offline=True))
    freshness = next(
        item for item in first["checks"] if item["check"] == "legacy workbook freshness"
    )
    assert freshness["ok"] is True
    assert freshness["detail"]["health_sheet_rows"] == {"健康测量": 1}
    assert freshness["detail"]["unknown_sheet_count"] == 1
    assert freshness["detail"]["has_unimported_changes"] is False
    assert "SYNTHETIC-ROW-MUST-NOT-LEAK" not in str(freshness)
    assert "SYNTHETIC-FINANCE-CELL-MUST-NOT-LEAK" not in str(freshness)

    finance_changed = load_workbook(legacy)
    finance_changed["财务记录"].append(["SYNTHETIC-FINANCE-ONLY-CHANGE"])
    finance_changed.save(legacy)
    finance_changed.close()
    finance_only = health_cli.command_doctor(
        Namespace(home=str(config.home_path), offline=True)
    )
    finance_freshness = next(
        item
        for item in finance_only["checks"]
        if item["check"] == "legacy workbook freshness"
    )
    assert finance_freshness["ok"] is True
    assert finance_freshness["detail"]["has_unimported_changes"] is False
    assert "SYNTHETIC-FINANCE-ONLY-CHANGE" not in str(finance_freshness)

    changed = load_workbook(legacy)
    changed["健康测量"].append(
        ["synthetic-id-2", "2026-01-03", "synthetic metric", 43, "unit", "hidden"]
    )
    changed.save(legacy)
    changed.close()
    second = health_cli.command_doctor(Namespace(home=str(config.home_path), offline=True))
    changed_freshness = next(
        item for item in second["checks"] if item["check"] == "legacy workbook freshness"
    )
    assert second["status"] == "needs_attention"
    assert changed_freshness["ok"] is False
    assert changed_freshness["detail"]["health_sheet_rows"] == {"健康测量": 2}
    assert changed_freshness["detail"]["has_unimported_changes"] is True


def test_doctor_reports_active_legacy_writer_without_exposing_job_payload(
    tmp_path: Path, monkeypatch
) -> None:
    host_home = tmp_path / "synthetic-host"
    scripts = host_home / ".hermes" / "scripts"
    cron = host_home / ".hermes" / "cron"
    scripts.mkdir(parents=True)
    cron.mkdir(parents=True)
    (scripts / "sync_ghealth_hourly.py").write_text("# synthetic", encoding="utf-8")
    (scripts / "import_health_ledger.py").write_text("# synthetic", encoding="utf-8")
    jobs_path = cron / "jobs.json"
    jobs_path.write_text(
        json.dumps(
            [
                {
                    "id": "synthetic-job",
                    "name": "ghealth每小时同步到iCloud健康台账",
                    "prompt": "PRIVATE-SYNTHETIC-PROMPT-MUST-NOT-LEAK",
                    "script": "sync_ghealth_hourly.py",
                    "enabled": True,
                },
                {
                    "id": "unrelated-job",
                    "name": "unrelated",
                    "prompt": "UNRELATED-PRIVATE-PROMPT-MUST-NOT-LEAK",
                    "enabled": True,
                },
            ]
        ),
        encoding="utf-8",
    )
    config = _initialized_runtime(tmp_path / "private", tmp_path / "target.xlsx")
    monkeypatch.setattr(
        health_cli,
        "scheduler_status",
        lambda _config: {"installed": False, "platform": "synthetic"},
    )
    monkeypatch.setattr(
        health_cli,
        "legacy_writer_status",
        lambda _host_home: legacy_writer_status(host_home),
    )

    result = health_cli.command_doctor(Namespace(home=str(config.home_path), offline=True))

    writer = next(
        item for item in result["checks"] if item["check"] == "legacy health writer"
    )
    assert result["status"] == "needs_attention"
    assert writer["ok"] is False
    assert writer["required"] is True
    assert writer["detail"] == {
        "script_artifact_count": 2,
        "matching_hermes_job_count": 1,
        "active_hermes_job_count": 1,
        "paused_hermes_job_count": 0,
        "unknown_hermes_job_count": 0,
        "legacy_launchagent_definition_count": 0,
        "unknown_legacy_launchagent_definition_count": 0,
        "active_writer_detected": True,
        "state": "active",
    }
    serialized = str(writer)
    assert "PRIVATE-SYNTHETIC-PROMPT-MUST-NOT-LEAK" not in serialized
    assert "UNRELATED-PRIVATE-PROMPT-MUST-NOT-LEAK" not in serialized

    jobs = json.loads(jobs_path.read_text(encoding="utf-8"))
    jobs[0]["enabled"] = False
    jobs_path.write_text(json.dumps(jobs), encoding="utf-8")
    paused = legacy_writer_status(host_home)
    assert paused["ok"] is True
    assert paused["required"] is False
    assert paused["detail"]["paused_hermes_job_count"] == 1
    assert paused["detail"]["state"] == "paused"


def test_bare_legacy_script_artifact_is_unverified_and_unsafe(tmp_path: Path) -> None:
    host_home = tmp_path / "synthetic-host"
    scripts = host_home / ".hermes" / "scripts"
    scripts.mkdir(parents=True)
    (scripts / "sync_ghealth_hourly.py").write_text("# synthetic", encoding="utf-8")

    status = legacy_writer_status(host_home)

    assert status["ok"] is False
    assert status["required"] is True
    assert status["detail"]["state"] == "unverified_artifacts"


def test_unreadable_launchagent_definition_is_unknown_and_unsafe(
    tmp_path: Path, monkeypatch
) -> None:
    host_home = tmp_path / "synthetic-host"
    launchagents = host_home / "Library" / "LaunchAgents"
    launchagents.mkdir(parents=True)
    definition = launchagents / "synthetic.plist"
    definition.write_text("synthetic", encoding="utf-8")
    real_read_bytes = Path.read_bytes

    def guarded_read_bytes(path: Path) -> bytes:
        if path == definition:
            raise PermissionError("synthetic unreadable definition")
        return real_read_bytes(path)

    monkeypatch.setattr(Path, "read_bytes", guarded_read_bytes)

    status = legacy_writer_status(host_home)

    assert status["ok"] is False
    assert status["required"] is True
    assert status["detail"]["state"] == "unknown"
    assert status["detail"]["unknown_legacy_launchagent_definition_count"] == 1


@pytest.mark.parametrize("sync_status", [None, "partial", "empty", "failed"])
def test_scheduler_install_requires_latest_full_success(
    tmp_path: Path, monkeypatch, sync_status: str | None
) -> None:
    config = _initialized_runtime(tmp_path / "private", tmp_path / "target.xlsx")
    if sync_status is not None:
        with HealthDatabase(config.database) as database:
            database.begin_sync("synthetic-sync", "2026-01-01", "2026-01-02")
            database.finish_sync(
                "synthetic-sync", sync_status, 1, 0, 0, "2026-01-02T00:00:00Z", []
            )
    writer = {"ok": True, "required": False, "detail": {"state": "absent"}}
    with HealthDatabase(config.database) as database:
        with pytest.raises(RuntimeError, match="full successful manual sync"):
            require_scheduler_install_convergence(config, database, writer)


def test_scheduler_install_requires_legacy_writer_retirement_then_proceeds(
    tmp_path: Path, monkeypatch
) -> None:
    config = _initialized_runtime(tmp_path / "private", tmp_path / "target.xlsx")
    today = datetime.now(ZoneInfo("UTC")).date().isoformat()
    with HealthDatabase(config.database) as database:
        database.begin_sync("synthetic-success", today, today)
        database.finish_sync(
            "synthetic-success", "success", 1, 0, 0, f"{today}T00:00:00Z", []
        )
    writer_state = {"ok": False, "required": True, "detail": {"state": "active"}}
    with HealthDatabase(config.database) as database:
        with pytest.raises(RuntimeError, match="legacy health writer"):
            require_scheduler_install_convergence(config, database, writer_state)
        writer_state.update(
            {"ok": True, "required": False, "detail": {"state": "paused"}}
        )
        require_scheduler_install_convergence(config, database, writer_state)


def test_scheduler_install_rejects_a_success_created_by_a_scheduled_run(
    tmp_path: Path, monkeypatch
) -> None:
    config = _initialized_runtime(tmp_path / "private", tmp_path / "target.xlsx")
    with HealthDatabase(config.database) as database:
        database.begin_sync(
            "synthetic-scheduled-success",
            "2026-01-01",
            "2026-01-02",
            trigger="scheduled",
        )
        database.finish_sync(
            "synthetic-scheduled-success",
            "success",
            1,
            0,
            0,
            "2026-01-02T00:00:00Z",
            [],
        )
    writer = {"ok": True, "required": False, "detail": {"state": "absent"}}
    with HealthDatabase(config.database) as database:
        with pytest.raises(RuntimeError, match="full successful manual sync"):
            require_scheduler_install_convergence(config, database, writer)

def test_scheduler_install_rejects_manual_success_not_linked_to_latest_migration(
    tmp_path: Path, monkeypatch
) -> None:
    config = _initialized_runtime(tmp_path / "private", tmp_path / "target.xlsx")
    legacy = tmp_path / "legacy.xlsx"
    workbook = Workbook()
    workbook.active.title = "健康测量"
    workbook.active.append(["记录ID", "日期", "指标", "数值", "单位"])
    workbook.active.append(
        ["synthetic-id", "2026-01-02", "synthetic metric", 42, "unit"]
    )
    workbook.save(legacy)
    workbook.close()
    today = datetime.now(ZoneInfo("UTC")).date().isoformat()
    with HealthDatabase(config.database) as database:
        database.begin_sync("synthetic-before-migration", today, today)
        database.finish_sync(
            "synthetic-before-migration", "success", 1, 0, 0, f"{today}T00:00:00Z", []
        )
        _insert_completed_migration_audit(database, config, legacy)
    writer = {"ok": True, "required": False, "detail": {"state": "paused"}}
    with HealthDatabase(config.database) as database:
        with pytest.raises(RuntimeError, match="full successful manual sync"):
            require_scheduler_install_convergence(config, database, writer)

    with HealthDatabase(config.database) as database:
        database.begin_sync(
            "synthetic-after-migration",
            today,
            today,
            migration_id="synthetic-migration",
        )
        database.finish_sync(
            "synthetic-after-migration", "success", 1, 0, 0, f"{today}T00:00:00Z", []
        )
    with HealthDatabase(config.database) as database:
        require_scheduler_install_convergence(config, database, writer)


def test_scheduler_install_rejects_changed_legacy_health_after_migration(
    tmp_path: Path, monkeypatch
) -> None:
    config = _initialized_runtime(tmp_path / "private", tmp_path / "target.xlsx")
    legacy = tmp_path / "legacy.xlsx"
    workbook = Workbook()
    measurement = workbook.active
    measurement.title = "健康测量"
    measurement.append(["记录ID", "日期", "指标", "数值", "单位"])
    measurement.append(
        ["synthetic-id", "2026-01-02", "synthetic metric", 42, "unit"]
    )
    workbook.save(legacy)
    workbook.close()
    today = datetime.now(ZoneInfo("UTC")).date().isoformat()
    with HealthDatabase(config.database) as database:
        _insert_completed_migration_audit(database, config, legacy)
        database.begin_sync(
            "synthetic-after-migration",
            today,
            today,
            migration_id="synthetic-migration",
        )
        database.finish_sync(
            "synthetic-after-migration", "success", 1, 0, 0, f"{today}T00:00:00Z", []
        )
    changed = load_workbook(legacy)
    changed["健康测量"].append(
        ["synthetic-id-2", "2026-01-03", "synthetic metric", 43, "unit"]
    )
    changed.save(legacy)
    changed.close()
    writer = {"ok": True, "required": False, "detail": {"state": "paused"}}
    with HealthDatabase(config.database) as database:
        with pytest.raises(RuntimeError, match="unchanged legacy health sheets"):
            require_scheduler_install_convergence(config, database, writer)


def test_scheduler_install_rejects_pending_legacy_migration_postconditions(
    tmp_path: Path, monkeypatch
) -> None:
    config = _initialized_runtime(tmp_path / "private", tmp_path / "target.xlsx")
    today = datetime.now(ZoneInfo("UTC")).date().isoformat()
    with HealthDatabase(config.database) as database:
        database.connection.execute(
            """
            INSERT INTO audit_events(occurred_at, action, kind, record_id, payload_json)
            VALUES (?, 'legacy_workbook_migration_pending', NULL, 'synthetic-pending', ?)
            """,
            (utc_now(), canonical_json({"completion": "pending_postconditions"})),
        )
        database.connection.commit()
        database.begin_sync("synthetic-manual", today, today)
        database.finish_sync(
            "synthetic-manual", "success", 1, 0, 0, f"{today}T00:00:00Z", []
        )
    writer = {"ok": True, "required": False, "detail": {"state": "paused"}}
    with HealthDatabase(config.database) as database:
        with pytest.raises(RuntimeError, match="pending legacy migration"):
            require_scheduler_install_convergence(config, database, writer)


def test_scheduler_install_rejects_structurally_invalid_migration_audit(
    tmp_path: Path, monkeypatch
) -> None:
    config = _initialized_runtime(tmp_path / "private", tmp_path / "target.xlsx")
    today = datetime.now(ZoneInfo("UTC")).date().isoformat()
    with HealthDatabase(config.database) as database:
        database.connection.execute(
            """
            INSERT INTO audit_events(occurred_at, action, kind, record_id, payload_json)
            VALUES (?, 'legacy_workbook_migration', NULL, 'synthetic-invalid', ?)
            """,
            (
                utc_now(),
                canonical_json(
                    {
                        "completion": "complete",
                        "target_database": str(config.database),
                        "target_workbook": str(config.workbook),
                    }
                ),
            ),
        )
        database.connection.commit()
        database.begin_sync("synthetic-manual", today, today)
        database.finish_sync(
            "synthetic-manual", "success", 1, 0, 0, f"{today}T00:00:00Z", []
        )
    writer = {"ok": True, "required": False, "detail": {"state": "paused"}}
    with HealthDatabase(config.database) as database:
        with pytest.raises(RuntimeError, match="migration audit is invalid"):
            require_scheduler_install_convergence(config, database, writer)
