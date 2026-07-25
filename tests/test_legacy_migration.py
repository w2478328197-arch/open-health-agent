from __future__ import annotations

import hashlib
import json
import os
import sqlite3
import subprocess
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
import pytest

import oha.legacy_migration as legacy_migration
from oha.config import create_config, save_config
from oha.database import HealthDatabase, stable_record_id
from oha.legacy_migration import migrate_legacy_workbook, parse_legacy_workbook


ROOT = Path(__file__).resolve().parents[1]
CLI = ROOT / "skills" / "open-health-agent" / "scripts" / "health_agent.py"
TEMPLATE = ROOT / "skills" / "open-health-agent" / "assets" / "health-ledger.xlsx"


def _digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _run_cli(*arguments: str) -> tuple[subprocess.CompletedProcess[str], dict]:
    completed = subprocess.run(
        [sys.executable, str(CLI), *arguments],
        cwd=ROOT,
        check=False,
        capture_output=True,
        text=True,
    )
    stream = completed.stdout if completed.returncode == 0 else completed.stderr
    return completed, json.loads(stream)


def _make_legacy_workbook(path: Path, record_id: str | None = "legacy-measurement-1") -> None:
    workbook = Workbook()
    finance = workbook.active
    finance.title = "财务记录"
    finance.append(["日期", "类型", "金额"])
    finance.append(
        ["2026-01-01", "SYNTHETIC-FINANCE-CELL-MUST-NOT-LEAK", 12.5]
    )
    custom = workbook.create_sheet("用户自建")
    custom.append(["保留", "公式"])
    custom.append(["synthetic", "=1+1"])
    measurements = workbook.create_sheet("健康测量")
    measurements.append(
        [
            "记录ID",
            "日期",
            "时间",
            "指标",
            "数值",
            "第二数值",
            "单位",
            "来源",
            "记录方式",
            "置信度",
            "原话",
            "备注",
            "导入时间",
        ]
    )
    import_log = workbook.create_sheet("健康导入日志")
    import_log.append(["时间", "状态", "原始载荷"])
    import_log.append(
        [
            "2026-01-02T00:00:00Z",
            "synthetic",
            "SYNTHETIC-PRIVATE-LOG-PAYLOAD-MUST-NOT-LEAK",
        ]
    )
    measurements.append(
        [
            record_id,
            "2026-01-02",
            "08:30",
            "synthetic metric",
            42,
            None,
            "unit",
            "synthetic source",
            "legacy import",
            "confirmed",
            "synthetic statement",
            None,
            "2026-01-02T08:31:00Z",
        ]
    )
    workbook.save(path)
    workbook.close()


def test_legacy_workbook_dry_run_is_read_only_and_reports_sheet_counts(
    tmp_path: Path,
) -> None:
    home = tmp_path / "private"
    target = tmp_path / "target.xlsx"
    source = tmp_path / "legacy.xlsx"
    config = create_config(home, workbook=target, timezone="UTC")
    save_config(config)
    with HealthDatabase(config.database):
        pass
    target_workbook = Workbook()
    target_workbook.active.title = "target sentinel"
    target_workbook.save(target)
    target_workbook.close()
    _make_legacy_workbook(source)

    before = {
        "source": _digest(source),
        "database": _digest(config.database),
        "target": _digest(target),
    }
    completed, result = _run_cli(
        "--home",
        str(home),
        "migrate",
        "legacy-workbook",
        "--source",
        str(source),
        "--dry-run",
    )

    assert completed.returncode == 0, completed.stderr
    assert result["status"] == "dry_run"
    assert result["target_database"] == str(config.database)
    assert result["sheets"]["健康测量"] == {
        "rows_read": 1,
        "importable": 1,
        "skipped": 0,
        "conflicts": 0,
        "warnings": 0,
    }
    assert result["sheets"]["财务记录"]["rows_read"] == 0
    assert result["sheets"]["财务记录"]["importable"] == 0
    assert result["sheets"]["财务记录"]["warnings"] == 1
    assert before == {
        "source": _digest(source),
        "database": _digest(config.database),
        "target": _digest(target),
    }
    serialized = json.dumps(result, ensure_ascii=False)
    assert "SYNTHETIC-FINANCE-CELL-MUST-NOT-LEAK" not in serialized
    assert "SYNTHETIC-PRIVATE-LOG-PAYLOAD-MUST-NOT-LEAK" not in serialized


def test_legacy_dry_run_does_not_create_an_absent_database_or_target(
    tmp_path: Path,
) -> None:
    home = tmp_path / "private"
    source = tmp_path / "legacy.xlsx"
    target = tmp_path / "target.xlsx"
    config = create_config(home, workbook=target, timezone="UTC")
    save_config(config)
    _make_legacy_workbook(source)

    completed, result = _run_cli(
        "--home",
        str(home),
        "migrate",
        "legacy-workbook",
        "--source",
        str(source),
        "--dry-run",
    )

    assert completed.returncode == 0, completed.stderr
    assert result["status"] == "dry_run"
    assert not config.database.exists()
    assert not target.exists()
    assert list((home / "backups").glob("legacy-source-*.xlsx")) == []


def test_legacy_source_hardlinked_to_target_is_rejected(tmp_path: Path) -> None:
    home = tmp_path / "private"
    source = tmp_path / "legacy.xlsx"
    target = tmp_path / "target.xlsx"
    _make_legacy_workbook(source)
    os.link(source, target)
    config = create_config(home, workbook=target, timezone="UTC")
    save_config(config)
    before = _digest(source)

    completed, result = _run_cli(
        "--home",
        str(home),
        "migrate",
        "legacy-workbook",
        "--source",
        str(source),
        "--dry-run",
    )

    assert completed.returncode == 1
    assert result["code"] == "invalid_input"
    assert _digest(source) == before


def test_missing_backup_flag_is_rejected_before_database_creation(tmp_path: Path) -> None:
    home = tmp_path / "private"
    source = tmp_path / "legacy.xlsx"
    target = tmp_path / "target.xlsx"
    config = create_config(home, workbook=target, timezone="UTC")
    save_config(config)
    _make_legacy_workbook(source)

    completed, result = _run_cli(
        "--home",
        str(home),
        "migrate",
        "legacy-workbook",
        "--source",
        str(source),
    )

    assert completed.returncode == 1
    assert result["code"] == "invalid_input"
    assert not config.database.exists()
    assert not target.exists()
    assert list((home / "backups").glob("legacy-source-*.xlsx")) == []


def test_legacy_fact_mapping_preserves_stable_ids_and_blank_values(tmp_path: Path) -> None:
    source = tmp_path / "legacy-facts.xlsx"
    workbook = Workbook()
    measurement = workbook.active
    measurement.title = "健康测量"
    measurement.append(
        [
            "记录ID",
            "日期",
            "时间",
            "指标",
            "数值",
            "第二数值",
            "单位",
            "来源",
            "记录方式",
            "置信度",
            "原话",
            "备注",
            "导入时间",
        ]
    )
    measurement.append(
        [
            "stable-measurement",
            "2026-01-02",
            "8:30:59",
            "synthetic metric",
            42,
            None,
            "unit",
            "synthetic source",
            "legacy text",
            "confirmed",
            "synthetic measurement statement",
            "synthetic note",
            "2026-01-02T08:31:00Z",
        ]
    )
    workout = workbook.create_sheet("训练记录")
    workout.append(
        [
            "记录ID",
            "日期",
            "开始时间",
            "训练类型",
            "时长_min",
            "距离_km",
            "热量_kcal",
            "平均心率_bpm",
            "最大心率_bpm",
            "强度/RPE",
            "训练肌群",
            "训练来源",
            "外部ID",
            "记录方式",
            "置信度",
            "原话",
            "备注",
            "导入时间",
        ]
    )
    workout.append(
        [
            "stable-workout",
            "2026-01-02",
            "18:05:44",
            "synthetic workout",
            30,
            None,
            120,
            100,
            130,
            4,
            "synthetic group",
            "synthetic source",
            "external-synthetic-1",
            "legacy import",
            "confirmed",
            "synthetic workout statement",
            None,
            "2026-01-02T18:40:00Z",
        ]
    )
    food = workbook.create_sheet("饮食记录")
    food.append(
        [
            "记录ID",
            "日期",
            "时间",
            "餐次",
            "食物",
            "食用量估计_g",
            "食用量下限_g",
            "食用量上限_g",
            "热量_kcal",
            "蛋白质_g",
            "数据来源",
            "估算说明",
            "置信度",
            "图片引用",
            "原话",
            "备注",
            "录入时间",
        ]
    )
    food.append(
        [
            "stable-food",
            "2026-01-02",
            "12:00:12",
            "synthetic meal",
            "synthetic food",
            None,
            None,
            None,
            300,
            None,
            "synthetic source",
            "synthetic estimate",
            "medium",
            "synthetic-reference",
            "synthetic food statement",
            None,
            "2026-01-02T12:01:00Z",
        ]
    )
    nutrients = workbook.create_sheet("饮食营养明细")
    nutrients.append(
        [
            "记录ID",
            "营养素ID",
            "营养素",
            "数值",
            "单位",
            "数据来源",
            "每100g或每份",
            "检索日期",
            "备注",
        ]
    )
    nutrients.append(
        [
            "stable-food",
            "synthetic-nutrient",
            "synthetic nutrient",
            7,
            "unit",
            "synthetic source",
            "per serving",
            "2026-01-02",
            None,
        ]
    )
    summary = workbook.create_sheet("每日营养汇总")
    summary.append(["日期", "已记录条目数", "已记录热量_kcal"])
    summary.append(["2026-01-02", 1, 300])
    workbook.save(source)
    workbook.close()

    plan = parse_legacy_workbook(source)
    records = {(kind, payload["record_id"]): payload for kind, payload in plan["records"]}

    mapped_measurement = records[("measurement", "stable-measurement")]
    assert mapped_measurement["time"] == "08:30"
    assert mapped_measurement["second_value"] is None
    assert mapped_measurement["original_text"] == "synthetic measurement statement"
    mapped_workout = records[("workout", "stable-workout")]
    assert mapped_workout["start_time"] == "18:05"
    assert mapped_workout["distance_km"] is None
    assert mapped_workout["external_id"] == "external-synthetic-1"
    mapped_food = records[("food", "stable-food")]
    assert mapped_food["time"] == "12:00"
    assert mapped_food["estimated_grams"] is None
    assert mapped_food["summary_nutrients"]["energy_kcal"] == 300
    assert mapped_food["summary_nutrients"]["protein_g"] is None
    assert mapped_food["nutrients"] == [
        {
            "id": "synthetic-nutrient",
            "name": "synthetic nutrient",
            "amount": 7.0,
            "unit": "unit",
            "source": "synthetic source",
            "basis": "per serving",
            "retrieved_at": "2026-01-02",
            "notes": "",
        }
    ]
    assert plan["sheets"]["每日营养汇总"]["importable"] == 0
    assert plan["sheets"]["每日营养汇总"]["skipped"] == 1


def test_legacy_food_schema_maps_consumed_grams_and_nutrient_details(tmp_path: Path) -> None:
    source = tmp_path / "legacy-food.xlsx"
    workbook = Workbook()
    default = workbook.active
    assert default is not None
    workbook.remove(default)
    food = workbook.create_sheet("饮食记录")
    food.append(
        [
            "记录ID",
            "日期",
            "时间",
            "餐次",
            "食物",
            "摄入量_g",
            "热量_kcal",
            "数据来源",
            "原话",
            "备注",
            "导入时间",
        ]
    )
    food.append(
        [
            "legacy-food-1",
            "2026-01-02",
            "12:00",
            "午餐",
            "synthetic food",
            125,
            300,
            "synthetic source",
            "synthetic original",
            "",
            "2026-01-02T12:01:00Z",
        ]
    )
    nutrients = workbook.create_sheet("饮食营养明细")
    nutrients.append(
        [
            "记录ID",
            "日期",
            "食物",
            "营养素ID",
            "营养素名称",
            "每100g含量",
            "单位",
            "摄入量_g",
            "本次估算量",
            "数据来源",
            "备注",
        ]
    )
    nutrients.append(
        [
            "legacy-food-1",
            "2026-01-02",
            "synthetic food",
            "calcium",
            "钙",
            80,
            "mg",
            125,
            100,
            "synthetic source",
            "synthetic note",
        ]
    )
    workbook.save(source)
    workbook.close()

    plan = parse_legacy_workbook(source)
    record = next(payload for kind, payload in plan["records"] if kind == "food")

    assert record["estimated_grams"] == 125
    assert record["recorded_at"] == "2026-01-02T12:01:00Z"
    assert record["nutrients"] == [
        {
            "id": "calcium",
            "name": "钙",
            "amount": 100.0,
            "unit": "mg",
            "source": "synthetic source",
            "basis": "80",
            "retrieved_at": "",
            "notes": "synthetic note",
        }
    ]
    assert plan["sheets"]["饮食营养明细"]["importable"] == 1
    assert plan["sheets"]["饮食营养明细"]["warnings"] == 0


def test_real_migration_is_idempotent_and_generates_deterministic_ids(
    tmp_path: Path,
) -> None:
    home = tmp_path / "private"
    target = tmp_path / "target.xlsx"
    source = tmp_path / "legacy-idless.xlsx"
    config = create_config(home, workbook=target, timezone="UTC")
    save_config(config)
    with HealthDatabase(config.database):
        pass
    _make_legacy_workbook(source, record_id=None)
    source_digest = _digest(source)

    first_completed, first = _run_cli(
        "--home",
        str(home),
        "migrate",
        "legacy-workbook",
        "--source",
        str(source),
        "--backup-source",
    )
    second_completed, second = _run_cli(
        "--home",
        str(home),
        "migrate",
        "legacy-workbook",
        "--source",
        str(source),
        "--backup-source",
    )

    assert first_completed.returncode == 0, first_completed.stderr
    assert second_completed.returncode == 0, second_completed.stderr
    assert first["status"] == "migrated"
    assert second["status"] == "already_migrated"
    backup = Path(first["verified_source_backup"])
    assert backup.is_file()
    assert _digest(backup) == source_digest == _digest(source)
    expected_id = stable_record_id(
        "measurement",
        [
            "2026-01-02",
            "08:30",
            "synthetic metric",
                "unit",
                "synthetic source",
                "legacy import",
                "synthetic statement",
        ],
    )
    with HealthDatabase(config.database) as database:
        rows = database.list_records("measurement")
        assert [row["record_id"] for row in rows] == [expected_id]
        migration_audits = database.connection.execute(
            "SELECT COUNT(*) FROM audit_events WHERE action='legacy_workbook_migration'"
        ).fetchone()[0]
        assert migration_audits == 1
    assert target.is_file()


def test_completed_migration_repeat_ignores_source_mtime_only_change(
    tmp_path: Path,
) -> None:
    home = tmp_path / "private"
    source = tmp_path / "legacy-mtime-completed.xlsx"
    config = create_config(home, workbook=tmp_path / "target.xlsx", timezone="UTC")
    save_config(config)
    _make_legacy_workbook(source)
    original_digest = _digest(source)

    with HealthDatabase(config.database) as database:
        first = migrate_legacy_workbook(
            config,
            database,
            source,
            backup_source=True,
            template=TEMPLATE,
        )
        stat = source.stat()
        os.utime(
            source,
            ns=(stat.st_atime_ns, stat.st_mtime_ns + 1_000_000_000),
        )
        repeated = migrate_legacy_workbook(
            config,
            database,
            source,
            backup_source=True,
            template=TEMPLATE,
        )

    assert _digest(source) == original_digest
    assert first["status"] == "migrated"
    assert repeated["status"] == "already_migrated"


def test_real_migration_rejects_malformed_completed_audit_instead_of_claiming_idempotency(
    tmp_path: Path,
) -> None:
    home = tmp_path / "private"
    source = tmp_path / "legacy-malformed-completed.xlsx"
    config = create_config(home, workbook=tmp_path / "target.xlsx", timezone="UTC")
    save_config(config)
    _make_legacy_workbook(source)
    migration_id = stable_record_id(
        "legacy_migration", [str(source.resolve()), _digest(source)]
    )
    with HealthDatabase(config.database) as database:
        database.connection.execute(
            """
            INSERT INTO audit_events(occurred_at, action, kind, record_id, payload_json)
            VALUES ('2026-01-02T00:00:00Z', 'legacy_workbook_migration', NULL, ?, ?)
            """,
            (
                migration_id,
                json.dumps(
                    {"migration_id": migration_id, "completion": "complete"},
                    ensure_ascii=False,
                ),
            ),
        )
        database.connection.commit()

        with pytest.raises(ValueError, match="completed legacy migration audit is invalid"):
            migrate_legacy_workbook(
                config,
                database,
                source,
                backup_source=True,
                template=TEMPLATE,
            )

    assert not (home / "backups").exists()


def test_real_migration_rejects_malformed_pending_audit_instead_of_finalizing_it(
    tmp_path: Path,
) -> None:
    home = tmp_path / "private"
    source = tmp_path / "legacy-malformed-pending.xlsx"
    target = tmp_path / "target.xlsx"
    config = create_config(home, workbook=target, timezone="UTC")
    save_config(config)
    _make_legacy_workbook(source)
    migration_id = stable_record_id(
        "legacy_migration", [str(source.resolve()), _digest(source)]
    )
    with HealthDatabase(config.database) as database:
        database.connection.execute(
            """
            INSERT INTO audit_events(occurred_at, action, kind, record_id, payload_json)
            VALUES ('2026-01-02T00:00:00Z', 'legacy_workbook_migration_pending', NULL, ?, ?)
            """,
            (
                migration_id,
                json.dumps(
                    {
                        "migration_id": migration_id,
                        "completion": "pending_postconditions",
                        "target_database": str(config.database),
                        "target_workbook": str(config.workbook),
                    },
                    ensure_ascii=False,
                ),
            ),
        )
        database.connection.commit()

        with pytest.raises(ValueError, match="pending legacy migration audit is invalid"):
            migrate_legacy_workbook(
                config,
                database,
                source,
                backup_source=True,
                template=TEMPLATE,
            )

        actions = [
            row["action"]
            for row in database.connection.execute(
                "SELECT action FROM audit_events WHERE record_id=?", (migration_id,)
            ).fetchall()
        ]
    assert actions == ["legacy_workbook_migration_pending"]
    assert not target.exists()


def test_legacy_sleep_total_minus_awake_maps_to_actual_asleep_with_provenance(
    tmp_path: Path,
) -> None:
    source = tmp_path / "legacy-sleep-derived.xlsx"
    workbook = Workbook()
    daily = workbook.active
    daily.title = "健康日报"
    daily.append(
        [
            "日期",
            "步数",
            "睡眠总时长_h",
            "清醒_min",
            "睡眠来源",
            "健康数据来源",
            "数据质量",
            "备注",
        ]
    )
    daily.append(
        [
            "2026-01-02",
            1234,
            8.0,
            30,
            "synthetic sleep source",
            "synthetic health source",
            "legacy-complete",
            "synthetic note",
        ]
    )
    workbook.save(source)
    workbook.close()

    plan = parse_legacy_workbook(source)
    daily_records = [payload for kind, payload in plan["records"] if kind == "daily"]

    assert len(daily_records) == 1
    mapped = daily_records[0]
    assert mapped["record_id"] == "2026-01-02"
    assert mapped["sleep_hours"] == 7.5
    assert mapped["awake_minutes"] == 30
    assert mapped["steps"] == 1234
    assert "legacy_sleep_derived_from_total_minus_awake" in mapped["quality"]
    assert "睡眠总时长_h - 清醒_min" in mapped["notes"]
    assert plan["sheets"]["健康日报"]["importable"] == 1


def test_legacy_sleep_total_without_awake_stays_null_and_warns(tmp_path: Path) -> None:
    source = tmp_path / "legacy-sleep-uncertain.xlsx"
    workbook = Workbook()
    daily = workbook.active
    daily.title = "健康日报"
    daily.append(["日期", "睡眠总时长_h", "数据质量"])
    daily.append(["2026-01-03", 8.0, "legacy-complete"])
    workbook.save(source)
    workbook.close()

    plan = parse_legacy_workbook(source)
    mapped = next(payload for kind, payload in plan["records"] if kind == "daily")

    assert mapped["sleep_hours"] is None
    assert "legacy_sleep_asleep_unknown" in mapped["quality"]
    assert plan["sheets"]["健康日报"]["warnings"] == 1
    assert plan["warnings"] == [
        {
            "code": "legacy_sleep_total_without_awake",
            "sheet": "健康日报",
            "row": 2,
        }
    ]


def test_stable_id_conflict_is_counted_and_blocks_real_migration(tmp_path: Path) -> None:
    home = tmp_path / "private"
    target = tmp_path / "target.xlsx"
    source = tmp_path / "legacy-conflict.xlsx"
    config = create_config(home, workbook=target, timezone="UTC")
    save_config(config)
    with HealthDatabase(config.database) as database:
        database.upsert(
            "measurement",
            {
                "record_id": "legacy-measurement-1",
                "date": "2026-01-02",
                "time": "08:30",
                "metric": "synthetic metric",
                "value": 41,
                "second_value": None,
                "unit": "unit",
                "source": "synthetic source",
                "method": "legacy import",
                "confidence": "confirmed",
                "original_text": "synthetic statement",
                "notes": "",
                "imported_at": "2026-01-02T08:31:00Z",
            },
        )
    _make_legacy_workbook(source)
    database_before = _digest(config.database)
    source_before = _digest(source)

    dry_completed, dry = _run_cli(
        "--home",
        str(home),
        "migrate",
        "legacy-workbook",
        "--source",
        str(source),
        "--dry-run",
    )
    real_completed, real = _run_cli(
        "--home",
        str(home),
        "migrate",
        "legacy-workbook",
        "--source",
        str(source),
        "--backup-source",
    )

    assert dry_completed.returncode == 0, dry_completed.stderr
    assert dry["sheets"]["健康测量"]["importable"] == 0
    assert dry["sheets"]["健康测量"]["conflicts"] == 1
    assert dry["conflicts"] == [
        {
            "code": "stable_id_payload_conflict",
            "kind": "measurement",
            "sheet": "健康测量",
        }
    ]
    assert real_completed.returncode == 1
    assert real["code"] == "legacy_migration_conflict"
    assert _digest(config.database) == database_before
    assert _digest(source) == source_before
    assert not target.exists()


def test_conflicting_duplicate_stable_ids_inside_source_block_migration(
    tmp_path: Path,
) -> None:
    home = tmp_path / "private"
    source = tmp_path / "legacy-internal-conflict.xlsx"
    config = create_config(home, workbook=tmp_path / "target.xlsx", timezone="UTC")
    save_config(config)
    with HealthDatabase(config.database):
        pass
    _make_legacy_workbook(source, record_id="same-source-id")
    workbook = load_workbook(source)
    workbook["健康测量"].append(
        [
            "same-source-id",
            "2026-01-02",
            "08:30",
            "synthetic metric",
            99,
            None,
            "unit",
            "synthetic source",
            "legacy import",
            "confirmed",
            "synthetic statement",
            None,
            "2026-01-02T08:31:00Z",
        ]
    )
    workbook.save(source)
    workbook.close()

    dry_completed, dry = _run_cli(
        "--home",
        str(home),
        "migrate",
        "legacy-workbook",
        "--source",
        str(source),
        "--dry-run",
    )
    real_completed, real = _run_cli(
        "--home",
        str(home),
        "migrate",
        "legacy-workbook",
        "--source",
        str(source),
        "--backup-source",
    )

    assert dry_completed.returncode == 0, dry_completed.stderr
    assert dry["sheets"]["健康测量"]["conflicts"] == 1
    assert dry["conflicts"] == [
        {
            "code": "stable_id_payload_conflict",
            "kind": "measurement",
            "sheet": "健康测量",
        }
    ]
    assert real_completed.returncode == 1
    assert real["code"] == "legacy_migration_conflict"
    with HealthDatabase(config.database) as database:
        assert database.list_records("measurement") == []


def test_food_semantic_duplicate_with_small_estimate_change_is_skipped(
    tmp_path: Path,
) -> None:
    home = tmp_path / "private"
    source = tmp_path / "legacy-food-duplicate.xlsx"
    config = create_config(home, workbook=tmp_path / "target.xlsx", timezone="UTC")
    save_config(config)
    with HealthDatabase(config.database) as database:
        database.upsert(
            "food",
            {
                "record_id": "canonical-food-id",
                "date": "2026-01-02",
                "time": "12:00",
                "meal": "synthetic meal",
                "food_name": "synthetic food",
                "estimated_grams": 100.0,
                "grams_low": None,
                "grams_high": None,
                "summary_nutrients": {"energy_kcal": 300.0},
                "nutrients": [],
                "source": "synthetic source",
                "estimation_notes": "synthetic estimate",
                "confidence": "medium",
                "image_reference": "",
                "original_text": "synthetic food statement",
                "notes": "",
                "recorded_at": "2026-01-02T12:01:00Z",
                "consumed": True,
            },
        )
    workbook = Workbook()
    food = workbook.active
    food.title = "饮食记录"
    food.append(
        [
            "记录ID",
            "日期",
            "时间",
            "餐次",
            "食物",
            "食用量估计_g",
            "热量_kcal",
            "数据来源",
            "估算说明",
            "置信度",
            "原话",
            "录入时间",
        ]
    )
    food.append(
        [
            "different-legacy-id",
            "2026-01-02",
            "12:00",
            "synthetic meal",
            "synthetic food",
            102.0,
            305.0,
            "synthetic source",
            "synthetic estimate",
            "medium",
            "synthetic food statement",
            "2026-01-02T12:01:00Z",
        ]
    )
    workbook.save(source)
    workbook.close()

    completed, result = _run_cli(
        "--home",
        str(home),
        "migrate",
        "legacy-workbook",
        "--source",
        str(source),
        "--dry-run",
    )

    assert completed.returncode == 0, completed.stderr
    assert result["sheets"]["饮食记录"]["importable"] == 0
    assert result["sheets"]["饮食记录"]["skipped"] == 1
    assert result["sheets"]["饮食记录"]["conflicts"] == 0
    assert {
        "code": "semantic_duplicate_estimate_variation",
        "kind": "food",
        "sheet": "饮食记录",
    } in result["warnings"]
    with HealthDatabase(config.database) as database:
        assert [row["record_id"] for row in database.list_records("food")] == [
            "canonical-food-id"
        ]


def test_food_semantic_identity_keeps_distinct_sources_separate(
    tmp_path: Path,
) -> None:
    home = tmp_path / "private"
    source = tmp_path / "legacy-food-distinct-source.xlsx"
    config = create_config(home, workbook=tmp_path / "target.xlsx", timezone="UTC")
    save_config(config)
    with HealthDatabase(config.database) as database:
        database.upsert(
            "food",
            {
                "record_id": "canonical-food-id",
                "date": "2026-01-02",
                "time": "12:00",
                "meal": "synthetic meal",
                "food_name": "synthetic food",
                "estimated_grams": 100.0,
                "grams_low": None,
                "grams_high": None,
                "summary_nutrients": {"energy_kcal": 300.0},
                "nutrients": [],
                "source": "synthetic source one",
                "estimation_notes": "synthetic estimate",
                "confidence": "medium",
                "image_reference": "",
                "original_text": "synthetic food statement",
                "notes": "",
                "recorded_at": "2026-01-02T12:01:00Z",
                "consumed": True,
            },
        )
    workbook = Workbook()
    food = workbook.active
    food.title = "饮食记录"
    food.append(
        [
            "记录ID",
            "日期",
            "时间",
            "餐次",
            "食物",
            "食用量估计_g",
            "热量_kcal",
            "数据来源",
            "估算说明",
            "置信度",
            "原话",
            "录入时间",
        ]
    )
    food.append(
        [
            "different-legacy-id",
            "2026-01-02",
            "12:00",
            "synthetic meal",
            "synthetic food",
            100.0,
            300.0,
            "synthetic source two",
            "synthetic estimate",
            "medium",
            "synthetic food statement",
            "2026-01-02T12:01:00Z",
        ]
    )
    workbook.save(source)
    workbook.close()

    completed, result = _run_cli(
        "--home",
        str(home),
        "migrate",
        "legacy-workbook",
        "--source",
        str(source),
        "--dry-run",
    )

    assert completed.returncode == 0, completed.stderr
    assert result["sheets"]["饮食记录"]["importable"] == 1
    assert result["sheets"]["饮食记录"]["skipped"] == 0
    assert result["sheets"]["饮食记录"]["conflicts"] == 0


def test_idless_food_records_from_distinct_sources_get_distinct_ids(
    tmp_path: Path,
) -> None:
    home = tmp_path / "private"
    source = tmp_path / "legacy-idless-food-sources.xlsx"
    config = create_config(home, workbook=tmp_path / "target.xlsx", timezone="UTC")
    save_config(config)
    workbook = Workbook()
    food = workbook.active
    food.title = "饮食记录"
    food.append(
        [
            "记录ID",
            "日期",
            "时间",
            "餐次",
            "食物",
            "食用量估计_g",
            "数据来源",
            "原话",
        ]
    )
    for source_name in ("synthetic source one", "synthetic source two"):
        food.append(
            [
                None,
                "2026-01-02",
                "12:00",
                "synthetic meal",
                "synthetic food",
                100.0,
                source_name,
                "synthetic food statement",
            ]
        )
    workbook.save(source)
    workbook.close()

    plan = parse_legacy_workbook(source)
    food_ids = [
        payload["record_id"]
        for kind, payload in plan["records"]
        if kind == "food"
    ]
    completed, result = _run_cli(
        "--home",
        str(home),
        "migrate",
        "legacy-workbook",
        "--source",
        str(source),
        "--dry-run",
    )

    assert len(food_ids) == 2
    assert len(set(food_ids)) == 2
    assert completed.returncode == 0, completed.stderr
    assert result["sheets"]["饮食记录"]["importable"] == 2
    assert result["sheets"]["饮食记录"]["conflicts"] == 0


def test_food_missing_estimate_is_not_close_to_an_arbitrary_large_value(
    tmp_path: Path,
) -> None:
    home = tmp_path / "private"
    source = tmp_path / "legacy-food-large-change.xlsx"
    config = create_config(home, workbook=tmp_path / "target.xlsx", timezone="UTC")
    save_config(config)
    with HealthDatabase(config.database) as database:
        database.upsert(
            "food",
            {
                "record_id": "canonical-food-id",
                "date": "2026-01-02",
                "time": "12:00",
                "meal": "synthetic meal",
                "food_name": "synthetic food",
                "estimated_grams": None,
                "grams_low": None,
                "grams_high": None,
                "summary_nutrients": {},
                "nutrients": [],
                "source": "synthetic source",
                "estimation_notes": "",
                "confidence": "legacy",
                "image_reference": "",
                "original_text": "synthetic food statement",
                "notes": "",
                "recorded_at": "2026-01-02T12:01:00Z",
                "consumed": True,
            },
        )
    workbook = Workbook()
    food = workbook.active
    food.title = "饮食记录"
    food.append(
        [
            "记录ID",
            "日期",
            "时间",
            "餐次",
            "食物",
            "食用量估计_g",
            "热量_kcal",
            "数据来源",
            "置信度",
            "原话",
            "录入时间",
        ]
    )
    food.append(
        [
            "different-legacy-id",
            "2026-01-02",
            "12:00",
            "synthetic meal",
            "synthetic food",
            10000.0,
            5000.0,
            "synthetic source",
            "legacy",
            "synthetic food statement",
            "2026-01-02T12:01:00Z",
        ]
    )
    workbook.save(source)
    workbook.close()

    completed, result = _run_cli(
        "--home",
        str(home),
        "migrate",
        "legacy-workbook",
        "--source",
        str(source),
        "--dry-run",
    )

    assert completed.returncode == 0, completed.stderr
    assert result["sheets"]["饮食记录"]["skipped"] == 0
    assert result["sheets"]["饮食记录"]["conflicts"] == 1


def test_food_detail_difference_is_not_discarded_as_estimate_variation(
    tmp_path: Path,
) -> None:
    home = tmp_path / "private"
    source = tmp_path / "legacy-food-detail-change.xlsx"
    config = create_config(home, workbook=tmp_path / "target.xlsx", timezone="UTC")
    save_config(config)
    with HealthDatabase(config.database) as database:
        database.upsert(
            "food",
            {
                "record_id": "canonical-food-id",
                "date": "2026-01-02",
                "time": "12:00",
                "meal": "synthetic meal",
                "food_name": "synthetic food",
                "estimated_grams": 100.0,
                "grams_low": None,
                "grams_high": None,
                "summary_nutrients": {"energy_kcal": 300.0},
                "nutrients": [],
                "source": "synthetic source",
                "estimation_notes": "synthetic estimate",
                "confidence": "medium",
                "image_reference": "",
                "original_text": "synthetic food statement",
                "notes": "",
                "recorded_at": "2026-01-02T12:01:00Z",
                "consumed": True,
            },
        )
    workbook = Workbook()
    food = workbook.active
    food.title = "饮食记录"
    food.append(
        [
            "记录ID",
            "日期",
            "时间",
            "餐次",
            "食物",
            "食用量估计_g",
            "热量_kcal",
            "数据来源",
            "估算说明",
            "置信度",
            "原话",
            "录入时间",
        ]
    )
    food.append(
        [
            "different-legacy-id",
            "2026-01-02",
            "12:00",
            "synthetic meal",
            "synthetic food",
            102.0,
            305.0,
            "synthetic source",
            "synthetic estimate",
            "medium",
            "synthetic food statement",
            "2026-01-02T12:01:00Z",
        ]
    )
    detail = workbook.create_sheet("饮食营养明细")
    detail.append(
        [
            "记录ID",
            "营养素ID",
            "营养素",
            "数值",
            "单位",
            "数据来源",
            "每100g或每份",
            "检索日期",
            "备注",
        ]
    )
    detail.append(
        [
            "different-legacy-id",
            "synthetic-nutrient",
            "synthetic nutrient",
            5.0,
            "mg",
            "synthetic source",
            "per serving",
            "2026-01-02",
            "",
        ]
    )
    workbook.save(source)
    workbook.close()

    completed, result = _run_cli(
        "--home",
        str(home),
        "migrate",
        "legacy-workbook",
        "--source",
        str(source),
        "--dry-run",
    )

    assert completed.returncode == 0, completed.stderr
    assert result["sheets"]["饮食记录"]["skipped"] == 0
    assert result["sheets"]["饮食记录"]["conflicts"] == 1


def test_invalid_negative_nutrient_detail_is_skipped_with_warning(tmp_path: Path) -> None:
    source = tmp_path / "legacy-invalid-nutrient.xlsx"
    workbook = Workbook()
    food = workbook.active
    food.title = "饮食记录"
    food.append(["记录ID", "日期", "食物", "数据来源", "原话"])
    food.append(
        [
            "synthetic-food-id",
            "2026-01-02",
            "synthetic food",
            "synthetic source",
            "synthetic food statement",
        ]
    )
    detail = workbook.create_sheet("饮食营养明细")
    detail.append(["记录ID", "营养素ID", "营养素", "数值", "单位"])
    detail.append(
        [
            "synthetic-food-id",
            "synthetic-nutrient",
            "synthetic nutrient",
            -5.0,
            "mg",
        ]
    )
    workbook.save(source)
    workbook.close()

    plan = parse_legacy_workbook(source)
    mapped = next(payload for kind, payload in plan["records"] if kind == "food")

    assert mapped["nutrients"] == []
    assert plan["sheets"]["饮食营养明细"]["importable"] == 0
    assert plan["sheets"]["饮食营养明细"]["skipped"] == 1
    assert plan["sheets"]["饮食营养明细"]["warnings"] == 1


def test_legacy_import_log_is_counted_without_materializing_payload_rows(
    tmp_path: Path, monkeypatch
) -> None:
    source = tmp_path / "legacy-import-log.xlsx"
    workbook = Workbook()
    log = workbook.active
    log.title = "健康导入日志"
    log.append(["时间", "状态", "原始载荷"])
    log.append(
        [
            "2026-01-02T00:00:00Z",
            "synthetic",
            "SYNTHETIC-PRIVATE-LOG-PAYLOAD-MUST-NOT-BE-READ",
        ]
    )
    workbook.save(source)
    workbook.close()
    real_rows = legacy_migration._rows

    def guarded_rows(sheet):
        if sheet.title == "健康导入日志":
            raise AssertionError("legacy import log payload was materialized")
        return real_rows(sheet)

    monkeypatch.setattr(legacy_migration, "_rows", guarded_rows)

    plan = parse_legacy_workbook(source)

    assert plan["sheets"]["健康导入日志"] == {
        "rows_read": 1,
        "importable": 0,
        "skipped": 1,
        "conflicts": 0,
        "warnings": 0,
    }
    assert "SYNTHETIC-PRIVATE-LOG-PAYLOAD-MUST-NOT-BE-READ" not in json.dumps(
        plan, ensure_ascii=False
    )


def test_idless_measurement_matches_existing_semantic_identity(tmp_path: Path) -> None:
    home = tmp_path / "private"
    source = tmp_path / "legacy-measurement-semantic.xlsx"
    config = create_config(home, workbook=tmp_path / "target.xlsx", timezone="UTC")
    save_config(config)
    with HealthDatabase(config.database) as database:
        database.upsert(
            "measurement",
            {
                "record_id": "canonical-custom-id",
                "date": "2026-01-02",
                "time": "08:30",
                "metric": "synthetic metric",
                "value": 42.0,
                "second_value": None,
                "unit": "unit",
                "source": "synthetic source",
                "method": "legacy import",
                "confidence": "confirmed",
                "original_text": "synthetic statement",
                "notes": "",
                "imported_at": "2026-01-02T08:31:00Z",
            },
        )
    _make_legacy_workbook(source, record_id=None)

    completed, result = _run_cli(
        "--home",
        str(home),
        "migrate",
        "legacy-workbook",
        "--source",
        str(source),
        "--dry-run",
    )

    assert completed.returncode == 0, completed.stderr
    assert result["sheets"]["健康测量"]["importable"] == 0
    assert result["sheets"]["健康测量"]["skipped"] == 1
    assert result["sheets"]["健康测量"]["conflicts"] == 0


def test_daily_migration_merges_complementary_existing_fields(tmp_path: Path) -> None:
    home = tmp_path / "private"
    source = tmp_path / "legacy-daily-complement.xlsx"
    config = create_config(home, workbook=tmp_path / "target.xlsx", timezone="UTC")
    save_config(config)
    with HealthDatabase(config.database) as database:
        database.upsert(
            "daily",
            {
                "record_id": "2026-01-02",
                "date": "2026-01-02",
                "weight_kg": 70.0,
                "quality": "synthetic manual measurement",
            },
        )
    workbook = Workbook()
    daily = workbook.active
    daily.title = "健康日报"
    daily.append(["日期", "步数", "实际睡眠时长_h", "数据质量"])
    daily.append(["2026-01-02", 1234, 7.5, "synthetic legacy import"])
    workbook.save(source)
    workbook.close()

    dry_completed, dry = _run_cli(
        "--home",
        str(home),
        "migrate",
        "legacy-workbook",
        "--source",
        str(source),
        "--dry-run",
    )
    real_completed, real = _run_cli(
        "--home",
        str(home),
        "migrate",
        "legacy-workbook",
        "--source",
        str(source),
        "--backup-source",
    )

    assert dry_completed.returncode == 0, dry_completed.stderr
    assert dry["sheets"]["健康日报"]["importable"] == 1
    assert dry["sheets"]["健康日报"]["conflicts"] == 0
    assert real_completed.returncode == 0, real_completed.stderr
    assert real["status"] == "migrated"
    with HealthDatabase(config.database) as database:
        rows = database.list_records("daily")
        assert len(rows) == 1
        assert rows[0]["weight_kg"] == 70.0
        assert rows[0]["steps"] == 1234
        assert rows[0]["sleep_hours"] == 7.5


def test_migration_rolls_back_all_records_when_a_planned_insert_fails(
    tmp_path: Path,
) -> None:
    home = tmp_path / "private"
    source = tmp_path / "legacy-rollback.xlsx"
    target = tmp_path / "target.xlsx"
    config = create_config(home, workbook=target, timezone="UTC")
    save_config(config)
    workbook = Workbook()
    measurement = workbook.active
    measurement.title = "健康测量"
    measurement.append(
        [
            "记录ID",
            "日期",
            "时间",
            "指标",
            "数值",
            "单位",
            "来源",
            "原话",
        ]
    )
    measurement.append(
        [
            "rollback-one",
            "2026-01-02",
            "08:00",
            "synthetic metric one",
            1,
            "unit",
            "synthetic source",
            "synthetic first statement",
        ]
    )
    measurement.append(
        [
            "rollback-two",
            "2026-01-02",
            "09:00",
            "synthetic metric two",
            2,
            "unit",
            "synthetic source",
            "synthetic second statement",
        ]
    )
    workbook.save(source)
    workbook.close()
    source_before = _digest(source)

    with HealthDatabase(config.database) as database:
        database.connection.execute(
            """
            CREATE TRIGGER synthetic_migration_failure
            BEFORE INSERT ON records
            WHEN NEW.record_id='rollback-two'
            BEGIN
                SELECT RAISE(ABORT, 'synthetic migration failure');
            END
            """
        )
        database.connection.commit()

        try:
            migrate_legacy_workbook(
                config,
                database,
                source,
                backup_source=True,
                template=TEMPLATE,
            )
        except sqlite3.IntegrityError:
            pass
        else:
            raise AssertionError("synthetic migration failure was not raised")

        assert database.list_records("measurement") == []
        assert database.connection.execute(
            "SELECT COUNT(*) FROM audit_events WHERE action='legacy_workbook_migration'"
        ).fetchone()[0] == 0

    assert _digest(source) == source_before
    assert not target.exists()


def test_successful_migration_exports_and_reports_only_integrity_flags(
    tmp_path: Path,
) -> None:
    home = tmp_path / "private"
    source = tmp_path / "legacy-integrity.xlsx"
    target = tmp_path / "target.xlsx"
    config = create_config(home, workbook=target, timezone="UTC")
    save_config(config)
    with HealthDatabase(config.database):
        pass
    _make_legacy_workbook(source)

    completed, result = _run_cli(
        "--home",
        str(home),
        "migrate",
        "legacy-workbook",
        "--source",
        str(source),
        "--backup-source",
    )

    assert completed.returncode == 0, completed.stderr
    assert result["integrity"] == {"sqlite": "ok", "context": "ok"}
    assert target.is_file()
    serialized = json.dumps(result, ensure_ascii=False)
    assert "synthetic statement" not in serialized
    assert '"value": 42' not in serialized
    assert "SYNTHETIC-FINANCE-CELL-MUST-NOT-LEAK" not in serialized
    assert "SYNTHETIC-PRIVATE-LOG-PAYLOAD-MUST-NOT-LEAK" not in serialized
    with HealthDatabase(config.database) as database:
        migration_audit = database.connection.execute(
            """
            SELECT payload_json FROM audit_events
            WHERE action='legacy_workbook_migration'
            ORDER BY event_id DESC LIMIT 1
            """
        ).fetchone()[0]
    assert "SYNTHETIC-FINANCE-CELL-MUST-NOT-LEAK" not in migration_audit
    assert "SYNTHETIC-PRIVATE-LOG-PAYLOAD-MUST-NOT-LEAK" not in migration_audit


def test_export_failure_leaves_migration_resumable_instead_of_completed(
    tmp_path: Path, monkeypatch
) -> None:
    home = tmp_path / "private"
    source = tmp_path / "legacy-resumable.xlsx"
    target = tmp_path / "target.xlsx"
    config = create_config(home, workbook=target, timezone="UTC")
    save_config(config)
    _make_legacy_workbook(source)

    with HealthDatabase(config.database) as database:
        real_export = legacy_migration.export_workbook

        def fail_export(*_args, **_kwargs):
            raise RuntimeError("synthetic export failure")

        monkeypatch.setattr(legacy_migration, "export_workbook", fail_export)
        try:
            migrate_legacy_workbook(
                config,
                database,
                source,
                backup_source=True,
                template=TEMPLATE,
            )
        except RuntimeError as exc:
            assert str(exc) == "synthetic export failure"
        else:
            raise AssertionError("synthetic export failure was not raised")

        monkeypatch.setattr(legacy_migration, "export_workbook", real_export)
        resumed = migrate_legacy_workbook(
            config,
            database,
            source,
            backup_source=True,
            template=TEMPLATE,
        )
        repeated = migrate_legacy_workbook(
            config,
            database,
            source,
            backup_source=True,
            template=TEMPLATE,
        )

        assert resumed["status"] == "migrated"
        assert resumed["integrity"] == {"sqlite": "ok", "context": "ok"}
        assert repeated["status"] == "already_migrated"
        assert target.is_file()
        assert database.connection.execute(
            "SELECT COUNT(*) FROM audit_events WHERE action='legacy_workbook_migration'"
        ).fetchone()[0] == 1


def test_pending_migration_resume_ignores_source_mtime_only_change(
    tmp_path: Path, monkeypatch
) -> None:
    home = tmp_path / "private"
    source = tmp_path / "legacy-mtime-pending.xlsx"
    config = create_config(home, workbook=tmp_path / "target.xlsx", timezone="UTC")
    save_config(config)
    _make_legacy_workbook(source)
    original_digest = _digest(source)

    with HealthDatabase(config.database) as database:
        real_export = legacy_migration.export_workbook

        def fail_export(*_args, **_kwargs):
            raise RuntimeError("synthetic export failure")

        monkeypatch.setattr(legacy_migration, "export_workbook", fail_export)
        with pytest.raises(RuntimeError, match="synthetic export failure"):
            migrate_legacy_workbook(
                config,
                database,
                source,
                backup_source=True,
                template=TEMPLATE,
            )
        stat = source.stat()
        os.utime(
            source,
            ns=(stat.st_atime_ns, stat.st_mtime_ns + 1_000_000_000),
        )
        monkeypatch.setattr(legacy_migration, "export_workbook", real_export)

        resumed = migrate_legacy_workbook(
            config,
            database,
            source,
            backup_source=True,
            template=TEMPLATE,
        )

    assert _digest(source) == original_digest
    assert resumed["status"] == "migrated"


def test_pending_migration_resume_rejects_missing_imported_record(
    tmp_path: Path, monkeypatch
) -> None:
    home = tmp_path / "private"
    source = tmp_path / "legacy-pending-record-missing.xlsx"
    config = create_config(home, workbook=tmp_path / "target.xlsx", timezone="UTC")
    save_config(config)
    _make_legacy_workbook(source)

    with HealthDatabase(config.database) as database:
        real_export = legacy_migration.export_workbook

        def fail_export(*_args, **_kwargs):
            raise RuntimeError("synthetic export failure")

        monkeypatch.setattr(legacy_migration, "export_workbook", fail_export)
        with pytest.raises(RuntimeError, match="synthetic export failure"):
            migrate_legacy_workbook(
                config,
                database,
                source,
                backup_source=True,
                template=TEMPLATE,
            )
        assert database.delete(
            "measurement", "legacy-measurement-1", "synthetic interruption"
        )
        monkeypatch.setattr(legacy_migration, "export_workbook", real_export)

        with pytest.raises(
            ValueError, match="pending legacy migration records no longer match"
        ):
            migrate_legacy_workbook(
                config,
                database,
                source,
                backup_source=True,
                template=TEMPLATE,
            )

        pending_count = database.connection.execute(
            """
            SELECT COUNT(*) FROM audit_events
            WHERE action='legacy_workbook_migration_pending'
            """
        ).fetchone()[0]
        completed_count = database.connection.execute(
            """
            SELECT COUNT(*) FROM audit_events
            WHERE action='legacy_workbook_migration'
            """
        ).fetchone()[0]

    assert pending_count == 1
    assert completed_count == 0


def test_pending_migration_resume_rejects_semantic_replacement_with_different_id(
    tmp_path: Path, monkeypatch
) -> None:
    home = tmp_path / "private"
    source = tmp_path / "legacy-pending-record-replaced.xlsx"
    config = create_config(home, workbook=tmp_path / "target.xlsx", timezone="UTC")
    save_config(config)
    _make_legacy_workbook(source)

    with HealthDatabase(config.database) as database:
        real_export = legacy_migration.export_workbook

        def fail_export(*_args, **_kwargs):
            raise RuntimeError("synthetic export failure")

        monkeypatch.setattr(legacy_migration, "export_workbook", fail_export)
        with pytest.raises(RuntimeError, match="synthetic export failure"):
            migrate_legacy_workbook(
                config,
                database,
                source,
                backup_source=True,
                template=TEMPLATE,
            )
        original = database.get("measurement", "legacy-measurement-1")
        assert original is not None
        assert database.delete(
            "measurement", "legacy-measurement-1", "synthetic replacement"
        )
        replacement = dict(original)
        replacement["record_id"] = "synthetic-replacement-id"
        database.upsert("measurement", replacement)
        monkeypatch.setattr(legacy_migration, "export_workbook", real_export)

        with pytest.raises(
            ValueError, match="pending legacy migration records no longer match"
        ):
            migrate_legacy_workbook(
                config,
                database,
                source,
                backup_source=True,
                template=TEMPLATE,
            )

        assert database.get("measurement", "legacy-measurement-1") is None
        assert database.get("measurement", "synthetic-replacement-id") is not None
        pending_count = database.connection.execute(
            """
            SELECT COUNT(*) FROM audit_events
            WHERE action='legacy_workbook_migration_pending'
            """
        ).fetchone()[0]
        completed_count = database.connection.execute(
            """
            SELECT COUNT(*) FROM audit_events
            WHERE action='legacy_workbook_migration'
            """
        ).fetchone()[0]

    assert pending_count == 1
    assert completed_count == 0
