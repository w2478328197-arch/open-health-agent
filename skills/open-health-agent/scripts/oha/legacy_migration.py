from __future__ import annotations

import hashlib
import json
import math
import os
import re
import shutil
import sqlite3
import tempfile
from copy import deepcopy
from datetime import date, datetime, time, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .config import Config
from .constants import DAILY_FIELD_MAP, SUMMARY_NUTRIENT_FIELDS
from .context import build_context
from .database import HealthDatabase, canonical_json, stable_record_id
from .recording import normalize_record
from .workbook_store import export_workbook


LEGACY_HEALTH_SHEETS = (
    "健康日报",
    "健康测量",
    "训练记录",
    "饮食记录",
    "饮食营养明细",
    "每日营养汇总",
    "健康导入日志",
)
LEGACY_DERIVED_OR_LOG_SHEETS = {"每日营养汇总", "健康导入日志"}

_KIND_SHEETS = {
    "daily": "健康日报",
    "measurement": "健康测量",
    "workout": "训练记录",
    "food": "饮食记录",
}
_COMPARISON_METADATA = {"batch_id", "data_until", "imported_at", "recorded_at"}


class LegacyMigrationConflictError(RuntimeError):
    """A migration plan contains unresolved record collisions."""


def _empty_sheet_summary() -> dict[str, int]:
    return {
        "rows_read": 0,
        "importable": 0,
        "skipped": 0,
        "conflicts": 0,
        "warnings": 0,
    }


def _rows(sheet) -> tuple[list[str], list[tuple[Any, ...]]]:
    iterator = sheet.iter_rows(values_only=True)
    first = next(iterator, ())
    headers = [str(value).strip() if value is not None else "" for value in first]
    rows = [tuple(row) for row in iterator if any(value not in (None, "") for value in row)]
    return headers, rows


def _row_dict(headers: list[str], row: tuple[Any, ...]) -> dict[str, Any]:
    return {header: row[index] for index, header in enumerate(headers) if header and index < len(row)}


def _text_cell(value: Any) -> str:
    if value in (None, ""):
        return ""
    return value if isinstance(value, str) else str(value)


def _date_cell(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return _text_cell(value)


def _timestamp_cell(value: Any) -> str:
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    return _text_cell(value)


_TIME_WITH_SECONDS = re.compile(r"^(\d{1,2}):(\d{2}):\d{2}(?:\.\d+)?$")


def _time_cell(value: Any) -> str:
    if isinstance(value, datetime):
        return value.strftime("%H:%M")
    if isinstance(value, time):
        return value.strftime("%H:%M")
    result = _text_cell(value).strip()
    match = _TIME_WITH_SECONDS.fullmatch(result)
    return f"{int(match.group(1)):02d}:{match.group(2)}" if match else result


def _optional_number(value: Any, label: str) -> float | int | None:
    if value in (None, ""):
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"{label} must be numeric")
    if not math.isfinite(float(value)):
        raise ValueError(f"{label} must be finite")
    return value


_DAILY_NUMERIC_FIELDS = {
    "steps",
    "distance_km",
    "active_energy_kcal",
    "active_minutes",
    "sleep_hours",
    "deep_minutes",
    "light_minutes",
    "rem_minutes",
    "awake_minutes",
    "resting_heart_rate_bpm",
    "hrv_ms",
    "spo2_percent",
    "respiratory_rate",
    "vo2max",
    "weight_kg",
    "body_fat_percent",
    "height_cm",
}


def _append_provenance(existing: str, value: str) -> str:
    return "; ".join(item for item in (existing.strip(), value) if item)


def _daily(row: dict[str, Any]) -> tuple[dict[str, Any], str | None]:
    day = _date_cell(row.get("日期"))
    date.fromisoformat(day)
    payload: dict[str, Any] = {"record_id": day, "date": day}
    for field, header in DAILY_FIELD_MAP.items():
        if field in {"record_id", "date", "sleep_hours"} or header not in row:
            continue
        value = row.get(header)
        if field in _DAILY_NUMERIC_FIELDS:
            payload[field] = _optional_number(value, header)
        elif field in {"data_until", "imported_at"}:
            payload[field] = _timestamp_cell(value)
        else:
            payload[field] = _text_cell(value)

    explicit_asleep = _optional_number(row.get("实际睡眠时长_h"), "实际睡眠时长_h")
    total_session = _optional_number(row.get("睡眠总时长_h"), "睡眠总时长_h")
    warning = None
    if explicit_asleep is not None:
        payload["sleep_hours"] = explicit_asleep
    elif total_session is not None and payload.get("awake_minutes") is not None:
        payload["sleep_hours"] = max(
            float(total_session) * 60 - float(payload["awake_minutes"]), 0.0
        ) / 60
        payload["quality"] = _append_provenance(
            _text_cell(payload.get("quality")),
            "legacy_sleep_derived_from_total_minus_awake",
        )
        payload["notes"] = _append_provenance(
            _text_cell(payload.get("notes")),
            "legacy provenance: 睡眠总时长_h - 清醒_min",
        )
    elif total_session is not None:
        payload["sleep_hours"] = None
        payload["quality"] = _append_provenance(
            _text_cell(payload.get("quality")), "legacy_sleep_asleep_unknown"
        )
        payload["notes"] = _append_provenance(
            _text_cell(payload.get("notes")),
            "legacy provenance: total session retained without actual-asleep inference",
        )
        warning = "legacy_sleep_total_without_awake"
    return payload, warning


def _measurement(row: dict[str, Any]) -> dict[str, Any]:
    return normalize_record(
        "measurement",
        {
            "record_id": _text_cell(row.get("记录ID")),
            "date": _date_cell(row.get("日期")),
            "time": _time_cell(row.get("时间")),
            "metric": _text_cell(row.get("指标")),
            "value": row.get("数值"),
            "second_value": row.get("第二数值"),
            "unit": _text_cell(row.get("单位")),
            "source": _text_cell(row.get("来源")) or "legacy workbook",
            "method": _text_cell(row.get("记录方式")) or "legacy-workbook-migration",
            "confidence": _text_cell(row.get("置信度")) or "legacy",
            "original_text": _text_cell(row.get("原话")),
            "notes": _text_cell(row.get("备注")),
            "imported_at": _timestamp_cell(row.get("导入时间")),
        },
    )


def _workout(row: dict[str, Any]) -> dict[str, Any]:
    return normalize_record(
        "workout",
        {
            "record_id": _text_cell(row.get("记录ID")),
            "date": _date_cell(row.get("日期")),
            "start_time": _time_cell(row.get("开始时间")),
            "workout_type": _text_cell(row.get("训练类型")),
            "duration_minutes": row.get("时长_min"),
            "distance_km": row.get("距离_km"),
            "calories_kcal": row.get("热量_kcal"),
            "average_heart_rate_bpm": row.get("平均心率_bpm"),
            "max_heart_rate_bpm": row.get("最大心率_bpm"),
            "intensity_rpe": row.get("强度/RPE"),
            "muscle_groups": _text_cell(row.get("训练肌群")),
            "training_source": _text_cell(row.get("训练来源")) or "legacy workbook",
            "external_id": _text_cell(row.get("外部ID")),
            "method": _text_cell(row.get("记录方式")) or "legacy-workbook-migration",
            "confidence": _text_cell(row.get("置信度")) or "legacy",
            "original_text": _text_cell(row.get("原话")),
            "notes": _text_cell(row.get("备注")),
            "imported_at": _timestamp_cell(row.get("导入时间")),
        },
    )


def _food(row: dict[str, Any]) -> dict[str, Any]:
    nutrients = {
        field: row.get(header)
        for field, header in SUMMARY_NUTRIENT_FIELDS.items()
        if header in row
    }
    return normalize_record(
        "food",
        {
            "record_id": _text_cell(row.get("记录ID")),
            "date": _date_cell(row.get("日期")),
            "time": _time_cell(row.get("时间")),
            "meal": _text_cell(row.get("餐次")),
            "food_name": _text_cell(row.get("食物")),
            "estimated_grams": (
                row.get("食用量估计_g")
                if row.get("食用量估计_g") not in (None, "")
                else row.get("摄入量_g")
            ),
            "grams_low": row.get("食用量下限_g"),
            "grams_high": row.get("食用量上限_g"),
            "summary_nutrients": nutrients,
            "source": _text_cell(row.get("数据来源")) or "legacy workbook",
            "estimation_notes": _text_cell(row.get("估算说明")),
            "confidence": _text_cell(row.get("置信度")) or "legacy",
            "image_reference": _text_cell(row.get("图片引用")),
            "original_text": _text_cell(row.get("原话")),
            "notes": _text_cell(row.get("备注")),
            "recorded_at": _timestamp_cell(
                row.get("录入时间")
                if row.get("录入时间") not in (None, "")
                else row.get("导入时间")
            ),
            "consumed": True,
        },
    )


def _nutrient(row: dict[str, Any]) -> dict[str, Any]:
    amount = (
        row.get("数值")
        if row.get("数值") not in (None, "")
        else row.get("本次估算量")
    )
    if (
        isinstance(amount, bool)
        or not isinstance(amount, (int, float))
        or not math.isfinite(float(amount))
        or float(amount) < 0
    ):
        raise ValueError("legacy nutrient amount must be a finite non-negative number")
    return {
        "id": _text_cell(row.get("营养素ID")),
        "name": _text_cell(
            row.get("营养素")
            if row.get("营养素") not in (None, "")
            else row.get("营养素名称")
        ),
        "amount": float(amount),
        "unit": _text_cell(row.get("单位")),
        "source": _text_cell(row.get("数据来源")),
        "basis": _text_cell(
            row.get("每100g或每份")
            if row.get("每100g或每份") not in (None, "")
            else row.get("每100g含量")
        ),
        "retrieved_at": _date_cell(row.get("检索日期")),
        "notes": _text_cell(row.get("备注")),
    }


def parse_legacy_workbook(source: Path) -> dict[str, Any]:
    source = source.expanduser().resolve()
    if source.suffix.lower() != ".xlsx" or not source.is_file():
        raise ValueError("legacy source must be an existing .xlsx workbook")

    workbook = load_workbook(source, read_only=True, data_only=True, keep_links=False)
    try:
        summaries = {name: _empty_sheet_summary() for name in LEGACY_HEALTH_SHEETS}
        warnings: list[dict[str, Any]] = []
        records: list[tuple[str, dict[str, Any]]] = []
        foods_by_id: dict[str, dict[str, Any]] = {}
        nutrient_rows: list[tuple[int, dict[str, Any]]] = []
        for name in workbook.sheetnames:
            if name not in LEGACY_HEALTH_SHEETS:
                summary = _empty_sheet_summary()
                summary["warnings"] = 1
                summaries[name] = summary
                warnings.append({"code": "non_health_sheet_ignored", "sheet": name})
                continue
            summary = summaries[name]
            if name in {"每日营养汇总", "健康导入日志"}:
                row_count = max((workbook[name].max_row or 0) - 1, 0)
                summary["rows_read"] = row_count
                summary["skipped"] = row_count
                continue
            headers, rows = _rows(workbook[name])
            summary["rows_read"] = len(rows)
            for row_number, values in enumerate(rows, start=2):
                try:
                    row = _row_dict(headers, values)
                    row_warning = None
                    if name == "健康日报":
                        daily_payload, row_warning = _daily(row)
                        mapped = ("daily", daily_payload)
                    elif name == "健康测量":
                        mapped = ("measurement", _measurement(row))
                    elif name == "训练记录":
                        mapped = ("workout", _workout(row))
                    elif name == "饮食记录":
                        payload = _food(row)
                        mapped = ("food", payload)
                        foods_by_id[payload["record_id"]] = payload
                    elif name == "饮食营养明细":
                        nutrient_rows.append((row_number, row))
                        continue
                    else:
                        summary["skipped"] += 1
                        continue
                except (TypeError, ValueError):
                    summary["skipped"] += 1
                    summary["warnings"] += 1
                    warnings.append(
                        {
                            "code": "invalid_structured_health_row",
                            "sheet": name,
                            "row": row_number,
                        }
                    )
                else:
                    summary["importable"] += 1
                    records.append(mapped)
                    if row_warning:
                        summary["warnings"] += 1
                        warnings.append(
                            {"code": row_warning, "sheet": name, "row": row_number}
                        )

        nutrient_summary = summaries["饮食营养明细"]
        for row_number, row in nutrient_rows:
            record_id = _text_cell(row.get("记录ID"))
            food = foods_by_id.get(record_id)
            try:
                if food is None:
                    raise ValueError("legacy nutrient has no matching food record")
                nutrient = _nutrient(row)
                if not nutrient["id"] or not nutrient["name"] or not nutrient["unit"]:
                    raise ValueError("legacy nutrient identity is incomplete")
            except (TypeError, ValueError):
                nutrient_summary["skipped"] += 1
                nutrient_summary["warnings"] += 1
                warnings.append(
                    {
                        "code": "unmatched_or_invalid_nutrient_detail",
                        "sheet": "饮食营养明细",
                        "row": row_number,
                    }
                )
            else:
                food.setdefault("nutrients", []).append(nutrient)
                nutrient_summary["importable"] += 1
    finally:
        workbook.close()

    return {"records": records, "sheets": summaries, "warnings": warnings}


def _existing_records_read_only(path: Path) -> list[tuple[str, dict[str, Any]]]:
    if not path.is_file():
        return []
    connection = sqlite3.connect(f"{path.resolve().as_uri()}?mode=ro", uri=True)
    connection.row_factory = sqlite3.Row
    try:
        rows = connection.execute("SELECT kind, payload_json FROM records").fetchall()
    finally:
        connection.close()
    return [(row["kind"], json.loads(row["payload_json"])) for row in rows]


def _business_payload(payload: dict[str, Any]) -> dict[str, Any]:
    return {key: value for key, value in payload.items() if key not in _COMPARISON_METADATA}


def _semantic_business_payload(payload: dict[str, Any]) -> dict[str, Any]:
    return {
        key: value
        for key, value in payload.items()
        if key not in _COMPARISON_METADATA
        and key != "record_id"
        and not (
            key in {"source_event_id", "source_event_item_id"}
            and value in (None, "")
        )
        and not (key == "method" and value in (None, "", "manual"))
    }


def _semantic_identity(kind: str, payload: dict[str, Any]) -> str:
    if kind == "daily":
        parts = [payload.get("date")]
    elif kind == "measurement":
        parts = [
            payload.get("date"),
            payload.get("time"),
            payload.get("metric"),
            payload.get("unit"),
            payload.get("source"),
            payload.get("original_text"),
        ]
    elif kind == "workout":
        external_id = payload.get("external_id")
        parts = (
            ["external", external_id]
            if external_id
            else [
                payload.get("date"),
                payload.get("start_time"),
                payload.get("workout_type"),
                payload.get("training_source"),
                payload.get("original_text"),
            ]
        )
    else:
        parts = [
            payload.get("date"),
            payload.get("time"),
            payload.get("meal"),
            payload.get("food_name"),
            payload.get("source"),
            payload.get("original_text"),
        ]
    return canonical_json([kind, *parts])


def _numbers_close(first: Any, second: Any) -> bool:
    if first is None or second is None:
        return first is None and second is None
    if isinstance(first, bool) or isinstance(second, bool):
        return first == second
    if not isinstance(first, (int, float)) or not isinstance(second, (int, float)):
        return first == second
    return math.isclose(float(first), float(second), rel_tol=0.10, abs_tol=1.0)


def _food_estimates_close(first: dict[str, Any], second: dict[str, Any]) -> bool:
    estimate_fields = {
        "record_id",
        "estimated_grams",
        "grams_low",
        "grams_high",
        "summary_nutrients",
        *_COMPARISON_METADATA,
    }
    def fixed(payload: dict[str, Any]) -> dict[str, Any]:
        return {
            key: value
            for key, value in payload.items()
            if key not in estimate_fields
            and not (
                key in {"source_event_id", "source_event_item_id"}
                and value in (None, "")
            )
            and not (key == "method" and value in (None, "", "manual"))
        }

    first_fixed = fixed(first)
    second_fixed = fixed(second)
    if canonical_json(first_fixed) != canonical_json(second_fixed):
        return False
    for field in ("estimated_grams", "grams_low", "grams_high"):
        if not _numbers_close(first.get(field), second.get(field)):
            return False
    first_summary = first.get("summary_nutrients") or {}
    second_summary = second.get("summary_nutrients") or {}
    if not isinstance(first_summary, dict) or not isinstance(second_summary, dict):
        return False
    return all(
        _numbers_close(first_summary.get(field), second_summary.get(field))
        for field in set(first_summary) | set(second_summary)
    )


def _is_blank(value: Any) -> bool:
    return value is None or value == "" or value == []


def _merge_daily(
    existing: dict[str, Any], incoming: dict[str, Any]
) -> tuple[dict[str, Any], list[str]]:
    merged = dict(existing)
    conflicts: list[str] = []
    provenance_fields = {"quality", "notes", "sources"}
    metadata_fields = {"batch_id", "data_until", "imported_at"}
    for field, value in incoming.items():
        prior = merged.get(field)
        if _is_blank(value):
            continue
        if _is_blank(prior):
            merged[field] = value
            continue
        if prior == value:
            continue
        if field in provenance_fields:
            merged[field] = _append_provenance(_text_cell(prior), _text_cell(value))
        elif field in metadata_fields:
            continue
        elif field not in {"record_id", "date"}:
            conflicts.append(field)
    return merged, conflicts


def _assess_plan(
    plan: dict[str, Any], existing_records: list[tuple[str, dict[str, Any]]]
) -> dict[str, Any]:
    assessed = {
        "records": [],
        "sheets": deepcopy(plan["sheets"]),
        "warnings": list(plan["warnings"]),
        "conflicts": [],
    }
    existing_by_id = {
        (kind, str(payload.get("record_id") or "")): payload
        for kind, payload in existing_records
        if payload.get("record_id")
    }
    existing_by_identity: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for kind, payload in existing_records:
        existing_by_identity.setdefault(
            (kind, _semantic_identity(kind, payload)), []
        ).append(payload)

    def accept(kind: str, payload: dict[str, Any]) -> None:
        assessed["records"].append((kind, payload))
        existing_by_id[(kind, str(payload["record_id"]))] = payload
        existing_by_identity.setdefault(
            (kind, _semantic_identity(kind, payload)), []
        ).append(payload)

    for kind, payload in plan["records"]:
        sheet = _KIND_SHEETS[kind]
        existing = existing_by_id.get((kind, str(payload["record_id"])))
        if existing is None:
            semantic_matches = existing_by_identity.get(
                (kind, _semantic_identity(kind, payload)), []
            )
            if kind == "food" and any(
                _food_estimates_close(candidate, payload)
                for candidate in semantic_matches
            ):
                assessed["sheets"][sheet]["importable"] -= 1
                assessed["sheets"][sheet]["skipped"] += 1
                assessed["sheets"][sheet]["warnings"] += 1
                assessed["warnings"].append(
                    {
                        "code": "semantic_duplicate_estimate_variation",
                        "kind": kind,
                        "sheet": sheet,
                    }
                )
                continue
            if any(
                canonical_json(_semantic_business_payload(candidate))
                == canonical_json(_semantic_business_payload(payload))
                for candidate in semantic_matches
            ):
                assessed["sheets"][sheet]["importable"] -= 1
                assessed["sheets"][sheet]["skipped"] += 1
                continue
            if semantic_matches:
                assessed["sheets"][sheet]["importable"] -= 1
                assessed["sheets"][sheet]["conflicts"] += 1
                assessed["conflicts"].append(
                    {
                        "code": "semantic_identity_payload_conflict",
                        "kind": kind,
                        "sheet": sheet,
                    }
                )
                continue
            accept(kind, payload)
            continue
        if kind == "daily":
            merged, conflicting_fields = _merge_daily(existing, payload)
            if not conflicting_fields:
                if canonical_json(_business_payload(existing)) == canonical_json(
                    _business_payload(merged)
                ):
                    assessed["sheets"][sheet]["importable"] -= 1
                    assessed["sheets"][sheet]["skipped"] += 1
                else:
                    accept(kind, merged)
                continue
            assessed["sheets"][sheet]["importable"] -= 1
            assessed["sheets"][sheet]["conflicts"] += 1
            assessed["conflicts"].append(
                {
                    "code": "daily_field_conflict",
                    "kind": kind,
                    "sheet": sheet,
                    "fields": sorted(conflicting_fields),
                }
            )
            continue
        assessed["sheets"][sheet]["importable"] -= 1
        if canonical_json(_business_payload(existing)) == canonical_json(
            _business_payload(payload)
        ):
            assessed["sheets"][sheet]["skipped"] += 1
            continue
        assessed["sheets"][sheet]["conflicts"] += 1
        assessed["conflicts"].append(
            {
                "code": "stable_id_payload_conflict",
                "kind": kind,
                "sheet": sheet,
            }
        )
    deduplicated: dict[tuple[str, str], tuple[str, dict[str, Any]]] = {}
    for kind, payload in assessed["records"]:
        deduplicated[(kind, str(payload["record_id"]))] = (kind, payload)
    assessed["records"] = list(deduplicated.values())
    return assessed


def _source_is_target(source: Path, target: Path) -> bool:
    if source == target:
        return True
    try:
        return source.samefile(target)
    except (FileNotFoundError, OSError):
        return False


def validate_legacy_source(config: Config, source: Path) -> Path:
    source = source.expanduser().resolve()
    if source.suffix.lower() != ".xlsx" or not source.is_file():
        raise ValueError("legacy source must be an existing .xlsx workbook")
    if _source_is_target(source, config.workbook):
        raise ValueError("legacy source and managed target workbook must be different files")
    return source


def preview_legacy_workbook(config: Config, source: Path) -> dict[str, Any]:
    source = validate_legacy_source(config, source)
    plan = _assess_plan(
        parse_legacy_workbook(source), _existing_records_read_only(config.database)
    )

    return {
        "status": "dry_run",
        "source": str(source),
        "target_database": str(config.database),
        "target_workbook": str(config.workbook),
        "sheets": plan["sheets"],
        "warnings": plan["warnings"],
        "conflicts": plan["conflicts"],
    }


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _fingerprint_cell(value: Any) -> Any:
    if isinstance(value, datetime):
        return {"type": "datetime", "value": value.isoformat()}
    if isinstance(value, date):
        return {"type": "date", "value": value.isoformat()}
    if isinstance(value, time):
        return {"type": "time", "value": value.isoformat()}
    if isinstance(value, float) and not math.isfinite(value):
        return {"type": "float", "value": repr(value)}
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return {"type": type(value).__name__, "value": str(value)}


def legacy_health_fingerprint(path: Path) -> dict[str, Any]:
    """Hash only recognized health-cell values, never unrelated sheet payloads."""

    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    digest = hashlib.sha256()
    sheet_rows: dict[str, int] = {}
    try:
        for name in LEGACY_HEALTH_SHEETS:
            if name not in workbook.sheetnames:
                continue
            if name in LEGACY_DERIVED_OR_LOG_SHEETS:
                row_count = max((workbook[name].max_row or 0) - 1, 0)
                sheet_rows[name] = row_count
                digest.update(canonical_json([name, row_count]).encode("utf-8"))
                continue
            populated = [
                [_fingerprint_cell(value) for value in row]
                for row in workbook[name].iter_rows(values_only=True)
                if any(value not in (None, "") for value in row)
            ]
            sheet_rows[name] = max(len(populated) - 1, 0)
            digest.update(canonical_json([name, populated]).encode("utf-8"))
        return {
            "sha256": digest.hexdigest(),
            "sheet_rows": sheet_rows,
            "unknown_sheet_count": sum(
                name not in LEGACY_HEALTH_SHEETS for name in workbook.sheetnames
            ),
        }
    finally:
        workbook.close()


def _source_fingerprint(source: Path) -> dict[str, Any]:
    stat = source.stat()
    return {
        "path": str(source),
        "size": stat.st_size,
        "mtime_ns": stat.st_mtime_ns,
        "sha256": _sha256(source),
        "health": legacy_health_fingerprint(source),
    }


def _migration_id(fingerprint: dict[str, Any]) -> str:
    return stable_record_id(
        "legacy_migration", [fingerprint["path"], fingerprint["sha256"]]
    )


def _completed_migration(database: HealthDatabase, migration_id: str) -> dict[str, Any] | None:
    row = database.connection.execute(
        """
        SELECT payload_json FROM audit_events
        WHERE action='legacy_workbook_migration' AND record_id=?
        ORDER BY event_id DESC LIMIT 1
        """,
        (migration_id,),
    ).fetchone()
    if row is None:
        return None
    try:
        payload = json.loads(row["payload_json"])
    except (TypeError, json.JSONDecodeError):
        return {"_invalid": True}
    return payload if isinstance(payload, dict) else {"_invalid": True}


def _pending_migration(database: HealthDatabase, migration_id: str) -> dict[str, Any] | None:
    row = database.connection.execute(
        """
        SELECT payload_json FROM audit_events
        WHERE action='legacy_workbook_migration_pending' AND record_id=?
        ORDER BY event_id DESC LIMIT 1
        """,
        (migration_id,),
    ).fetchone()
    if row is None:
        return None
    try:
        payload = json.loads(row["payload_json"])
    except (TypeError, json.JSONDecodeError):
        return {"_invalid": True}
    return payload if isinstance(payload, dict) else {"_invalid": True}


def migration_audit_core_valid(
    payload: dict[str, Any],
    record_id: str,
    *,
    completion: str,
) -> bool:
    def sha256(value: Any) -> bool:
        return isinstance(value, str) and re.fullmatch(r"[0-9a-f]{64}", value) is not None

    def absolute_path(value: Any) -> bool:
        return isinstance(value, str) and bool(value) and Path(value).is_absolute()

    def nonnegative_integer(value: Any) -> bool:
        return isinstance(value, int) and not isinstance(value, bool) and value >= 0

    source = payload.get("source")
    health = source.get("health") if isinstance(source, dict) else None
    sheet_rows = health.get("sheet_rows") if isinstance(health, dict) else None
    return bool(
        isinstance(record_id, str)
        and record_id
        and payload.get("migration_id") == record_id
        and payload.get("completion") == completion
        and absolute_path(payload.get("target_database"))
        and absolute_path(payload.get("target_workbook"))
        and isinstance(source, dict)
        and absolute_path(source.get("path"))
        and nonnegative_integer(source.get("size"))
        and nonnegative_integer(source.get("mtime_ns"))
        and sha256(source.get("sha256"))
        and isinstance(health, dict)
        and sha256(health.get("sha256"))
        and isinstance(sheet_rows, dict)
        and all(
            isinstance(name, str) and nonnegative_integer(count)
            for name, count in sheet_rows.items()
        )
    )


def _migration_summary_valid(payload: dict[str, Any]) -> bool:
    sheets = payload.get("sheets")
    warning_codes = payload.get("warning_codes")
    imported_records = payload.get("imported_records")
    count_fields = {"rows_read", "importable", "skipped", "conflicts", "warnings"}

    def nonnegative_integer(value: Any) -> bool:
        return isinstance(value, int) and not isinstance(value, bool) and value >= 0

    return bool(
        isinstance(sheets, dict)
        and bool(sheets)
        and all(
            isinstance(name, str)
            and isinstance(summary, dict)
            and count_fields.issubset(summary)
            and all(nonnegative_integer(summary[field]) for field in count_fields)
            for name, summary in sheets.items()
        )
        and isinstance(warning_codes, list)
        and all(isinstance(code, str) for code in warning_codes)
        and isinstance(imported_records, dict)
        and all(
            isinstance(kind, str) and nonnegative_integer(count)
            for kind, count in imported_records.items()
        )
    )


def _migration_resume_payload_valid(
    payload: dict[str, Any],
    *,
    migration_id: str,
    completion: str,
    fingerprint: dict[str, Any],
    config: Config,
) -> bool:
    if not migration_audit_core_valid(
        payload, migration_id, completion=completion
    ) or not _migration_summary_valid(payload):
        return False
    stored_source = payload.get("source")
    content_fields = ("path", "size", "sha256", "health")
    stored_content = {
        field: stored_source.get(field)
        for field in content_fields
        if isinstance(stored_source, dict)
    }
    current_content = {field: fingerprint.get(field) for field in content_fields}
    if (
        canonical_json(stored_content) != canonical_json(current_content)
        or payload.get("target_database") != str(config.database)
        or payload.get("target_workbook") != str(config.workbook)
    ):
        return False
    backup_value = payload.get("verified_source_backup")
    if not isinstance(backup_value, str) or not Path(backup_value).is_absolute():
        return False
    backup = Path(backup_value).resolve()
    if backup.parent != (config.home_path / "backups").resolve():
        return False
    return backup.is_file() and _sha256(backup) == fingerprint["sha256"]


def _verified_source_backup(config: Config, source: Path, digest: str) -> Path:
    directory = config.home_path / "backups"
    directory.mkdir(parents=True, exist_ok=True)
    suffix = digest[:16]
    for existing in sorted(directory.glob(f"legacy-source-*-{suffix}.xlsx")):
        if _sha256(existing) == digest:
            return existing

    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    destination = directory / f"legacy-source-{stamp}-{suffix}.xlsx"
    fd, raw_temp = tempfile.mkstemp(
        prefix=".legacy-source-", suffix=".tmp.xlsx", dir=directory
    )
    os.close(fd)
    temp = Path(raw_temp)
    try:
        shutil.copy2(source, temp)
        try:
            temp.chmod(0o600)
        except OSError:
            pass
        if _sha256(temp) != digest:
            raise RuntimeError("legacy source backup verification failed")
        os.replace(temp, destination)
        if _sha256(destination) != digest:
            destination.unlink(missing_ok=True)
            raise RuntimeError("legacy source backup verification failed")
        return destination
    finally:
        temp.unlink(missing_ok=True)


def migrate_legacy_workbook(
    config: Config,
    database: HealthDatabase,
    source: Path,
    *,
    backup_source: bool,
    template: Path,
) -> dict[str, Any]:
    if not backup_source:
        raise ValueError("a verified source backup is required for a real legacy migration")
    source = validate_legacy_source(config, source)

    fingerprint = _source_fingerprint(source)
    migration_id = _migration_id(fingerprint)
    completed = _completed_migration(database, migration_id)
    if completed is not None:
        if not _migration_resume_payload_valid(
            completed,
            migration_id=migration_id,
            completion="complete",
            fingerprint=fingerprint,
            config=config,
        ):
            raise ValueError("completed legacy migration audit is invalid")
        return {
            "status": "already_migrated",
            "target_database": completed["target_database"],
            "target_workbook": completed["target_workbook"],
            "migration_id": migration_id,
        }

    backup = _verified_source_backup(config, source, fingerprint["sha256"])
    pending = _pending_migration(database, migration_id)
    if pending is not None and not _migration_resume_payload_valid(
        pending,
        migration_id=migration_id,
        completion="pending_postconditions",
        fingerprint=fingerprint,
        config=config,
    ):
        raise ValueError("pending legacy migration audit is invalid")
    if pending is not None:
        pending_imported = pending.get("imported_records")
        expected_witnesses = sum(pending_imported.values())
        if not database.migration_record_witnesses_valid(
            migration_id, expected_witnesses
        ):
            raise ValueError(
                "pending legacy migration records no longer match the imported batch"
            )
    existing_records = [
        (kind, payload)
        for kind in _KIND_SHEETS
        for payload in database.list_records(kind)
    ]
    plan = _assess_plan(parse_legacy_workbook(backup), existing_records)
    if plan["conflicts"]:
        raise LegacyMigrationConflictError(
            f"legacy migration has {len(plan['conflicts'])} unresolved conflict(s)"
        )
    if pending is not None and plan["records"]:
        raise ValueError("pending legacy migration records no longer match the imported batch")
    if pending is None:
        imported: dict[str, int] = {}
        for kind, _payload in plan["records"]:
            imported[kind] = imported.get(kind, 0) + 1
        audit_payload = {
            "migration_id": migration_id,
            "completion": "pending_postconditions",
            "source": fingerprint,
            "verified_source_backup": str(backup),
            "target_database": str(config.database),
            "target_workbook": str(config.workbook),
            "sheets": plan["sheets"],
            "warning_codes": sorted({item["code"] for item in plan["warnings"]}),
            "imported_records": imported,
        }
        database.upsert_many_with_audit(
            plan["records"],
            audit_action="legacy_workbook_migration_pending",
            audit_record_id=migration_id,
            audit_payload=audit_payload,
            witness_migration_id=migration_id,
        )
    else:
        audit_payload = pending
        imported = dict(pending.get("imported_records") or {})
    database.prune_history(
        audit_limit=config.audit_event_retention,
        sync_limit=config.sync_run_retention,
    )
    workbook = export_workbook(config, database, template)
    sqlite_integrity = database.connection.execute("PRAGMA quick_check").fetchone()[0]
    if sqlite_integrity != "ok":
        raise RuntimeError("post-migration SQLite integrity check failed")
    context = build_context(config, database)
    if not isinstance(context, dict) or not {
        "generated_at",
        "selected_date",
        "today",
        "trends",
    }.issubset(context):
        raise RuntimeError("post-migration context integrity check failed")
    audit_payload = dict(audit_payload)
    audit_payload["completion"] = "complete"
    database.finalize_audit_event(
        record_id=migration_id,
        pending_action="legacy_workbook_migration_pending",
        completed_action="legacy_workbook_migration",
        payload=audit_payload,
    )
    return {
        "status": "migrated",
        "migration_id": migration_id,
        "verified_source_backup": str(backup),
        "target_database": str(config.database),
        "target_workbook": str(workbook),
        "sheets": audit_payload["sheets"],
        "warnings": plan["warnings"],
        "imported_records": imported,
        "integrity": {"sqlite": "ok", "context": "ok"},
    }
