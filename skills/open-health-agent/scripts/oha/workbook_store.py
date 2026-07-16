from __future__ import annotations

import json
import hashlib
import os
import shutil
import tempfile
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.table import Table, TableStyleInfo

from .config import Config
from .constants import (
    DAILY_FIELD_MAP,
    HEALTH_NOTES,
    SHEET_HEADERS,
    SUMMARY_NUTRIENT_FIELDS,
)
from .database import HealthDatabase
from .energy import tef_from_macros

HEADER_FILL = "1F4E78"
INPUT_FILL = "FFF2CC"
TABLE_NAMES = {
    "健康日报": "DailyHealthTable",
    "健康测量": "MeasurementsTable",
    "训练记录": "WorkoutsTable",
    "饮食记录": "FoodLogTable",
    "饮食营养明细": "NutrientsTable",
    "每日营养汇总": "DailyNutritionTable",
    "目标历史": "GoalsTable",
    "同步日志": "SyncRunsTable",
    "健康说明": "HealthNotesTable",
}


def _xml_safe_text(value: str) -> str:
    """Replace characters forbidden by XML 1.0 while keeping SQLite exact.

    XLSX cells are serialized as XML. User speech/OCR can contain C0 control
    characters that SQLite and JSON accept but openpyxl correctly refuses to
    serialize. The readable workbook uses U+FFFD for each invalid code point;
    the canonical record in SQLite remains byte-for-byte unchanged.
    """

    output: list[str] = []
    for character in value:
        codepoint = ord(character)
        valid = (
            codepoint in (0x09, 0x0A, 0x0D)
            or 0x20 <= codepoint <= 0xD7FF
            or 0xE000 <= codepoint <= 0xFFFD
            or 0x10000 <= codepoint <= 0x10FFFF
        ) and codepoint not in (0xFFFE, 0xFFFF)
        output.append(character if valid else "\uFFFD")
    return "".join(output)


def as_excel_text(value: str) -> str:
    """Make externally controlled text safe for XLSX XML and formulas.

    The canonical, exact value remains in SQLite. Excel receives an escaped
    display value for any leading formula/control marker, including after
    whitespace that spreadsheet software may ignore.
    """
    safe_value = _xml_safe_text(value)
    stripped = safe_value.lstrip(" \t\r\n")
    if stripped.startswith(("=", "+", "-", "@")):
        prefix_length = len(safe_value) - len(stripped)
        return safe_value[:prefix_length] + "'" + stripped
    return safe_value


def as_excel_date(value: Any) -> Any:
    if isinstance(value, str) and len(value) >= 10:
        try:
            return date.fromisoformat(value[:10])
        except ValueError:
            return value
    return value


def style_sheet(ws, headers: list[str]) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 28
    for index, header in enumerate(headers, 1):
        cell = ws.cell(1, index, header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        width = min(max(len(str(header)) * 1.7 + 4, 12), 26)
        ws.column_dimensions[get_column_letter(index)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    for index, header in enumerate(headers, 1):
        column = get_column_letter(index)
        if header == "日期" or header == "生效日期" or header == "检索日期":
            for cell in ws[column][1:]:
                cell.number_format = "yyyy-mm-dd"
        if header in {"原话", "备注", "错误摘要", "估算说明", "说明", "安全约束"}:
            ws.column_dimensions[column].width = 38 if header != "说明" else 80
            for cell in ws[column][1:]:
                cell.alignment = Alignment(vertical="top", wrap_text=True)


def ensure_health_sheet(workbook, name: str, headers: list[str]):
    if name not in workbook.sheetnames:
        ws = workbook.create_sheet(name)
        ws.append(headers)
    else:
        ws = workbook[name]
        existing = [ws.cell(1, col).value for col in range(1, len(headers) + 1)]
        if any(existing) and existing != headers:
            raise RuntimeError(f"sheet {name} has incompatible headers; refusing to overwrite")
        extras = [ws.cell(1, col).value for col in range(len(headers) + 1, ws.max_column + 1)]
        if any(value not in (None, "") for value in extras):
            raise RuntimeError(
                f"sheet {name} contains unmanaged columns; move custom fields to a separate sheet before export"
            )
        if not any(existing):
            ws.append(headers)
    for table_name in list(ws.tables.keys()):
        del ws.tables[table_name]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    return ws


def append_dict_rows(ws, headers: list[str], rows: list[dict[str, Any]]) -> None:
    for row in rows:
        values: list[Any] = []
        for header in headers:
            value = row.get(header)
            if header in {"日期", "生效日期", "检索日期"}:
                value = as_excel_date(value)
            if isinstance(value, (dict, list)):
                value = json.dumps(value, ensure_ascii=False, sort_keys=True)
            if isinstance(value, str):
                value = as_excel_text(value)
            values.append(value)
        ws.append(values)


def add_table(ws, headers: list[str], name: str) -> None:
    if ws.max_row < 2:
        return
    reference = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
    table = Table(displayName=name, ref=reference)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def daily_rows(database: HealthDatabase) -> list[dict[str, Any]]:
    output: list[dict[str, Any]] = []
    for payload in database.list_records("daily"):
        output.append({header: payload.get(field) for field, header in DAILY_FIELD_MAP.items()})
    return output


def measurement_rows(database: HealthDatabase) -> list[dict[str, Any]]:
    output = []
    mapping = {
        "record_id": "记录ID",
        "date": "日期",
        "time": "时间",
        "metric": "指标",
        "value": "数值",
        "second_value": "第二数值",
        "unit": "单位",
        "source": "来源",
        "method": "记录方式",
        "confidence": "置信度",
        "original_text": "原话",
        "notes": "备注",
        "imported_at": "导入时间",
    }
    for payload in database.list_records("measurement"):
        output.append({header: payload.get(field) for field, header in mapping.items()})
    return output


def workout_rows(database: HealthDatabase) -> list[dict[str, Any]]:
    mapping = {
        "record_id": "记录ID",
        "date": "日期",
        "start_time": "开始时间",
        "workout_type": "训练类型",
        "duration_minutes": "时长_min",
        "distance_km": "距离_km",
        "calories_kcal": "热量_kcal",
        "average_heart_rate_bpm": "平均心率_bpm",
        "max_heart_rate_bpm": "最大心率_bpm",
        "intensity_rpe": "强度/RPE",
        "muscle_groups": "训练肌群",
        "training_source": "训练来源",
        "external_id": "外部ID",
        "method": "记录方式",
        "confidence": "置信度",
        "original_text": "原话",
        "notes": "备注",
        "imported_at": "导入时间",
    }
    return [
        {header: payload.get(field) for field, header in mapping.items()}
        for payload in database.list_records("workout")
    ]


def food_rows(database: HealthDatabase) -> list[dict[str, Any]]:
    mapping = {
        "record_id": "记录ID",
        "date": "日期",
        "time": "时间",
        "meal": "餐次",
        "food_name": "食物",
        "estimated_grams": "食用量估计_g",
        "grams_low": "食用量下限_g",
        "grams_high": "食用量上限_g",
        "source": "数据来源",
        "estimation_notes": "估算说明",
        "confidence": "置信度",
        "image_reference": "图片引用",
        "original_text": "原话",
        "notes": "备注",
        "recorded_at": "录入时间",
    }
    output = []
    for payload in database.list_records("food"):
        row = {header: payload.get(field) for field, header in mapping.items()}
        nutrients = payload.get("summary_nutrients") or {}
        if isinstance(nutrients, dict):
            for field, header in SUMMARY_NUTRIENT_FIELDS.items():
                row[header] = nutrients.get(field)
        output.append(row)
    return output


def nutrient_rows(database: HealthDatabase) -> list[dict[str, Any]]:
    output = []
    for payload in database.list_records("food"):
        record_id = payload.get("record_id")
        nutrients = payload.get("nutrients") or []
        if not isinstance(nutrients, list):
            continue
        for nutrient in nutrients:
            if not isinstance(nutrient, dict):
                continue
            output.append(
                {
                    "记录ID": record_id,
                    "营养素ID": nutrient.get("id"),
                    "营养素": nutrient.get("name"),
                    "数值": nutrient.get("amount"),
                    "单位": nutrient.get("unit"),
                    "数据来源": nutrient.get("source") or payload.get("source"),
                    "每100g或每份": nutrient.get("basis"),
                    "检索日期": nutrient.get("retrieved_at"),
                    "备注": nutrient.get("notes"),
                }
            )
    return output


def nutrition_summary_rows(database: HealthDatabase) -> list[dict[str, Any]]:
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for payload in database.list_records("food"):
        if payload.get("date"):
            grouped[str(payload["date"])].append(payload)
    rows = []
    for day, foods in sorted(grouped.items()):
        fields = [
            "energy_kcal",
            "protein_g",
            "fat_g",
            "carbohydrate_g",
            "fiber_g",
            "sugar_g",
            "sodium_mg",
            "potassium_mg",
            "calcium_mg",
            "iron_mg",
            "magnesium_mg",
        ]
        totals: dict[str, float | None] = {}
        coverage: dict[str, int] = {}
        for field in fields:
            values = []
            for food in foods:
                nutrients = food.get("summary_nutrients") or {}
                value = nutrients.get(field) if isinstance(nutrients, dict) else None
                if isinstance(value, (int, float)):
                    values.append(float(value))
            totals[field] = round(sum(values), 4) if values else None
            coverage[field] = len(values)
        complete_macros = min(coverage["protein_g"], coverage["carbohydrate_g"], coverage["fat_g"])
        tef = None
        if complete_macros == len(foods):
            tef = tef_from_macros(
                totals.get("protein_g") or 0,
                totals.get("carbohydrate_g") or 0,
                totals.get("fat_g") or 0,
            )
        coverage_labels = {
            "energy_kcal": "热量",
            "protein_g": "蛋白质",
            "fat_g": "脂肪",
            "carbohydrate_g": "碳水",
            "fiber_g": "纤维",
            "sodium_mg": "钠",
            "potassium_mg": "钾",
            "calcium_mg": "钙",
            "iron_mg": "铁",
            "magnesium_mg": "镁",
        }
        coverage_text = "、".join(
            f"{coverage_labels[field]} {coverage[field]}/{len(foods)}" for field in coverage_labels
        )
        notes = (
            f"已记录 {len(foods)} 项；字段覆盖(有值/条目)：{coverage_text}；"
            f"三大营养素完整项至少 {complete_macros}/{len(foods)}。总计仅代表已记录摄入。"
        )
        rows.append(
            {
                "日期": day,
                "已记录条目数": len(foods),
                "已记录热量_kcal": totals["energy_kcal"],
                "蛋白质_g": totals["protein_g"],
                "脂肪_g": totals["fat_g"],
                "碳水_g": totals["carbohydrate_g"],
                "膳食纤维_g": totals["fiber_g"],
                "糖_g": totals["sugar_g"],
                "钠_mg": totals["sodium_mg"],
                "钾_mg": totals["potassium_mg"],
                "钙_mg": totals["calcium_mg"],
                "铁_mg": totals["iron_mg"],
                "镁_mg": totals["magnesium_mg"],
                "估算TEF下限_kcal": round(tef.low, 2) if tef is not None else None,
                "估算TEF中值_kcal": round(tef.midpoint, 2) if tef is not None else None,
                "估算TEF上限_kcal": round(tef.high, 2) if tef is not None else None,
                "记录覆盖说明": notes,
                "更新时间": datetime.now().isoformat(timespec="seconds"),
            }
        )
    return rows


def goal_rows(database: HealthDatabase) -> list[dict[str, Any]]:
    mapping = {
        "record_id": "目标ID",
        "effective_date": "生效日期",
        "status": "状态",
        "priority": "优先级",
        "original_text": "用户原话",
        "safety_constraints": "安全约束",
        "recorded_at": "记录时间",
    }
    return [
        {header: payload.get(field) for field, header in mapping.items()}
        for payload in database.list_records("goal")
    ]


def sync_rows(database: HealthDatabase) -> list[dict[str, Any]]:
    output = []
    for run in reversed(database.list_sync_runs(limit=500)):
        output.append(
            {
                "批次ID": run["batch_id"],
                "运行时间": run["started_at"],
                "查询起始日期": run["from_date"],
                "查询结束日期": run["to_date"],
                "日报记录数": run["daily_count"],
                "训练记录数": run["workout_count"],
                "测量记录数": run["measurement_count"],
                "状态": run["status"],
                "数据截止时间": run["data_until"],
                "错误摘要": " | ".join(run["errors"]),
            }
        )
    return output


def health_notes_rows() -> list[dict[str, Any]]:
    return [{"项目": item, "说明": explanation} for item, explanation in HEALTH_NOTES]


def workbook_payloads(database: HealthDatabase) -> dict[str, list[dict[str, Any]]]:
    return {
        "健康日报": daily_rows(database),
        "健康测量": measurement_rows(database),
        "训练记录": workout_rows(database),
        "饮食记录": food_rows(database),
        "饮食营养明细": nutrient_rows(database),
        "每日营养汇总": nutrition_summary_rows(database),
        "目标历史": goal_rows(database),
        "同步日志": sync_rows(database),
        "健康说明": health_notes_rows(),
    }


def prune_backups(directory: Path, prefix: str, keep: int) -> None:
    backups = sorted(
        directory.glob(f"{prefix}-*.xlsx"),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    for old in backups[max(keep, 0) :]:
        old.unlink(missing_ok=True)


def inspect_workbook_schema(path: Path) -> dict[str, Any]:
    """Read-only validation for doctor; never returns cell content or file paths."""

    if path.suffix.lower() != ".xlsx":
        return {"ok": False, "reason": "unsupported_format"}
    if not path.is_file():
        return {"ok": False, "reason": "missing"}
    try:
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        return {
            "ok": False,
            "reason": "unreadable",
            "error_type": type(exc).__name__,
        }
    try:
        missing = [name for name in SHEET_HEADERS if name not in workbook.sheetnames]
        incompatible: list[str] = []
        unmanaged_columns: list[str] = []
        for name, headers in SHEET_HEADERS.items():
            if name not in workbook.sheetnames:
                continue
            sheet = workbook[name]
            actual = [sheet.cell(1, column).value for column in range(1, len(headers) + 1)]
            if actual != headers:
                incompatible.append(name)
            extras = [
                sheet.cell(1, column).value
                for column in range(len(headers) + 1, sheet.max_column + 1)
            ]
            if any(value not in (None, "") for value in extras):
                unmanaged_columns.append(name)
        return {
            "ok": not missing and not incompatible and not unmanaged_columns,
            "reason": "valid"
            if not missing and not incompatible and not unmanaged_columns
            else "managed_schema_mismatch",
            "missing_managed_sheets": missing,
            "incompatible_managed_sheets": incompatible,
            "managed_sheets_with_extra_columns": unmanaged_columns,
        }
    finally:
        workbook.close()


def export_workbook(config: Config, database: HealthDatabase, template: Path) -> Path:
    destination = config.workbook
    if destination.suffix.lower() != ".xlsx":
        raise ValueError("the managed workbook must use .xlsx; macro-enabled .xlsm files are not supported")
    destination.parent.mkdir(parents=True, exist_ok=True)
    backup_dir = config.home_path / "backups"
    backup_dir.mkdir(parents=True, exist_ok=True)

    original_non_health_rows: dict[str, int] = {}
    if destination.exists():
        source_check = load_workbook(destination, read_only=True, data_only=False)
        try:
            original_non_health_rows = {}
            for name in source_check.sheetnames:
                if name in SHEET_HEADERS:
                    continue
                sheet = source_check[name]
                if sheet.max_row is None:
                    sheet.calculate_dimension(force=True)
                original_non_health_rows[name] = sheet.max_row or 0
        finally:
            source_check.close()

    fd, raw_temp = tempfile.mkstemp(
        prefix=f".{destination.stem}.open-health-agent.",
        suffix=".tmp.xlsx",
        dir=destination.parent,
    )
    os.close(fd)
    temp = Path(raw_temp)
    try:
        if destination.exists():
            shutil.copy2(destination, temp)
        elif template.exists():
            shutil.copy2(template, temp)
        else:
            workbook = Workbook()
            try:
                workbook.active.title = "使用说明"
                workbook.save(temp)
            finally:
                workbook.close()
        # ``copy2`` carries the source mode. Tighten the temporary file before
        # any private database rows are serialized into it; otherwise a public
        # template or pre-existing workbook can create a brief 0644 health file
        # in a shared/synchronized directory.
        os.chmod(temp, 0o600)

        workbook = load_workbook(temp)
        try:
            payloads = workbook_payloads(database)
            for name, headers in SHEET_HEADERS.items():
                ws = ensure_health_sheet(workbook, name, headers)
                append_dict_rows(ws, headers, payloads[name])
                style_sheet(ws, headers)
                add_table(ws, headers, TABLE_NAMES[name])
            if workbook.calculation is None:
                workbook.calculation = CalcProperties(calcMode="auto")
            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.forceFullCalc = True
            workbook.save(temp)
        finally:
            workbook.close()

        check = load_workbook(temp, read_only=True, data_only=False)
        try:
            missing = [name for name in SHEET_HEADERS if name not in check.sheetnames]
            if missing:
                raise RuntimeError(f"temporary workbook is missing sheets: {', '.join(missing)}")
            for name, row_count in original_non_health_rows.items():
                if name not in check.sheetnames or check[name].max_row < row_count:
                    raise RuntimeError(f"non-health sheet changed unexpectedly: {name}")
            for name, headers in SHEET_HEADERS.items():
                actual = [check[name].cell(1, col).value for col in range(1, len(headers) + 1)]
                if actual != headers:
                    raise RuntimeError(f"header validation failed: {name}")
        finally:
            check.close()

        backup_prefix = "oha-" + hashlib.sha256(str(destination).encode("utf-8")).hexdigest()[:16]
        if destination.exists():
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
            backup = backup_dir / f"{backup_prefix}-{timestamp}.xlsx"
            shutil.copy2(destination, backup)
            try:
                os.chmod(backup, 0o600)
            except OSError:
                pass
        with temp.open("rb") as handle:
            try:
                os.fsync(handle.fileno())
            except OSError:
                pass
        os.replace(temp, destination)
        try:
            os.chmod(destination, 0o600)
        except OSError:
            pass
        try:
            directory_descriptor = os.open(destination.parent, os.O_RDONLY)
            try:
                os.fsync(directory_descriptor)
            finally:
                os.close(directory_descriptor)
        except OSError:
            pass
        try:
            # The private backups directory is application-managed. Prune all
            # OHA-prefixed workbook backups so changing the workbook path does
            # not leave an unbounded set under obsolete path hashes.
            prune_backups(backup_dir, "oha", config.backup_retention)
        except OSError:
            # The workbook has already been safely replaced. A retention
            # cleanup failure must not misreport the export as failed.
            pass
        return destination
    finally:
        if temp.exists():
            temp.unlink()
