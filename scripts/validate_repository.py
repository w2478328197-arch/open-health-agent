#!/usr/bin/env python3
"""Validate the portable Skill package and its public, privacy-safe assets."""

from __future__ import annotations

import json
import py_compile
import re
import subprocess
import sys
from pathlib import Path
from urllib.parse import unquote


ROOT = Path(__file__).resolve().parents[1]
SKILL_ROOT = ROOT / "skills" / "open-health-agent"

REQUIRED_FILES = (
    ROOT / "README.md",
    ROOT / "README_EN.md",
    ROOT / "LICENSE",
    ROOT / "PRIVACY.md",
    ROOT / "PRIVACY.zh-CN.md",
    ROOT / "SECURITY.md",
    ROOT / "SECURITY.zh-CN.md",
    ROOT / "install.sh",
    ROOT / "requirements-dev.txt",
    SKILL_ROOT / "SKILL.md",
    SKILL_ROOT / "agents" / "openai.yaml",
    SKILL_ROOT / "assets" / "health-ledger.xlsx",
    SKILL_ROOT / "assets" / "AGENTS.md.template",
    SKILL_ROOT / "assets" / "profile.example.json",
    SKILL_ROOT / "scripts" / "health_agent.py",
    SKILL_ROOT / "scripts" / "requirements.txt",
    ROOT / "tests" / "fixtures" / "README.md",
)

EXPECTED_HEALTH_SHEETS = {
    "健康日报",
    "健康测量",
    "训练记录",
    "饮食记录",
    "饮食营养明细",
    "每日营养汇总",
    "目标历史",
    "同步日志",
    "健康说明",
}


class Validation:
    def __init__(self) -> None:
        self.errors: list[str] = []

    def require(self, condition: bool, message: str) -> None:
        if not condition:
            self.errors.append(message)


def parse_frontmatter(path: Path, validation: Validation) -> tuple[dict[str, str], str]:
    text = path.read_text(encoding="utf-8")
    validation.require(text.startswith("---\n"), "SKILL.md must start with YAML frontmatter")
    parts = text.split("---", 2)
    if len(parts) != 3:
        validation.errors.append("SKILL.md frontmatter is not closed")
        return {}, text
    raw, body = parts[1], parts[2]
    metadata: dict[str, str] = {}
    for line in raw.splitlines():
        if not line.strip():
            continue
        match = re.fullmatch(r"([A-Za-z0-9_-]+):\s*(.*)", line)
        if not match:
            validation.errors.append(f"unsupported SKILL.md frontmatter line: {line!r}")
            continue
        metadata[match.group(1)] = match.group(2).strip().strip("'\"")
    return metadata, body


def validate_skill(validation: Validation) -> None:
    path = SKILL_ROOT / "SKILL.md"
    if not path.exists():
        return
    metadata, body = parse_frontmatter(path, validation)
    validation.require(set(metadata) == {"name", "description"}, "SKILL.md frontmatter must contain only name and description")
    validation.require(metadata.get("name") == "open-health-agent", "SKILL.md name must be open-health-agent")
    description = metadata.get("description", "")
    validation.require(20 <= len(description) <= 1024, "SKILL.md description must be 20-1024 characters")
    validation.require("<" not in description and ">" not in description, "SKILL.md description cannot contain angle brackets")
    validation.require(len(body.splitlines()) <= 500, "SKILL.md body must stay at or below 500 lines")
    validation.require("TODO" not in body, "SKILL.md still contains a TODO marker")

    lowered = body.casefold()
    concept_groups = {
        "first-use explanation": ("先解释", "首次说明", "explain first", "explain the skill before", "before any setup"),
        "Google Health data": ("ghealth",),
        "goals in agent rules": ("agents.md",),
        "thermic effect of food": ("tef", "食物热效应"),
        "lean-mass energy basis": ("瘦体重", "lean mass"),
        "WHO fallback": ("who", "世界卫生组织"),
        "privacy disclosure": ("隐私", "privacy"),
        "ledger-before-advice rule": ("建议前", "before advice", "before every recommendation", "read context"),
    }
    for label, alternatives in concept_groups.items():
        validation.require(
            any(alternative.casefold() in lowered for alternative in alternatives),
            f"SKILL.md is missing required concept: {label}",
        )


def validate_openai_yaml(validation: Validation) -> None:
    path = SKILL_ROOT / "agents" / "openai.yaml"
    if not path.exists():
        return
    text = path.read_text(encoding="utf-8")
    for key in ("display_name:", "short_description:", "default_prompt:"):
        validation.require(key in text, f"agents/openai.yaml is missing {key[:-1]}")
    validation.require("$open-health-agent" in text, "default_prompt must explicitly mention $open-health-agent")


def validate_project_metadata(validation: Validation) -> None:
    constants_path = SKILL_ROOT / "scripts" / "oha" / "constants.py"
    if not constants_path.exists():
        return
    match = re.search(
        r'^APP_VERSION\s*=\s*["\']([^"\']+)["\']',
        constants_path.read_text(encoding="utf-8"),
        flags=re.MULTILINE,
    )
    validation.require(match is not None, "constants.py must declare APP_VERSION")
    if match is None:
        return

    version = match.group(1)
    for path in (ROOT / "README.md", ROOT / "README_EN.md"):
        if not path.exists():
            continue
        text = path.read_text(encoding="utf-8")
        validation.require(
            f"version-{version}" in text,
            f"{path.name} version badge must match APP_VERSION {version}",
        )
        validation.require("Python-3.10%2B" in text, f"{path.name} must show the Python 3.10+ requirement")
        validation.require("Apache--2.0" in text, f"{path.name} must show the Apache-2.0 license")
        validation.require("README_EN.md" in text and "README.md" in text, f"{path.name} must link both languages")
        validation.require("SECURITY" in text and "PRIVACY" in text, f"{path.name} must link security and privacy notices")


def validate_public_markdown_links(validation: Validation) -> None:
    paths = (
        ROOT / "README.md",
        ROOT / "README_EN.md",
        ROOT / "SECURITY.md",
        ROOT / "SECURITY.zh-CN.md",
        ROOT / "PRIVACY.md",
        ROOT / "PRIVACY.zh-CN.md",
    )
    for path in paths:
        if not path.exists():
            continue
        text = path.read_text(encoding="utf-8")
        for target in re.findall(r"\[[^\]]*\]\(([^)]+)\)", text):
            target = target.strip()
            if not target or target.startswith(("http://", "https://", "mailto:", "#")):
                continue
            local_target = unquote(target.split("#", 1)[0])
            if not local_target:
                continue
            validation.require(
                (path.parent / local_target).exists(),
                f"{path.name} has a broken relative link: {target}",
            )


def validate_profile_example(validation: Validation) -> None:
    path = SKILL_ROOT / "assets" / "profile.example.json"
    if not path.exists():
        return
    try:
        profile = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        validation.errors.append(f"profile.example.json is invalid JSON: {exc}")
        return
    validation.require(isinstance(profile, dict), "profile.example.json must contain an object")
    if not isinstance(profile, dict):
        return
    for identity_key in ("name", "email", "phone", "wechat", "address"):
        validation.require(not profile.get(identity_key), f"profile.example.json must not include {identity_key}")
    validation.require(profile.get("lean_mass_kg") is None, "profile.example.json lean_mass_kg must be null")
    validation.require(profile.get("goals", []) == [], "profile.example.json goals must be empty")
    validation.require(profile.get("health_constraints", []) == [], "profile.example.json health_constraints must be empty")


def validate_workbook(validation: Validation) -> None:
    path = SKILL_ROOT / "assets" / "health-ledger.xlsx"
    if not path.exists():
        return
    try:
        from openpyxl import load_workbook

        # Some spreadsheet writers omit cached worksheet dimensions. Normal
        # mode asks openpyxl to derive dimensions from cells instead of
        # returning max_column=None in read-only mode.
        workbook = load_workbook(path, read_only=False, data_only=False)
    except Exception as exc:
        validation.errors.append(f"health-ledger.xlsx cannot be opened: {exc}")
        return
    try:
        sys.path.insert(0, str(SKILL_ROOT / "scripts"))
        from oha.constants import SHEET_HEADERS

        validation.require(
            EXPECTED_HEALTH_SHEETS.issubset(workbook.sheetnames),
            "health-ledger.xlsx is missing one or more health sheets",
        )
        validation.require("使用说明" in workbook.sheetnames, "health-ledger.xlsx must include 使用说明")
        for name in EXPECTED_HEALTH_SHEETS:
            if name in workbook.sheetnames:
                validation.require(workbook[name].max_column >= 2, f"workbook sheet {name} has no usable schema")
                expected = SHEET_HEADERS[name]
                actual = [workbook[name].cell(1, column).value for column in range(1, len(expected) + 1)]
                validation.require(actual == expected, f"workbook sheet {name} headers do not match the runtime schema")
        validation.require(not getattr(workbook, "_external_links", []), "health-ledger.xlsx must not contain external links")
        if "能量估算" in workbook.sheetnames:
            energy_sheet = workbook["能量估算"]
            formulas = [
                cell.value
                for row in energy_sheet.iter_rows()
                for cell in row
                if isinstance(cell.value, str) and cell.value.startswith("=")
            ]
            validation.require(len(formulas) >= 7, "能量估算 must retain its documented formulas")
            validation.require(
                all("COUNT(B10:B12)<3" in str(energy_sheet[cell].value) for cell in ("B18", "B19", "B20")),
                "能量估算 TEF formulas must require all three recorded macros",
            )
    finally:
        workbook.close()


def validate_python(validation: Validation) -> None:
    for path in sorted((SKILL_ROOT / "scripts").rglob("*.py")):
        try:
            py_compile.compile(str(path), doraise=True)
        except py_compile.PyCompileError as exc:
            validation.errors.append(f"Python compile failed for {path.relative_to(ROOT)}: {exc}")


def validate_secret_scan(validation: Validation) -> None:
    completed = subprocess.run(
        [sys.executable, str(ROOT / "scripts" / "check_secrets.py")],
        cwd=ROOT,
        check=False,
        capture_output=True,
        text=True,
    )
    validation.require(
        completed.returncode == 0,
        "secret/privacy scan failed:\n" + (completed.stdout + completed.stderr).strip(),
    )


def main() -> int:
    validation = Validation()
    for path in REQUIRED_FILES:
        validation.require(path.exists(), f"missing required file: {path.relative_to(ROOT)}")
    validate_skill(validation)
    validate_openai_yaml(validation)
    validate_project_metadata(validation)
    validate_public_markdown_links(validation)
    validate_profile_example(validation)
    validate_workbook(validation)
    validate_python(validation)
    validate_secret_scan(validation)

    if validation.errors:
        print("Repository validation failed:")
        for error in validation.errors:
            print(f"- {error}")
        return 1
    print("Repository validation passed.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
